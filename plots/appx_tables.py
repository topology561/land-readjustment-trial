# -*- coding: utf-8 -*-
"""
W-P.003　附表 4、14、15、18、19、21、25（照手冊表樣·KL 2026-10-24 02:14／14:55）。

命令列：python plots/appx_tables.py --out <目錄>
輸出（皆寫至 --out，⛔ 入倉）：見 OUT_NAMES。
所據：出圖線/orders/W-P.003_附表4-14-15-18-19-21-25.md（§一、§二）。
  附表 4　抵充土地清冊：verify/case_params_UC9898.json 之 財務接線_v3.抵充地_m2（KL 確認 0）⇒ 無列、合計 0。
  附表 14　他項權利清冊：U_LAND 之「權利種類」（現存之他項權利）；示範資料無 ⇒ 無列。
  附表 15　限制登記清冊：匯入之地籍無限制登記之欄（「重劃區地籍」之「有無限制登記」皆空）⇒ 無列。
  附表 18　重劃前地價評議表：每宗一列；區段與宗地地價同 W-P.001（財務接線_v3·個位四捨五入·手冊第 2804〜2806 列）。
  附表 19　重劃後地價評議表：每街廓一列（財務接線_v3.後街廓_面積單價·全精度存值·顯示二位小數）。
  附表 21　重劃後可分配土地街廓面積地價計算表：每街廓一列；總價 ＝ 面積 × 單價（二位四捨五入）。
  附表 25　街角地側面道路負擔百分率累積表：手冊第 4048〜4078 列之值（法定之表·照錄）。
缺值一律報錯停止（⛔ 靜默補值）；匯入資料所無之欄留空（⛔ 杜撰）。
"""
import argparse
import json
import os
import sys

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO, "plots"))

import prelim_maps as PM  # noqa: E402  （讀取與表樣之共用函式·W-P.001／002）

CASE_PARAMS = PM.CASE_PARAMS
TITLE_HEAD = "○○縣(市)第○○期○○○○市地重劃區"
TITLE_HEAD_SHORT = "○○縣(市)第○○期○○市地重劃區"
SIGN_A4 = ("日期：", "編造：", "校對：", "承辦人：", "科長：")
SIGN_A14 = ("日期：", "承辦：", "股長：", "科長：", "局(處)長：")
SIGN_A21 = ("製表日期:　年　月　日", "承辦:", "股長:", "科長:")
NOTE_A14 = "附註：本表未列事項詳原土地登記簿記載"
REVIEW_DATE = "評議日期：　　年　　月　　日"

# 附表 25（手冊第 4053〜4076 列·W 公尺 → 十分位之 Rw %）
RW_TABLE = {
    1: (17.36, 18.08, 18.80, 19.52, 20.23, 20.94, 21.64, 22.35, 23.05, 23.75),
    2: (24.44, 25.14, 25.83, 26.52, 27.20, 27.88, 28.56, 29.24, 29.91, 30.58),
    3: (31.25, 31.92, 32.58, 33.24, 33.89, 34.55, 35.20, 35.85, 36.49, 37.14),
    4: (37.78, 38.42, 39.05, 39.68, 40.31, 40.94, 41.56, 42.18, 42.80, 43.42),
    5: (44.03, 44.64, 45.24, 45.85, 46.45, 47.05, 47.64, 48.24, 48.83, 49.42),
    6: (50.00, 50.58, 51.16, 51.74, 52.31, 52.88, 53.45, 54.02, 54.58, 55.14),
    7: (55.69, 56.25, 56.80, 57.35, 57.89, 58.44, 58.98, 59.52, 60.05, 60.58),
    8: (61.11, 61.64, 62.16, 62.68, 63.20, 63.72, 64.23, 64.74, 65.24, 65.75),
    9: (66.25, 66.75, 67.24, 67.74, 68.23, 68.72, 69.20, 69.68, 70.16, 70.64),
    10: (71.11, 71.58, 72.05, 72.52, 72.98, 73.44, 73.89, 74.35, 74.80, 75.25),
    11: (75.69, 76.14, 76.58, 77.02, 77.45, 77.88, 78.31, 78.74, 79.16, 79.58),
    12: (80.00, 80.42, 80.83, 81.24, 81.64, 82.05, 82.45, 82.85, 83.24, 83.64),
    13: (84.03, 84.42, 84.80, 85.18, 85.56, 85.94, 86.31, 86.68, 87.05, 87.42),
    14: (87.78, 88.14, 88.49, 88.85, 89.20, 89.55, 89.89, 90.24, 90.58, 90.92),
    15: (91.25, 91.59, 91.91, 92.24, 92.56, 92.88, 93.20, 93.52, 93.83, 94.14),
    16: (94.44, 94.75, 95.05, 95.35, 95.64, 95.94, 96.23, 96.52, 96.80, 97.08),
    17: (97.36, 97.64, 97.91, 98.18, 98.45, 98.72, 98.98, 99.24, 99.49, 99.75),
    18: (100.00,),
}

OUT_NAMES = {
    "a4": "附表04_抵充土地清冊.xlsx",
    "a14": "附表14_他項權利清冊.xlsx",
    "a15": "附表15_限制登記清冊.xlsx",
    "a18": "附表18_重劃前地價評議表.xlsx",
    "a19": "附表19_重劃後地價評議表.xlsx",
    "a21": "附表21_重劃後可分配土地街廓面積地價計算表.xlsx",
    "a25": "附表25_街角地側面道路負擔百分率累積表.xlsx",
}


# ---------------------------------------------------------------- 共用
def _center():
    from openpyxl.styles import Alignment
    return Alignment(horizontal="center", vertical="center", wrap_text=True)


def _book():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    return wb, ws


def _title(ws, row, text, ncol):
    ws.cell(row, 1, text).alignment = _center()
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)


def _hdr(ws, r1, c1, r2, c2, text):
    ws.cell(r1, c1, text).alignment = _center()
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def _sign(ws, row, labels, ncol):
    step = max(1, ncol // len(labels))
    for i, lab in enumerate(labels):
        ws.cell(row, 1 + i * step, lab)


def load_params(path=CASE_PARAMS):
    d = json.load(open(path, encoding="utf-8"))
    v = d["財務接線_v3"]
    for k in ("抵充地_m2", "後街廓_面積單價"):
        if k not in v:
            raise PM.SourceError(f"case_params 之 財務接線_v3 缺 {k}")
    zoning = d["global"].get("分區")
    if PM._blank(zoning):
        raise PM.SourceError("case_params 之 global.分區 為空")
    return v, zoning


# ---------------------------------------------------------------- 附表 4
A4_NCOL = 10
A4_HEAD = ("鄉鎮市區", "段", "小段", "地號", "登記面積（㎡）", "抵充面積（㎡）", "權屬", "管理機關", "代管機關", "備註")


def write_a4(path, offset_area):
    if offset_area != 0:
        raise PM.SourceError(f"抵充地_m2＝{offset_area}≠0：本單唯據 KL 確認之 0（抵充之宗之清單須另單）")
    wb, ws = _book()
    ws.title = "抵充土地清冊"
    _title(ws, 1, TITLE_HEAD_SHORT + "抵充土地清冊", A4_NCOL)
    for c, lab in enumerate(A4_HEAD, start=1):
        ws.cell(2, c, lab).alignment = _center()
    PM._grid(ws, 2, 2, A4_NCOL)
    r = 3
    ws.cell(r, 1, "合計道路、溝渠、河川等土地抵充面積：")
    ws.cell(r, 5, 0)
    ws.cell(r, 6, "平方公尺")
    ws.cell(r + 1, 1, "未登記土地面積：")
    ws.cell(r + 1, 5, 0)
    ws.cell(r + 1, 6, "平方公尺")
    ws.cell(r + 2, 1, "總計抵充土地面積：")
    ws.cell(r + 2, 5, 0)
    ws.cell(r + 2, 6, "平方公尺")
    _sign(ws, r + 4, SIGN_A4, A4_NCOL)
    PM._page(ws, "1:2", landscape=True)
    wb.save(path)
    return []


# ---------------------------------------------------------------- 附表 14、15
A14_NCOL = 21
A15_NCOL = 23


def _owner_block(ws, row, ncol):
    """表頭之所有權人欄（每一所有權人一頁之式）：值留空，由使用者填。"""
    ws.cell(row, 1, "土地所有權人姓名")
    ws.cell(row, 5, "住址")
    ws.cell(row, 10, "國民身分證統一編號")
    ws.cell(row, ncol - 4, "分配清冊　第　冊　第　頁　號")


def rights_rows(owners):
    """現存之他項權利（U_LAND 之「權利種類」非空者）之列。"""
    out = []
    for ln in sorted(owners, key=PM.landno_key):
        for seq in owners[ln].values():
            for o in seq:
                if not PM._blank(o[PM.C_RIGHT_KIND]):
                    out.append((ln, o))
    return out


def write_a14(path, owners):
    rows = rights_rows(owners)
    if rows:
        raise PM.SourceError(f"有現存之他項權利 {len(rows)} 列：本單未定其填法 ⇒ 停機上呈")
    wb, ws = _book()
    ws.title = "他項權利清冊"
    _title(ws, 1, TITLE_HEAD + "他項權利清冊", A14_NCOL)
    _owner_block(ws, 2, A14_NCOL)
    _hdr(ws, 3, 1, 3, 14, "重劃前原載情形")
    _hdr(ws, 3, 15, 3, 20, "重劃後變更情形")
    _hdr(ws, 3, 21, 5, 21, "備註")
    _hdr(ws, 4, 1, 4, 6, "土地標示")
    _hdr(ws, 4, 7, 4, 8, "權利人")
    for c, lab in zip(range(9, 15), ("主登記次序", "權利種類", "權利範圍", "權利價值", "存續期間", "其他登記事項")):
        _hdr(ws, 4, c, 5, c, lab)
    _hdr(ws, 4, 15, 4, 19, "土地標示")
    _hdr(ws, 4, 20, 5, 20, "協議內容")
    for c, lab in enumerate(("鄉鎮市區", "段", "小段", "地號", "面積（㎡）", "持分比率", "姓名", "住址"), start=1):
        ws.cell(5, c, lab).alignment = _center()
    for c, lab in enumerate(("段", "小段", "地號", "面積（㎡）", "持分比率"), start=15):
        ws.cell(5, c, lab).alignment = _center()
    PM._grid(ws, 3, 5, A14_NCOL)
    ws.cell(7, 1, NOTE_A14)
    _sign(ws, 8, SIGN_A14, A14_NCOL)
    PM._page(ws, "1:5", landscape=True)
    wb.save(path)
    return rows


def write_a15(path):
    wb, ws = _book()
    ws.title = "限制登記清冊"
    _title(ws, 1, TITLE_HEAD + "限制登記清冊", A15_NCOL)
    _owner_block(ws, 2, A15_NCOL)
    _hdr(ws, 3, 1, 3, 12, "重劃前原載情形")
    _hdr(ws, 3, 13, 3, 22, "重劃後轉載情形")
    _hdr(ws, 3, 23, 5, 23, "備註")
    _hdr(ws, 4, 1, 4, 6, "土地標示")
    _hdr(ws, 4, 7, 4, 8, "權利人")
    for c, lab in zip(range(9, 13), ("登記種類", "權利範圍", "內容", "限制登記文號及囑託機關")):
        _hdr(ws, 4, c, 5, c, lab)
    _hdr(ws, 4, 13, 4, 17, "土地標示")
    _hdr(ws, 4, 18, 4, 19, "權利人")
    for c, lab in zip(range(20, 23), ("登記種類", "權利範圍", "協議內容")):
        _hdr(ws, 4, c, 5, c, lab)
    for c, lab in enumerate(("鄉鎮市區", "段", "小段", "地號", "面積（㎡）", "持分比率", "姓名", "國民身分證統一編號"),
                            start=1):
        ws.cell(5, c, lab).alignment = _center()
    for c, lab in enumerate(("段", "小段", "地號", "面積（㎡）", "持分比率", "姓名", "國民身分證統一編號"), start=13):
        ws.cell(5, c, lab).alignment = _center()
    PM._grid(ws, 3, 5, A15_NCOL)
    ws.cell(7, 1, NOTE_A14)
    _sign(ws, 8, SIGN_A14, A15_NCOL)
    PM._page(ws, "1:5", landscape=True)
    wb.save(path)
    return []


# ---------------------------------------------------------------- 附表 18
A18_NCOL = 9


def write_a18(path, lots, zone_of, price):
    wb, ws = _book()
    ws.title = "重劃前地價評議表"
    _title(ws, 1, "○○縣（市）第○○期○○○○市地重劃區重劃前地價評議表", A18_NCOL)
    ws.cell(2, 1, REVIEW_DATE)
    _hdr(ws, 3, 1, 4, 1, "區段編號")
    _hdr(ws, 3, 2, 4, 2, "查估之重劃前每平方公尺單價")
    _hdr(ws, 3, 3, 3, 6, "土地坐落")
    for c, lab in enumerate(("鄉鎮市區", "段", "小段", "地號"), start=3):
        ws.cell(4, c, lab).alignment = _center()
    _hdr(ws, 3, 7, 4, 7, "地價及標準地價評議委員會修正意見、具體事實及理由")
    _hdr(ws, 3, 8, 4, 8, "評議結果（元/平方公尺）")
    _hdr(ws, 3, 9, 4, 9, "備註")
    rows = []
    r = 4
    for z in sorted(set(zone_of[x["地號"]] for x in lots)):
        grp = sorted([x for x in lots if zone_of[x["地號"]] == z], key=lambda x: PM.landno_key(x["地號"]))
        r0 = r + 1
        for x in grp:
            r += 1
            vals = (z, price[z], x["鄉鎮市區"], x["段"], x.get("小段"), x["地號"])
            for c, v in enumerate(vals, start=1):
                if c <= 2 and r > r0:
                    continue
                ws.cell(r, c, v)
            rows.append(vals)
        for c in (1, 2):
            ws.cell(r0, c).alignment = _center()
            if r > r0:
                ws.merge_cells(start_row=r0, start_column=c, end_row=r, end_column=c)
        ws.cell(r0, 2).number_format = "#,##0"
    PM._grid(ws, 3, r, A18_NCOL)
    PM._page(ws, "1:4", landscape=True)
    wb.save(path)
    return rows


# ---------------------------------------------------------------- 附表 19
A19_NCOL = 7
A19_HEAD = ("區段編號", "街路或地段名稱", "起訖點或範圍", "查估之路線價或區段價（元/平方公尺）",
            "地價及標準地價評議委員會修正意見、具體事實及理由", "評議結果（元/平方公尺）", "備註")


def block_key(b):
    return (int(b[1:]) if b[:1] == "R" and b[1:].isdigit() else 10 ** 9, b)


def write_a19(path, post):
    wb, ws = _book()
    ws.title = "重劃後地價評議表"
    _title(ws, 1, "○○縣（市）第○○期○○○○市地重劃區重劃後地價評議表", A19_NCOL)
    ws.cell(2, 1, REVIEW_DATE)
    for c, lab in enumerate(A19_HEAD, start=1):
        ws.cell(3, c, lab).alignment = _center()
    rows = []
    r = 3
    for b in sorted(post, key=block_key):
        u = post[b]["單價_元每m2"]
        if PM._blank(u):
            raise PM.SourceError(f"後街廓 {b} 之單價為空")
        r += 1
        vals = (b, None, f"街廓 {b}", float(u))
        for c, v in enumerate(vals, start=1):
            ws.cell(r, c, v)
        ws.cell(r, 4).number_format = "#,##0.00"
        rows.append(vals)
    PM._grid(ws, 3, r, A19_NCOL)
    PM._page(ws, "1:3", landscape=True)
    wb.save(path)
    return rows


# ---------------------------------------------------------------- 附表 21
A21_NCOL = 7


def write_a21(path, post, zoning):
    wb, ws = _book()
    ws.title = "街廓面積地價計算表"
    _title(ws, 1, TITLE_HEAD_SHORT, A21_NCOL)
    _title(ws, 2, "重劃後可分配土地街廓面積地價計算表", A21_NCOL)
    ws.cell(3, A21_NCOL - 2, "第○可分配土地街廓面積地價圖")
    _hdr(ws, 4, 1, 5, 1, "街廓編號")
    _hdr(ws, 4, 2, 5, 2, "都市計畫使用分區")
    _hdr(ws, 4, 3, 4, 5, "規劃測量人員計算之面積")
    for c, lab in enumerate(("編號", "計算式", "平方公尺"), start=3):
        ws.cell(5, c, lab).alignment = _center()
    _hdr(ws, 4, 6, 5, 6, "單價 × 面積")
    _hdr(ws, 4, 7, 5, 7, "總價")
    rows = []
    r = 5
    for b in sorted(post, key=block_key):
        a, u = post[b]["面積_m2"], post[b]["單價_元每m2"]
        if PM._blank(a) or PM._blank(u):
            raise PM.SourceError(f"後街廓 {b} 之面積或單價為空")
        a, u = float(a), float(u)
        total = PM.round_half_up(a * u, 2)
        r += 1
        vals = (b, zoning, None, None, a, f"{u:,.2f} × {a:,.2f}", total)
        for c, v in enumerate(vals, start=1):
            ws.cell(r, c, v)
        ws.cell(r, 5).number_format = "#,##0.00"
        ws.cell(r, 7).number_format = "#,##0.00"
        rows.append(vals)
    PM._grid(ws, 4, r, A21_NCOL)
    _sign(ws, r + 2, SIGN_A21, A21_NCOL)
    PM._page(ws, "1:5", landscape=True)
    wb.save(path)
    return rows


# ---------------------------------------------------------------- 附表 25
A25_NCOL = 12
A25_COLHEAD = ("0.0", "1", "2", "3", "4", "5", "6", "7", "8", "9")


def write_a25(path):
    wb, ws = _book()
    ws.title = "Rw累積表"
    _title(ws, 1, "街角地側面道路負擔百分率累積表", A25_NCOL)
    ws.cell(3, 1, "RW").alignment = _center()
    _hdr(ws, 2, 2, 2, 11, "公寸")
    for c, lab in enumerate(A25_COLHEAD, start=2):
        ws.cell(3, c, lab).alignment = _center()
    _hdr(ws, 2, 12, 3, 12, "備註")
    rows = []
    r = 3
    for w in sorted(RW_TABLE):
        r += 1
        ws.cell(r, 1, w).alignment = _center()
        for j, v in enumerate(RW_TABLE[w]):
            ws.cell(r, 2 + j, v).number_format = "0.00"
            rows.append((w, j, v))
    _hdr(ws, 4, 12, r, 12, "公分以下用內插法。")
    PM._grid(ws, 2, r, A25_NCOL)
    ws.cell(r + 1, 1, "註：本表係由前面負擔百分率簡表推演而來。")
    PM._page(ws, "1:3", landscape=False)
    wb.save(path)
    return rows


def main(argv=None):
    ap = argparse.ArgumentParser(description="W-P.003 附表 4、14、15、18、19、21、25")
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    os.makedirs(a.out, exist_ok=True)

    lots, owners, area = PM.load_cadastre()
    zone_of, unit, price = PM.load_price_zones()
    miss = {x["地號"] for x in lots} - set(zone_of)
    if miss:
        raise PM.SourceError(f"case_params 之原地號_區段 缺：{sorted(miss)}")
    v, zoning = load_params()
    post = v["後街廓_面積單價"]

    p = {k: os.path.join(a.out, n) for k, n in OUT_NAMES.items()}
    write_a4(p["a4"], v["抵充地_m2"])
    r14 = write_a14(p["a14"], owners)
    write_a15(p["a15"])
    r18 = write_a18(p["a18"], lots, zone_of, price)
    r19 = write_a19(p["a19"], post)
    r21 = write_a21(p["a21"], post, zoning)
    r25 = write_a25(p["a25"])
    print(f"附表4 抵充 {v['抵充地_m2']} ㎡；附表14 {len(r14)} 列；附表15 0 列")
    print(f"附表18 {len(r18)} 列（區段 {sorted(set(zone_of.values()))}·{price}）")
    print(f"附表19 {len(r19)} 列；附表21 {len(r21)} 列·Σ面積 {round(sum(x[4] for x in r21), 2)}"
          f"·Σ總價 {round(sum(x[6] for x in r21), 2)}")
    print(f"附表25 {len(r25)} 格")
    for k, f in p.items():
        print(f"輸出 {f}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
