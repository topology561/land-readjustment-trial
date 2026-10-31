# -*- coding: utf-8 -*-
"""
W-P.006　附表 13 重劃前後土地分配清冊、附表 24 土地分配計算表（照手冊表樣·先做格式與可填之欄）。

命令列：python plots/post_tables.py --state <配地結果_退縮*.json> --moedict <dict-revised.json> --out <目錄>
          [--map <地號對照表.xlsx>]
  --state ＝ plots/alloc_state.py 之輸出（W-P.006 版：含「財務」「歸戶表」「現金補償」及分配地之 G 式各項）。
  --map   ＝ 承辦人改過之地號對照表（plots/post_maps.py 之輸出）；無 ⇒ S 型之預設（同 post_maps）。
輸出（皆寫至 --out，⛔ 入倉）：OUT_FMT 之二檔。

所據：出圖線/orders/W-P.006_附表13-24.md（§一、§二）。
  手冊（參考/市地重劃作業手冊.txt）第 2702〜2733 列（附表 13 之編造與歸戶）、第 2769 列（附表 13 之表樣）、
  第 3755〜3800 列（附表 24 之編造）、第 3813〜3828 列（附表 13 之重劃後部分·抵費地與共同負擔公共設施）、
  第 3990〜4045 列（附表 24 之範例：①＝G×C÷(1−C)、②＝Σ(a＋a')−G−①、③＝平均公告現值）。
  KL 2026-10-28 23:31 問二 甲（未受分配而發給現金補償之戶照列·重劃後各欄留空·備註「未達最小分配面積標準
  二分之一，發給現金補償 ○○ 元」）；問三 甲（平均公告現值以匯入之公告現值暫填·表下加註）；
  23:37（正式地號依對照表帶入）；10/26 23:14 第 1 點 甲（重劃前單價取個位、重劃後單價全精度計算、顯示二位）；
  10/30 19:28（W-P.005 全定版·下一批附表 13、24）。
  本版之配地 ＝ 階段一之終（調配之前）：未分配之戶之重劃後各欄留空；各宗之 a 之併入（a'）之逐筆來源
  管線未出 ⇒ 附表 24 之重劃前唯列原位次之地號，附表 13 之「取配面積」留空（俟調配之終態）。
W-P.006_補令一（KL 2026-10-31 08:14·出圖線/KL答覆/2026-10-31_0814.md 之 1〜8）：
  一 縣有之表依管理機關分表（「縣有（抵費地）」管理人＝花蓮縣政府〔手冊第 4501 列「縣（市）主管機關」·
     平均地權條例第 3 條〕；「縣有（道路）」等公共設施各一表，管理人〔手冊第 4498 列「各該公共設施主管機關」〕
     資料所無 ⇒ 留空）；
  三 取配面積：一戶之分配地之 a 之和 ＝ 該戶重劃前各宗登記面積之和（差 ≤ 0.01×宗數）⇒ 各宗皆整筆參加分配，
     附表 13 之取配面積 ＝ 登記面積；否則留空。附表 24：其戶唯一筆重劃前土地且如上 ⇒ 分配面積 ＝ Σa ＝ 本宗之 a、
     Σa' ＝ 0（手冊第 3762 列「Σa：原位次分配土地…逕以重劃前面積加總」）；他者分配面積、Σa、Σa' 留空並加註；
  四 附表 24 之重劃前單價唯列配地計算所用之值（區段地價·未取整·顯示二位）並加註其義；
  五 表下加註計算式之取位；七 ③ 之暫值之註；八 欄名依手冊（附表 24 範例·手冊第 3990〜4045 列）。
缺值一律報錯停止（⛔ 靜默補值）；匯入資料所無之欄留空（⛔ 杜撰）。
"""
import argparse
import json
import os
import sys
from decimal import Decimal, ROUND_HALF_UP

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO, "plots"))

import prelim_maps as PM  # noqa: E402  （地籍、歸戶、表樣之共用·W-P.001／002）
import owner_index  # noqa: E402
import post_maps as PMAP  # noqa: E402  （暫編、S 型之預設、對照表·W-P.005）

KIND_LOT, KIND_OFFSET = PMAP.KIND_LOT, PMAP.KIND_OFFSET
COUNTY = "花蓮縣"   # 抵費地與共同負擔公共設施用地之所有權人（手冊第 3819〜3821 列·登記為縣有）
NOTE_PUBLIC = "依平均地權條例第 60 條規定，由土地所有權人共同負擔取得土地"
NOTE_OFFSET = "依平均地權條例第 60 條規定，由土地所有權人以未建築土地折價抵付之土地"
NOTE_CASH = "未達最小分配面積標準二分之一，發給現金補償 {:,} 元"
NOTE_A13 = "附註：重劃後分配面積與實地指界交接土地面積如有不符，應以地政機關實際登記面積為準。"
NOTE_A24 = ("③暫以匯入之地籍資料所載之公告土地現值計（暫值）；依手冊，以重劃土地分配結果公告期滿時之"
            "當期公告土地現值為準")
NOTE_ROUND = "計算式各數為取位後之顯示值，與重劃後應分配權利面積之差係取位所致"
NOTE_PRE = ("評定之重劃前單價、原位置重劃前地價：配地計算所用之重劃前地價區段之區段地價（未取整、顯示至小數二位）；"
            "附表 18 之評定單價為其取至元之值")
NOTE_SPLIT = ("本宗之 a 含同戶他筆土地或本筆土地分屬他宗之部分，其逐筆之來源待配地程式提供，"
              "分配面積、Σa、Σa'暫空")
MGR_OFFSET = "花蓮縣政府"    # 抵費地之管理機關（手冊第 4501 列「縣（市）主管機關」·平均地權條例第 3 條）
LEGEND_A24 = "①＝G×C÷(1−C)；②＝(Σa＋Σa')−G−①；③＝平均公告現值"
SIGN_A13 = ("日期：", "承辦：", "股長：", "科長：", "局(處)長：")
SIGN_A24 = ("製表日期：", "承辦：", "股長：", "科長：", "局（處）長：")
TITLE_A13 = "○○縣(市)第○○期○○○○市地重劃區重劃前後土地分配清冊"
TITLE_A24 = "○○縣（市）第○○期○○○○市地重劃區土地分配計算表"
BOOK_A13 = "共　　冊　第　　冊　第　　頁"
SHEET_COUNTY = "縣有（{}）"
A13_NCOL = 24
A13_FIRST = 8                       # 資料之首列
A24_NCOL = 6
OUT_FMT = {"a13": "附表13_重劃前後土地分配清冊_退縮{}m.xlsx", "a24": "附表24_土地分配計算表_退縮{}m.xlsx"}
NF = {"area": "#,##0.00", "yuan": "#,##0", "p2": "#,##0.00", "k8": "0.00000000", "k4": "0.0000",
      "m2": "0.00", "pct": "0.00%"}

# 附表 13 之表頭（列, 欄起, 列迄, 欄迄, 字）
A13_HEAD = (
    (5, 1, 5, 10, "重劃前土地情形"), (5, 11, 7, 11, "扣除各項負擔後應分配面積(㎡)"),
    (5, 12, 5, 23, "重劃後土地情形"), (5, 24, 7, 24, "備註"),
    (6, 1, 7, 1, "鄉鎮市區"), (6, 2, 7, 2, "段"), (6, 3, 7, 3, "小段"), (6, 4, 7, 4, "地號"),
    (6, 5, 7, 5, "登記面積(㎡)"), (6, 6, 7, 6, "取配面積(㎡)"), (6, 7, 6, 8, "所有權人"),
    (6, 9, 7, 9, "持分比率"), (6, 10, 7, 10, "評定單價(元／㎡)"),
    (7, 7, 7, 7, "姓名"), (7, 8, 7, 8, "國民身分證統一編號"),
    (6, 12, 7, 12, "使用分區"), (6, 13, 7, 13, "段"), (6, 14, 7, 14, "小段"), (6, 15, 7, 15, "地號"),
    (6, 16, 7, 16, "面積(㎡)"), (6, 17, 7, 17, "評定單價(元／㎡)"), (6, 18, 6, 19, "預計繳領差額地價"),
    (7, 18, 7, 18, "繳（元）"), (7, 19, 7, 19, "領（元）"), (6, 20, 6, 21, "所有權人"),
    (7, 20, 7, 20, "姓名"), (7, 21, 7, 21, "國民身分證統一編號"), (6, 22, 7, 22, "持分比率"), (6, 23, 7, 23, "建號"),
)
A13_OWNER = (   # 列 3、4：所有權人與管理人（標籤之欄起迄, 值之欄起迄）
    (3, (1, 3, "土地所有權人（首列登記人）姓名"), (4, 6)), (3, (7, 8, "國民身分證統一編號"), (9, 10)),
    (3, (11, 12, "住址"), (13, 18)),
    (4, (1, 3, "管理人姓名"), (4, 6)), (4, (7, 8, "國民身分證統一編號"), (9, 10)),
)
A24_PRE = ("段", "小段", "地號", "登記面積(㎡)", "分配面積(㎡)", "共有人申請單獨分配持分比例",
           "持分面積(㎡)", "法定空地面積(㎡)", "當期公告現值(元/㎡)", "評定之重劃前單價(元/㎡)")
A24_CALC = ("Σa(㎡)", "Σa'(㎡)", "Σ(a＋a')(㎡)", "原位置重劃前地價(元/㎡)",
            "預配位置重劃前原地平均單價(元/㎡)", "重劃後宗地地價(元/㎡)", "A＝重劃後宗地地價／預配位置重劃前地價",
            "B", "C", "1−C", "S(m)", "W(m)", "F(m)", "l1", "l2", "Rw")
A24_POST = ("評定之重劃後單價(元/㎡)", "重劃費用負擔總額(元)", "重劃後應分配權利面積(㎡)", "實配面積(㎡)",
            "分配率＝G／a", "段", "小段", "地號", "面積(㎡)", "土地使用分區", "預計繳納(領取)差額地價總額")
A24_BURDEN_HEAD = (("抵費地負擔", "①面積(㎡)", "重劃後單價(元/㎡)", "負擔額(元)"),
                   ("公共設施用地共同負擔", "②面積(㎡)", "③平均公告現值(元/㎡)", "負擔額(元)"))
A24_BURDEN = ("合計(元)", "預計應繳納差額地價(元)", "預計應領取差額地價(元)")


def D(x):
    return Decimal(repr(float(x)))


def q(x, nd):
    return x.quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP)


# ---------------------------------------------------------------- 讀取與核對
def load_case_zoning():
    d = json.load(open(PM.CASE_PARAMS, encoding="utf-8"))
    z = d["global"].get("分區")
    if PM._blank(z):
        raise PM.SourceError("case_params 之 global.分區 為空")
    return z


def check_state(doc):
    for k in ("財務", "歸戶表", "現金補償"):
        if k not in doc:
            raise PM.SourceError(f"配地結果缺 {k}（須 W-P.006 版之 plots/alloc_state.py）")
    f = doc["財務"]
    for k in ("B", "C", "重劃後單價", "重劃前單價"):
        if k not in f:
            raise PM.SourceError(f"配地結果之財務缺 {k}")
    for x in doc["宗"]:
        if x["類別"] != KIND_LOT:
            continue
        for k in ("a", "a原位次", "併入面積", "重劃前區段", "F", "W", "Rw", "l側", "l正"):
            if k not in x:
                raise PM.SourceError(f"{x['配地暫編']} 缺 {k}")
        if x["街廓"] not in f["重劃後單價"] or x["重劃前區段"] not in f["重劃前單價"]:
            raise PM.SourceError(f"{x['配地暫編']} 之街廓或區段無單價")
        if q(D(x["a原位次"]) + D(x["併入面積"]), 2) != q(D(x["a"]), 2):
            raise PM.SourceError(f"{x['配地暫編']} 之 a ≠ 原位次＋併入")


def scheme(moedict, doc, ents):
    """⇒ (lots, owners, area, key_of, hh_no, hh_lots, code_key)：重劃前之地籍與歸戶（W-P.001／002 之法）
    及配地之歸戶代號 → 重劃前之戶（一一對應·不一即停）。"""
    strokes = owner_index.load_strokes(moedict)
    lots, owners, area = PM.load_cadastre()
    key_of, _order = PM.households(lots, owners)
    hh_lots, members = PM.household_members(owners, key_of)
    hh_no = owner_index.number_households(hh_lots, members, strokes, PM.landno_key)
    table = doc["歸戶表"]
    code_key = {}
    for ln, k in key_of.items():
        if ln not in table:
            raise PM.SourceError(f"歸戶表無 {ln}")
        if code_key.setdefault(table[ln], k) != k:
            raise PM.SourceError(f"歸戶代號 {table[ln]} 對應重劃前之數戶")
    if len(set(code_key.values())) != len(code_key):
        raise PM.SourceError("重劃前之一戶對應數歸戶代號")
    extra = sorted(set(table) - set(key_of), key=PM.landno_key)
    for ln in extra:
        if table[ln] not in code_key:
            raise PM.SourceError(f"歸戶表之 {ln} 之代號 {table[ln]} 無重劃前之戶")
    if extra:
        print(f"歸戶表另有重劃前地籍所無之號（不用）：{extra}")
    for e in ents:
        if e["類別"] == KIND_LOT and code_key.get(e["歸戶"]) != key_of.get(e["原地號"]):
            raise PM.SourceError(f"{e['暫編']} 之歸戶 {e['歸戶']} 與其原地號之戶不一")
    return lots, owners, area, key_of, hh_no, hh_lots, code_key


def cash_of(doc, code_key, ents):
    """現金補償 ⇒ {戶鍵: 金額}；代號不明、金額非正整數、重複、或該戶已有分配地 ⇒ 停。"""
    have = {code_key[e["歸戶"]] for e in ents if e["類別"] == KIND_LOT}
    out = {}
    for c in doc["現金補償"]:
        k = code_key.get(c.get("歸戶"))
        amt = c.get("金額")
        if k is None:
            raise PM.SourceError(f"現金補償之歸戶 {c.get('歸戶')!r} 不明")
        if not isinstance(amt, int) or isinstance(amt, bool) or amt <= 0:
            raise PM.SourceError(f"現金補償之金額 {amt!r} 非正整數")
        if k in out:
            raise PM.SourceError(f"現金補償之歸戶 {c['歸戶']} 重複")
        if k in have:
            raise PM.SourceError(f"現金補償之歸戶 {c['歸戶']} 已有分配地")
        out[k] = amt
    return out


def shares(owners, ln):
    """⇒ [(姓名, 統編, 持分, 住址, 管理人姓名, 管理人統編)]（登記次序）。"""
    out = []
    for o in PM.owner_rows(owners, ln):
        den, num = int(o[PM.C_DEN]), int(o[PM.C_NUM])
        if den <= 0 or num <= 0:
            raise PM.SourceError(f"{ln} 之權利範圍不合：{num}/{den}")
        out.append((PM._val(o[PM.C_NAME]), PM._val(o[PM.C_UID]), f"{num}/{den}", PM._val(o[PM.C_ADDR]),
                    PM._val(o[PM.C_MGR_NAME]), PM._val(o[PM.C_MGR_UID])))
    return out


def household_shares(owners, lns):
    """一戶之各宗之（姓名, 統編, 持分）須同（重劃後按原共有狀態·一戶之分配地之持分即此）；不同 ⇒ 停。"""
    s = {ln: tuple(r[:3] for r in shares(owners, ln)) for ln in lns}
    if len(set(s.values())) != 1:
        raise PM.SourceError(f"戶 {lns} 之各宗之持分不同（須地價換算之持分·本單未定）")
    return next(iter(s.values()))


def curval(owners, ln):
    vals = {o[PM.C_CURVAL] for seq in owners[ln].values() for o in seq}
    if len(vals) != 1 or any(PM._blank(v) for v in vals):
        raise PM.SourceError(f"{ln} 之公告現值缺或不一：{vals}")
    return Decimal(str(next(iter(vals))))


def avg_curval(e, owners, hh_lns):
    """平均公告現值（KL 23:31 問三之讀：Σ(計入 G 之面積 × 現值) ÷ Σ 計入 G 之面積·個位）。
    無併入 ⇒ 原地號之現值；有併入（其逐筆來源未出）⇒ 該戶各宗之現值皆同者取之（加權平均必等於之），否則 None（留空）。"""
    if D(e["併入面積"]) == 0:
        return q(curval(owners, e["原地號"]), 0)
    vals = {curval(owners, ln) for ln in hh_lns}
    return q(next(iter(vals)), 0) if len(vals) == 1 else None


def full_households(ents, key_of, hh_lots, area):
    """⇒ {戶鍵}：該戶之分配地之 a 之和 ＝ 該戶重劃前各宗登記面積之和（差 ≤ 0.01×宗數·a 取二位所致）
    ⇒ 各宗皆整筆參加分配（取配面積 ＝ 登記面積）。"""
    tot = {}
    for e in ents:
        if e["類別"] == KIND_LOT:
            k = key_of[e["原地號"]]
            n, s = tot.get(k, (0, Decimal(0)))
            tot[k] = (n + 1, s + q(D(e["a"]), 2))
    return {k for k, (n, s) in tot.items()
            if abs(s - sum(D(area[ln]) for ln in hh_lots[k])) <= Decimal("0.01") * n}


# ---------------------------------------------------------------- 計算（附表 24）
def calc(e, fin, owners, hh_lns, single):
    post = D(fin["重劃後單價"][e["街廓"]])
    pre = D(fin["重劃前單價"][e["重劃前區段"]])
    B, C = D(fin["B"]), D(fin["C"])
    G, area = q(D(e["G"]), 2), q(D(e["面積"]), 2)
    tot = q(D(e["a原位次"]), 2) + q(D(e["併入面積"]), 2)
    a0, a1 = (tot, Decimal("0.00")) if single else (None, None)   # single：其戶唯一筆且整筆參加分配
    A = post / pre
    rw = q(D(e["Rw"]) / 100, 4)
    w = q(D(e["W"]), 2)
    b1 = q(G * C / (1 - C), 2)
    b2 = tot - G - b1
    c3 = avg_curval(e, owners, hh_lns)
    y1 = q(b1 * post, 0)
    y2 = None if c3 is None else q(b2 * c3, 0)
    diff = q((area - G) * post, 0)
    f8 = lambda x: f"{q(x, 8):.8f}"  # noqa: E731
    formula = (f"{G:.2f}=〔{tot:.2f}×(1−{f8(A)}×{f8(B)})−{q(D(e['F']), 2):.2f}×{q(D(e['l側']), 2):.2f}×{rw:.4f}"
               f"−{q(D(e['S']), 2):.2f}×{q(D(e['l正']), 2):.2f}〕×(1−{f8(C)})")
    return {"post": post, "pre": pre, "B": B, "C": C, "A": A, "G": G, "area": area,
            "a0": a0, "a1": a1, "tot": tot, "rw": rw, "w": w, "b1": b1, "b2": b2, "c3": c3, "y1": y1, "y2": y2,
            "sum": None if y2 is None else y1 + y2, "diff": diff, "formula": formula}


# ---------------------------------------------------------------- 寫表之共用
def _put(ws, r, c, v, nf=None, r2=None, c2=None):
    from openpyxl.styles import Alignment
    cell = ws.cell(r, c, v)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if nf:
        cell.number_format = NF[nf]
    if (r2 or r, c2 or c) != (r, c):
        ws.merge_cells(start_row=r, start_column=c, end_row=r2 or r, end_column=c2 or c)
    return cell


def _num(x):
    return None if x is None else float(x)


# ---------------------------------------------------------------- 附表 13
def a13_sheet(wb, name, head_owner, pre_rows, post_rows, note_first):
    ws = wb.create_sheet(name)
    _put(ws, 1, 1, TITLE_A13, c2=A13_NCOL)
    _put(ws, 2, 19, BOOK_A13, c2=A13_NCOL)
    vals = list(head_owner)
    for (r, (c1, c2, lab), (v1, v2)), v in zip(A13_OWNER, vals):
        _put(ws, r, c1, lab, c2=c2)
        _put(ws, r, v1, v, c2=v2)
    for r1, c1, r2, c2, lab in A13_HEAD:
        _put(ws, r1, c1, lab, r2=r2, c2=c2)
    n = max(len(pre_rows), len(post_rows), 1)
    for i in range(n):
        r = A13_FIRST + i
        if i < len(pre_rows):
            for c, (v, nf) in enumerate(pre_rows[i], start=1):
                _put(ws, r, c, v, nf)
        if i < len(post_rows):
            for c, (v, nf) in enumerate(post_rows[i], start=11):
                _put(ws, r, c, v, nf)
    if note_first:
        _put(ws, A13_FIRST, 24, note_first)
    last = A13_FIRST + n - 1
    PM._grid(ws, 3, last, A13_NCOL)
    ws.cell(last + 2, 1, NOTE_A13)
    step = A13_NCOL // len(SIGN_A13)
    for i, lab in enumerate(SIGN_A13):
        ws.cell(last + 4, 1 + i * step, lab)
    PM._page(ws, "1:7", landscape=True)
    return ws


def lot_post_rows(e, no, zoning, own3, fin):
    post = D(fin["重劃後單價"][e["街廓"]])
    G, area = q(D(e["G"]), 2), q(D(e["面積"]), 2)
    diff = q((area - G) * post, 0)
    pay, get = (diff, None) if diff >= 0 else (None, -diff)
    rows = []
    for nm, uid, sh in own3:
        rows.append([(float(G), "area"), (zoning, None), (None, None), (None, None), (no, None),
                     (float(area), "area"), (float(post), "p2"), (_num(pay), "yuan"), (_num(get), "yuan"),
                     (nm, None), (uid, None), (sh, None), (None, None)])
    return rows


def write_a13(path, doc, ents, numbers, sch, cash, zoning, pre_price, full):
    import openpyxl
    lots, owners, area, key_of, hh_no, hh_lots, code_key = sch
    fin = doc["財務"]
    loc = {x["地號"]: x for x in lots}
    by_hh = {}
    for e in ents:
        if e["類別"] == KIND_LOT:
            by_hh.setdefault(key_of[e["原地號"]], []).append(e)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    out = []
    for k in sorted(hh_lots, key=lambda kk: hh_no[kk]):
        lns = sorted(hh_lots[k], key=PM.landno_key)
        first = shares(owners, lns[0])[0]
        head = (first[0], first[1], first[3], first[4], first[5])
        pre_rows = []
        for ln in lns:
            for nm, uid, sh, *_ in shares(owners, ln):
                pre_rows.append([(loc[ln]["鄉鎮市區"], None), (loc[ln]["段"], None), (loc[ln].get("小段"), None),
                                 (ln, None), (area[ln], "area"), (area[ln] if k in full else None, "area"),
                                 (nm, None), (uid, None),
                                 (sh, None), (pre_price[ln], "yuan")])
        post_rows = []
        es = sorted(by_hh.get(k, []), key=lambda e: PM.landno_key(numbers[e["暫編"]]))
        if es:
            own3 = household_shares(owners, lns)
            for e in es:
                post_rows += lot_post_rows(e, numbers[e["暫編"]], zoning, own3, fin)
        note = NOTE_CASH.format(cash[k]) if k in cash else None
        a13_sheet(wb, f"歸戶號{hh_no[k]}", head, pre_rows, post_rows, note)
        out.append((hh_no[k], lns, [numbers[e["暫編"]] for e in es], note))
    groups = {}   # 依管理機關分表：抵費地一表、公共設施依其分類各一表（管理機關資料所無 ⇒ 分類即其別）
    for e in sorted((e for e in ents if e["類別"] != KIND_LOT), key=lambda e: PM.landno_key(numbers[e["暫編"]])):
        groups.setdefault(e["類別"], []).append(e)
    for cat in sorted(groups, key=lambda c: (c != KIND_OFFSET, PM.landno_key(numbers[groups[c][0]["暫編"]]))):
        off = cat == KIND_OFFSET
        rows = []
        for e in groups[cat]:
            price = (float(D(fin["重劃後單價"][e["街廓"]])), "p2") if off else (None, None)
            rows.append([(None, "area"), (zoning if off else e["類別"], None), (None, None), (None, None),
                         (numbers[e["暫編"]], None), (float(q(D(e["面積"]), 2)), "area"), price,
                         (None, "yuan"), (None, "yuan"), (COUNTY, None), (None, None), ("1/1", None), (None, None)])
        ws = a13_sheet(wb, SHEET_COUNTY.format(cat), (COUNTY, None, None, MGR_OFFSET if off else None, None),
                       [], rows, None)
        for i in range(len(rows)):
            _put(ws, A13_FIRST + i, 24, NOTE_OFFSET if off else NOTE_PUBLIC)
    wb.save(path)
    return out


# ---------------------------------------------------------------- 附表 24
def write_a24(path, doc, ents, numbers, sch, zoning, full):
    import openpyxl
    lots, owners, area, key_of, hh_no, hh_lots, code_key = sch
    fin = doc["財務"]
    loc = {x["地號"]: x for x in lots}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    out = []
    es = sorted((e for e in ents if e["類別"] == KIND_LOT), key=lambda e: PM.landno_key(numbers[e["暫編"]]))
    for e in es:
        no, ln = numbers[e["暫編"]], e["原地號"]
        k = key_of[ln]
        lns = sorted(hh_lots[k], key=PM.landno_key)
        single = k in full and len(lns) == 1
        x = calc(e, fin, owners, lns, single)
        own = shares(owners, ln)
        ws = wb.create_sheet(f"地號{no}")
        _put(ws, 1, 1, TITLE_A24, c2=A24_NCOL)
        _put(ws, 2, 1, "歸戶號")
        _put(ws, 2, 2, hh_no[k])
        _put(ws, 2, 3, "所有權(代表)人姓名")
        _put(ws, 2, 4, own[0][0] if len(own) == 1 else f"{own[0][0]}等{len(own)}人", c2=A24_NCOL)
        r = 3
        _put(ws, r, 1, "重劃前", c2=A24_NCOL)
        pre = (loc[ln]["段"], loc[ln].get("小段"), ln, (area[ln], "area"), (_num(x["a0"]), "area"), None, None, None,
               (float(curval(owners, ln)), "yuan"), (float(x["pre"]), "p2"))
        for lab, v in zip(A24_PRE, pre):
            r += 1
            _put(ws, r, 1, lab)
            v, nf = v if isinstance(v, tuple) else (v, None)
            _put(ws, r, 2, v, nf)
        r += 1
        _put(ws, r, 1, "計算分配有關資料", c2=A24_NCOL)
        calc_v = ((x["a0"], "area"), (x["a1"], "area"), (x["tot"], "area"), (x["pre"], "p2"), (None, None),
                  (x["post"], "p2"), (x["A"], "k8"), (x["B"], "k8"), (x["C"], "k8"), (1 - x["C"], "k8"),
                  (q(D(e["S"]), 2), "m2"), (x["w"], "m2"), (q(D(e["F"]), 2), "m2"), (q(D(e["l側"]), 2), "m2"),
                  (q(D(e["l正"]), 2), "m2"), (x["rw"], "k4"))
        for lab, (v, nf) in zip(A24_CALC, calc_v):
            r += 1
            _put(ws, r, 1, lab)
            _put(ws, r, 2, _num(v), nf)
        r += 1
        _put(ws, r, 1, "計算式")
        _put(ws, r, 2, x["formula"], c2=A24_NCOL)
        r += 1
        _put(ws, r, 1, "重劃後", c2=A24_NCOL)
        d = x["diff"]
        box = ("■繳納□領取" if d > 0 else "□繳納■領取" if d < 0 else "□繳納□領取") + f"新臺幣 {abs(d):,} 元"
        post_v = ((x["post"], "p2"), (x["sum"], "yuan"), (x["G"], "area"), (x["area"], "area"), (x["G"] / x["tot"], "pct"),
                  (None, None), (None, None), (no, None), (x["area"], "area"), (zoning, None), (box, None))
        for lab, (v, nf) in zip(A24_POST, post_v):
            r += 1
            _put(ws, r, 1, lab)
            _put(ws, r, 2, v if isinstance(v, str) or v is None else float(v), nf)
        for head, mark, (a_, p_, y_, pnf) in zip(A24_BURDEN_HEAD, ("①", "②"),
                                                 ((x["b1"], x["post"], x["y1"], "p2"), (x["b2"], x["c3"], x["y2"], "yuan"))):
            r += 1
            for c, lab in enumerate(head, start=1):
                _put(ws, r, c, lab)
            r += 1
            _put(ws, r, 1, mark)
            _put(ws, r, 2, _num(a_), "area")
            _put(ws, r, 3, _num(p_), pnf)
            _put(ws, r, 4, _num(y_), "yuan")
        burden = ((None, None, x["sum"]), (None, None, d if d > 0 else 0), (None, None, -d if d < 0 else 0))
        for lab, (a_, p_, y_) in zip(A24_BURDEN, burden):
            r += 1
            _put(ws, r, 1, lab)
            _put(ws, r, 2, _num(a_), "area")
            _put(ws, r, 3, _num(p_), "p2")
            _put(ws, r, 4, _num(y_), "yuan")
        PM._grid(ws, 2, r, A24_NCOL)
        notes = [LEGEND_A24, NOTE_A24, NOTE_ROUND, NOTE_PRE] + ([] if single else [NOTE_SPLIT])
        for i, s in enumerate(notes, start=1):
            ws.cell(r + i, 1, s)
        for i, lab in enumerate(SIGN_A24):
            ws.cell(r + len(notes) + 2, 1 + i, lab)
        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 16
        PM._page(ws, "1:2", landscape=False)
        out.append((no, hh_no[k], x))
    wb.save(path)
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description="W-P.006 附表 13、24")
    ap.add_argument("--state", required=True)
    ap.add_argument("--moedict", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--map")
    a = ap.parse_args(argv)
    doc = PMAP.load_state(a.state)
    check_state(doc)
    tag = PMAP.tag_of(doc)
    ents = PMAP.entries_of(doc)
    order = PMAP.s_order(doc, ents)
    numbers = PMAP.read_map(a.map, ents) if a.map else PMAP.default_numbers(order)
    sch = scheme(a.moedict, doc, ents)
    cash = cash_of(doc, sch[6], ents)
    zoning = load_case_zoning()
    zone_of, _unit, price = PM.load_price_zones()
    pre_price = {}
    for x in sch[0]:
        ln = x["地號"]
        if ln not in zone_of or zone_of[ln] not in price:
            raise PM.SourceError(f"{ln} 無重劃前地價區段或單價")
        pre_price[ln] = price[zone_of[ln]]
    os.makedirs(a.out, exist_ok=True)
    p = {k: os.path.join(a.out, v.format(tag)) for k, v in OUT_FMT.items()}
    full = full_households(ents, sch[3], sch[5], sch[2])
    r13 = write_a13(p["a13"], doc, ents, numbers, sch, cash, zoning, pre_price, full)
    r24 = write_a24(p["a24"], doc, ents, numbers, sch, zoning, full)
    n_lot = sum(1 for e in ents if e["類別"] == KIND_LOT)
    print(f"退縮 {tag} m：附表 13 戶 {len(r13)}（有分配地 {sum(1 for x in r13 if x[2])}、取配面積可填 {len(full)}、"
          f"現金補償 {len(cash)}）＋縣有 {len(ents) - n_lot}；附表 24 張 {len(r24)}"
          f"（Σa 可填 {sum(1 for x in r24 if x[2]['a0'] is not None)}）")
    for no, hh, x in r24:
        print(f"  地號 {no}（歸戶號 {hh}）：G {x['G']}、Σ {x['tot']}（Σa {x['a0']}、Σa' {x['a1']}）、① {x['b1']}、② {x['b2']}、"
              f"③ {x['c3']}、負擔額 {x['y1']}＋{x['y2']}＝{x['sum']}、差額 {x['diff']}")
    for k, v in p.items():
        print(f"輸出 {v}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
