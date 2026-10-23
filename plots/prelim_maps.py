# -*- coding: utf-8 -*-
"""
W-P.001　重劃前三圖層（地籍線·歸戶圖·地價圖）＋顏色對照表＋附表 2、3。

命令列：python plots/prelim_maps.py --out <目錄>
輸出（皆寫至 --out，⛔ 入倉）：
  重劃前圖.dxf（DXF R2010·三圖層）
  顏色對照表_重劃前.xlsx
  附表02_重劃前土地地號摘錄簿.xlsx
  附表03_重劃前土地清冊.xlsx

所據：出圖線/orders/W-P.001_重劃前三圖層與附表2-3.md（§一、§二）。
宗地幾何唯以 app.py 之三函式（經 verify/app_harvest.harvest 取之·⛔ 改 app.py）：
  parse_ua_cadastral_dxf → reconstruct_parcel_polygons → validate_parcel_assignments_by_area。
缺值一律報錯停止（⛔ 靜默補值）；匯入資料所無之欄留空（⛔ 杜撰）。
"""
import argparse
import colorsys
import json
import os
import sys
from decimal import Decimal, ROUND_HALF_UP

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO, "verify"))

DXF_SRC = os.path.join(REPO, "data", "V6_1.dxf")
XLSX_SRC = os.path.join(REPO, "data", "地籍資料來源_匿名版.xlsx")
CASE_PARAMS = os.path.join(REPO, "verify", "case_params_UC9898.json")

LAYER_CAD = "重劃前地籍線"
LAYER_OWN = "重劃前歸戶圖"
LAYER_PRICE = "重劃前地價圖"
LAYERS = (LAYER_CAD, LAYER_OWN, LAYER_PRICE)
TEXT_STYLE = "CJK"
TEXT_FONT = "msjh.ttc"
TEXT_HEIGHT = 1.0

SECTION_NAME = "0013德安段"   # U_LAND 之段小段（§一 4）
EXPECTED_SWAPS = {frozenset({"628-3", "628-1"}), frozenset({"628-44", "628-41"})}

# U_LAND 之欄位置（0 起算·表頭有重名故以位置取·§一 4）
C_SECTION, C_CODE, C_AREA, C_CURVAL = 0, 1, 6, 9
C_SEQ, C_DEN, C_NUM, C_DECL = 23, 28, 29, 31
C_UID, C_NAME, C_ADDR, C_MGR_UID, C_MGR_NAME = 32, 33, 34, 39, 40
C_RIGHT_KIND = 60


class SourceError(RuntimeError):
    """資料與單之所據不符 ⇒ 停機上呈。"""


def _blank(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", "null"))


def _val(v):
    return None if _blank(v) else v


def code_to_landno(code):
    m, s = divmod(int(code), 10000)
    return f"{m}" if s == 0 else f"{m}-{s}"


def landno_key(s):
    parts = str(s).split("-")
    return (int(parts[0]), int(parts[1]) if len(parts) > 1 else 0)


def round_half_up(x, nd=0):
    q = Decimal(1).scaleb(-nd)
    r = Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP)
    return int(r) if nd == 0 else float(r)


# ---------------------------------------------------------------- 讀取
def load_cadastre(xlsx=XLSX_SRC):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    rows = list(wb["重劃區地籍"].iter_rows(values_only=True))
    head = rows[0]
    i_town, i_sec, i_lot = head.index("鄉鎮市"), head.index("地段"), head.index("地號")
    lots = []
    for r in rows[1:]:
        if _blank(r[i_lot]):
            continue
        lots.append({"鄉鎮市區": _val(r[i_town]), "段": _val(r[i_sec]), "地號": str(r[i_lot]).strip()})
    lotset = {x["地號"] for x in lots}
    if len(lotset) != len(lots):
        raise SourceError("「重劃區地籍」之地號有重複")

    owners = {}   # 地號 -> {登記次序: row}
    area = {}
    for r in wb["U_LAND"].iter_rows(values_only=True):
        if r[C_SECTION] != SECTION_NAME or _blank(r[C_CODE]):
            continue
        try:
            ln = code_to_landno(r[C_CODE])
        except (TypeError, ValueError):
            continue
        if ln not in lotset:
            continue
        a = float(r[C_AREA])
        if ln in area and abs(area[ln] - a) > 1e-9:
            raise SourceError(f"{ln} 之 U_LAND 面積不一致：{area[ln]} vs {a}")
        area[ln] = a
        owners.setdefault(ln, {})
        seq = r[C_SEQ]
        if _blank(seq):
            raise SourceError(f"{ln} 有所有權登記次序為空之列")
        owners[ln].setdefault(seq, []).append(r)
    missing = lotset - set(owners)
    if missing:
        raise SourceError(f"U_LAND 無下列地號：{sorted(missing, key=landno_key)}")
    return lots, owners, area


def load_price_zones(path=CASE_PARAMS):
    v = json.load(open(path, encoding="utf-8"))["財務接線_v3"]
    zone_of = dict(v["原地號_區段"])
    unit = {z: d["單價_元每m2"] for z, d in v["重劃前區段_面積單價"].items()}
    price = {z: round_half_up(u, 0) for z, u in unit.items()}
    return zone_of, unit, price


def build_parcels(area):
    from app_harvest import harvest
    ns, _ = harvest()
    with open(DXF_SRC, "rb") as f:
        b = f.read()
    parsed = ns["parse_ua_cadastral_dxf"](b)
    polys = ns["reconstruct_parcel_polygons"](parsed["boundary_lines"], parsed["parcels"])
    res = ns["validate_parcel_assignments_by_area"](polys, area)
    swaps = {frozenset({s["parcel_a_orig"], s["parcel_b_orig"]}) for s in res["swaps"]}
    if len(res["swaps"]) != 2 or swaps != EXPECTED_SWAPS:
        raise SourceError(f"面積校正之交換不符單之 §二 2：{res['swaps']}")
    return res["parcel_polys"], res["swaps"]


# ---------------------------------------------------------------- 歸戶與顏色
def owner_rows(owners, ln):
    """每宗每所有權人一列（以所有權登記次序去重）。"""
    out = []
    for seq in sorted(owners[ln], key=lambda s: int(s)):
        out.append(owners[ln][seq][0])
    return out


def households(lots, owners):
    """歸戶鍵 ＝ 該宗全體所有權人之（統編, 姓名）之集合（手冊歸戶方法）。"""
    key_of = {}
    for x in lots:
        ln = x["地號"]
        key_of[ln] = frozenset((str(r[C_UID]), str(r[C_NAME])) for r in owner_rows(owners, ln))
    order = []
    for ln in sorted(key_of, key=landno_key):
        if key_of[ln] not in order:
            order.append(key_of[ln])
    return key_of, order


def distinct_palette(n):
    cols = []
    for i in range(n):
        h = (i * 0.618033988749895) % 1.0
        s = (0.55, 0.75, 0.40)[i % 3]
        v = (0.90, 0.72, 0.80)[(i // 3) % 3]
        r, g, b = colorsys.hsv_to_rgb(h, s, v)
        cols.append((int(round(r * 255)), int(round(g * 255)), int(round(b * 255))))
    if len(set(cols)) != n:
        raise RuntimeError("色盤產生重色")
    return cols


ZONE_COLORS = [(106, 70, 190), (46, 204, 64), (240, 160, 40), (60, 140, 220)]


# ---------------------------------------------------------------- 輸出
def write_dxf(path, parcels, key_of, hh_color, zone_of, zone_color, price):
    import ezdxf
    from ezdxf.enums import TextEntityAlignment
    from shapely.geometry import Polygon

    doc = ezdxf.new("R2010")
    # 系統圖層「0」「Defpoints」由 ezdxf 存檔時必建（後者刪亦補回）；其上⛔ 置圖元。
    for name in LAYERS:
        doc.layers.add(name)
    doc.styles.new(TEXT_STYLE, dxfattribs={"font": TEXT_FONT})
    msp = doc.modelspace()

    for p in sorted(parcels, key=lambda q: landno_key(q["parcel"])):
        ln = p["parcel"]
        coords = [(float(x), float(y)) for x, y in p["polygon_coords"]]
        if coords[0] == coords[-1]:
            coords = coords[:-1]
        rp = Polygon(coords).representative_point()
        msp.add_lwpolyline(coords, close=True, dxfattribs={"layer": LAYER_CAD})
        t = msp.add_text(ln, height=TEXT_HEIGHT, dxfattribs={"layer": LAYER_CAD, "style": TEXT_STYLE})
        t.set_placement((rp.x, rp.y), align=TextEntityAlignment.MIDDLE_CENTER)

        h = msp.add_hatch(dxfattribs={"layer": LAYER_OWN})
        h.set_solid_fill(rgb=hh_color[key_of[ln]])
        h.paths.add_polyline_path(coords, is_closed=True)

        z = zone_of[ln]
        h2 = msp.add_hatch(dxfattribs={"layer": LAYER_PRICE})
        h2.set_solid_fill(rgb=zone_color[z])
        h2.paths.add_polyline_path(coords, is_closed=True)
        t2 = msp.add_text(str(price[z]), height=TEXT_HEIGHT,
                          dxfattribs={"layer": LAYER_PRICE, "style": TEXT_STYLE})
        t2.set_placement((rp.x, rp.y), align=TextEntityAlignment.MIDDLE_CENTER)
    doc.saveas(path)


def _xlsx(path, sheets):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(name)
        ws.append(list(header))
        for r in rows:
            ws.append(list(r))
    wb.save(path)


def owner_label(r):
    return f"{r[C_NAME]}（{r[C_UID]}）"


def write_color_table(path, lots, owners, key_of, hh_order, hh_color, zone_of, zone_color, price):
    rows_own = []
    for k in hh_order:
        lns = sorted([ln for ln, kk in key_of.items() if kk == k], key=landno_key)
        first = owner_rows(owners, lns[0])
        names = "、".join(owner_label(r) for r in sorted(first, key=lambda r: str(r[C_UID])))
        rows_own.append((LAYER_OWN, "{},{},{}".format(*hh_color[k]), names, "、".join(lns)))
    rows_price = []
    for z in sorted(zone_color):
        lns = sorted([ln for ln, zz in zone_of.items() if zz == z and ln in key_of], key=landno_key)
        rows_price.append((LAYER_PRICE, "{},{},{}".format(*zone_color[z]), z, price[z], "、".join(lns)))
    _xlsx(path, [
        ("重劃前歸戶圖", ("圖層", "顏色 R,G,B", "所有權人", "地號"), rows_own),
        ("重劃前地價圖", ("圖層", "顏色 R,G,B", "區段", "宗地地價（元/㎡）", "地號"), rows_price),
    ])


APPX2_HEAD = ("鄉鎮市區", "段", "小段", "地號", "備註")
APPX3_HEAD = ("鄉鎮市區", "段", "小段", "地號", "面積（㎡）", "持分", "持分面積（㎡）", "姓名",
              "國民身分證統一編號", "住所", "當期公告現值（元/㎡）", "當期申報地價（元/㎡）",
              "原規定地價或前次移轉現值（元/㎡）", "管理人（代表人）", "他項權利", "建物",
              "限制登記", "備註")


def write_appx2(path, lots):
    rows = [(x["鄉鎮市區"], x["段"], None, x["地號"], None)
            for x in sorted(lots, key=lambda x: landno_key(x["地號"]))]
    _xlsx(path, [("附表2", APPX2_HEAD, rows)])
    return rows


def write_appx3(path, lots, owners, area):
    rows = []
    for x in sorted(lots, key=lambda x: landno_key(x["地號"])):
        ln = x["地號"]
        all_rows = [r for seq in owners[ln].values() for r in seq]
        has_right = any(not _blank(r[C_RIGHT_KIND]) for r in all_rows)
        for r in owner_rows(owners, ln):
            den, num = int(r[C_DEN]), int(r[C_NUM])
            if den <= 0:
                raise SourceError(f"{ln} 之權利範圍分母非正：{den}")
            rows.append((
                x["鄉鎮市區"], x["段"], None, ln, area[ln], f"{num}/{den}",
                round_half_up(area[ln] * num / den, 2),
                _val(r[C_NAME]), _val(r[C_UID]), _val(r[C_ADDR]),
                _val(r[C_CURVAL]), _val(r[C_DECL]), None,
                "ˇ" if not _blank(r[C_MGR_UID]) else None,
                "ˇ" if has_right else None,
                None, None, None,
            ))
    _xlsx(path, [("附表3", APPX3_HEAD, rows)])
    return rows


OUT_NAMES = {
    "dxf": "重劃前圖.dxf",
    "color": "顏色對照表_重劃前.xlsx",
    "appx2": "附表02_重劃前土地地號摘錄簿.xlsx",
    "appx3": "附表03_重劃前土地清冊.xlsx",
}


def main(argv=None):
    ap = argparse.ArgumentParser(description="W-P.001 重劃前三圖層＋顏色對照表＋附表 2、3")
    ap.add_argument("--out", required=True)
    a = ap.parse_args(argv)
    os.makedirs(a.out, exist_ok=True)

    lots, owners, area = load_cadastre()
    zone_of, unit, price = load_price_zones()
    miss = {x["地號"] for x in lots} - set(zone_of)
    if miss:
        raise SourceError(f"case_params 之原地號_區段 缺：{sorted(miss)}")
    parcels, swaps = build_parcels(area)
    pset = {p["parcel"] for p in parcels}
    if pset != {x["地號"] for x in lots} or len(parcels) != len(lots):
        raise SourceError(f"幾何之宗與「重劃區地籍」不符：{sorted(pset ^ {x['地號'] for x in lots})}")

    key_of, hh_order = households(lots, owners)
    hh_color = dict(zip(hh_order, distinct_palette(len(hh_order))))
    zones = sorted(set(zone_of[x["地號"]] for x in lots))
    if len(zones) > len(ZONE_COLORS):
        raise SourceError("區段數逾色盤")
    zone_color = dict(zip(zones, ZONE_COLORS))

    p = {k: os.path.join(a.out, v) for k, v in OUT_NAMES.items()}
    write_dxf(p["dxf"], parcels, key_of, hh_color, zone_of, zone_color, price)
    write_color_table(p["color"], lots, owners, key_of, hh_order, hh_color, zone_of, zone_color, price)
    r2 = write_appx2(p["appx2"], lots)
    r3 = write_appx3(p["appx3"], lots, owners, area)
    print(f"宗數 {len(parcels)}；交換 {[(s['parcel_a_orig'], s['parcel_b_orig']) for s in swaps]}")
    print(f"戶數 {len(hh_order)}；區段 {zones}；宗地地價 {price}（來源 verify/case_params_UC9898.json 財務接線_v3）")
    print(f"附表2 {len(r2)} 列；附表3 {len(r3)} 列")
    for k, v in p.items():
        print(f"輸出 {v}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
