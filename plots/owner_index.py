# -*- coding: utf-8 -*-
"""
W-P.002　歸戶號之編定與附表 26 土地所有權人索引表。

所據：出圖線/orders/W-P.002_歸戶號與附表26.md（§一 2、§二 1、§二 3）。
  - 筆劃 ＝ 教育部標準字體（KL 2026-10-24 11:43「甲」）：教育部《重編國語辭典修訂本》之單字之
    stroke_count 與 radical（plots/fetch_moedict.py 取之·⛔ 入倉）。
  - 排序（KL 2026-10-23 21:18「單獨在前」；2026-10-25 02:30「一甲 二乙 三甲 四可·甲」；02:33）：
    逐字比較，每字先比筆劃、次部首（康熙 214 部之序）、再 ord；非漢字排於漢字之後。
    單獨所有之戶依（群, 姓名, 最小地號）；共有戶依（人數, 首列登記人〔該戶自然人中姓名最前者〕, 最小地號）；
    單獨在前、共有在後，自 1 連號。群：自然人 0、公司與團體 1、公有 2。
  - 索引表：每（所有權人, 戶）一列，依（群, 姓名, 歸戶號）排序；「冊」「頁」留空，「附註」＝ 歸戶號。
缺值一律報錯停止（⛔ 補值、⛔ 自定）。
"""
import hashlib
import json
import unicodedata

MOE_JSON_SHA256 = "f68fc3df140a64d8de6eb905ff6576bfd9d260eed0bbf5cc9c41aba5a131a7ad"

# U_LAND 之欄位置（0 起算·§一 4）
C_KIND, C_PUB = 37, 38

GROUP_PERSON, GROUP_CORP, GROUP_PUBLIC = 0, 1, 2


class SourceError(RuntimeError):
    """資料與單之所據不符 ⇒ 停機上呈。"""


# ---------------------------------------------------------------- 筆劃之資料
def _radical_numbers():
    no = {}
    for i in range(214):
        no[unicodedata.normalize("NFKC", chr(0x2F00 + i))] = i + 1
    no["青"] = 174     # 資料之部首「青」以「靑」（174）對之（§二 1）
    return no


def load_strokes(path):
    """字 →（筆劃, 部首號）。JSON 之 sha256 不合、部首不在 214 部、一字二值 ⇒ SourceError。"""
    with open(path, "rb") as f:
        raw = f.read()
    sha = hashlib.sha256(raw).hexdigest()
    if sha != MOE_JSON_SHA256:
        raise SourceError(f"筆劃之資料之 sha256 不合：{sha}")
    rad_no = _radical_numbers()
    tab = {}
    for e in json.loads(raw):
        t = e.get("title", "")
        if len(t) != 1:
            continue
        rad = str(e.get("radical", "")).strip()
        if rad not in rad_no:
            raise SourceError(f"「{t}」之部首「{rad}」不在康熙 214 部")
        try:
            v = (int(e["stroke_count"]), rad_no[rad])
        except (KeyError, TypeError, ValueError) as ex:
            raise SourceError(f"「{t}」之筆劃缺或非整數：{e.get('stroke_count')!r}") from ex
        if t in tab and tab[t] != v:
            raise SourceError(f"「{t}」一字二值：{tab[t]} vs {v}")
        tab[t] = v
    return tab


def _is_han(c):
    return unicodedata.name(c, "").startswith(("CJK UNIFIED IDEOGRAPH", "CJK COMPATIBILITY IDEOGRAPH"))


def char_key(c, strokes):
    if _is_han(c):
        if c not in strokes:
            raise SourceError(f"漢字「{c}」（U+{ord(c):04X}）無筆劃之資料（⛔ 補值）")
        s, r = strokes[c]
        return (0, s, r, ord(c))
    return (1, ord(c), 0, 0)


def name_key(name, strokes):
    return tuple(char_key(c, strokes) for c in str(name))


# ---------------------------------------------------------------- 分群
def group_of(row):
    """自然人 0、公司與團體 1、公有 2（§一 4）；其餘之組合 ⇒ SourceError。"""
    kind, pub = str(row[C_KIND]).strip(), str(row[C_PUB]).strip()
    if pub == "公":
        return GROUP_PUBLIC
    if pub == "私" and kind == "1":
        return GROUP_PERSON
    if pub == "私" and kind == "7":
        return GROUP_CORP
    raise SourceError(f"「類別／公私別」＝（{kind}, {pub}）不在單之 §一 4 之三種")


# ---------------------------------------------------------------- 歸戶號
def number_households(hh_lots, members, strokes, landno_key):
    """
    hh_lots：戶鍵 → [地號]；members：戶鍵 → [(統編, 姓名, 群)]。
    回傳 戶鍵 → 歸戶號（1 起連號；單獨在前、共有在後）。
    """
    singles, joints = [], []
    for k, mem in members.items():
        if not mem:
            raise SourceError(f"戶 {sorted(hh_lots.get(k, []))} 無所有權人")
        first = min(landno_key(x) for x in hh_lots[k])
        if len(mem) == 1:
            _, nm, g = mem[0]
            singles.append(((g, name_key(nm, strokes), first), k))
        else:
            heads = [name_key(nm, strokes) for _, nm, g in mem if g == GROUP_PERSON]
            if not heads:
                raise SourceError(f"共有戶無自然人（其群與人數孰先未定·停而上呈）：地號 {sorted(hh_lots[k], key=landno_key)}")
            joints.append(((len(mem), min(heads), first), k))
    seq = [k for _, k in sorted(singles, key=lambda t: t[0])] + [k for _, k in sorted(joints, key=lambda t: t[0])]
    return {k: i + 1 for i, k in enumerate(seq)}


def index_rows(members, no_of_hh, strokes):
    """每（所有權人, 戶）一列，依（群, 姓名之鍵, 歸戶號）排序 ⇒ [(姓名, 歸戶號)]。"""
    rows = []
    for k, mem in members.items():
        for _, nm, g in mem:
            rows.append(((g, name_key(nm, strokes), no_of_hh[k]), str(nm)))
    rows.sort(key=lambda t: t[0])
    return [(nm, key[2]) for key, nm in rows]


# ---------------------------------------------------------------- 附表 26（手冊表樣）
APPX26_TITLE = "土地所有權人索引表"
APPX26_NCOL = 8
APPX26_SIGN = ("承辦:", "股長:", "科長:", "製表日期: 年 月 日")   # 置 A、C、E、G（照手冊·半形冒號）


def write_appx26(path, rows):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Side

    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "土地所有權人索引表"

    ws.cell(1, 1, APPX26_TITLE).alignment = center
    ws.merge_cells("A1:H1")
    for off in (0, 4):
        c = 1 + off
        ws.cell(2, c, "所有權人姓名").alignment = center
        ws.merge_cells(start_row=2, start_column=c, end_row=3, end_column=c)
        ws.cell(2, c + 1, "分配清冊").alignment = center
        ws.merge_cells(start_row=2, start_column=c + 1, end_row=2, end_column=c + 2)
        ws.cell(3, c + 1, "第 冊").alignment = center
        ws.cell(3, c + 2, "第 頁").alignment = center
        ws.cell(2, c + 3, "附註").alignment = center
        ws.merge_cells(start_row=2, start_column=c + 3, end_row=3, end_column=c + 3)

    r = 3
    for i, (nm, no) in enumerate(rows):
        r, j = 4 + i // 2, i % 2
        ws.cell(r, 1 + 4 * j, nm)          # 冊、頁（B、C／F、G）留空
        ws.cell(r, 4 + 4 * j, int(no))

    s = Side(style="thin")
    for rr in range(2, r + 1):
        for cc in range(1, APPX26_NCOL + 1):
            ws.cell(rr, cc).border = Border(left=s, right=s, top=s, bottom=s)
    for col, w in zip("ABCDEFGH", (16, 8, 8, 10) * 2):
        ws.column_dimensions[col].width = w

    for k, lab in enumerate(APPX26_SIGN):
        ws.cell(r + 2, 1 + 2 * k, lab)

    ws.print_title_rows = "1:3"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    wb.save(path)
    return rows
