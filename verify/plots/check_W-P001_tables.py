# -*- coding: utf-8 -*-
"""
W-P.001 補令一：附表 2、3 之手冊表樣之量測（獨立於 plots/prelim_maps.py：⛔ import 之；期望取自倉內原始資料與手冊）。

命令列：python verify/plots/check_W-P001_tables.py <附表02.xlsx> <附表03.xlsx>
任一不過 ⇒ 結束碼 1。T1〜T3 附表 2；T4〜T6 附表 3。
"""
import functools
import os
import sys
from collections import defaultdict

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
XLSX_SRC = os.path.join(REPO, "data", "地籍資料來源_匿名版.xlsx")
SIGN = ["編造：", "校對：", "承辦人：", "科長：", "日期："]
H2_ROW3 = ["地號", "備註"] * 4
H3_ROW5 = ["地號", "面積（㎡）", "持分", "面積（㎡）", "姓名", "國民身分證統一編號", "住所"]
H3_ROW4 = {1: "土地標示", 3: "持分面積", 5: "所有權人", 8: "當期公告現值（元/㎡）", 9: "當期申報地價（元/㎡）",
           10: "原規定地價或前次移轉現值（元/㎡）", 11: "管理人（代表人）", 12: "他項權利", 13: "建物",
           14: "限制登記", 15: "備註"}
H3_MERGES = {"A1:O1", "A2:O2", "A4:B4", "C4:D4", "E4:G4"} | {f"{c}4:{c}5" for c in "HIJKLMNO"}


def _blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def lkey(s):
    p = str(s).split("-")
    return (int(p[0]), int(p[1]) if len(p) > 1 else 0)


@functools.lru_cache(maxsize=None)
def truth():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX_SRC, read_only=True)
    rows = list(wb["重劃區地籍"].iter_rows(values_only=True))
    h = rows[0]
    it, isec, ilot = h.index("鄉鎮市"), h.index("地段"), h.index("地號")
    body = [r for r in rows[1:] if not _blank(r[ilot])]
    place = sorted({(r[it], r[isec]) for r in body})
    lots = sorted((str(r[ilot]).strip() for r in body), key=lkey)
    ls = set(lots)
    own = defaultdict(dict)
    area = {}
    for r in wb["U_LAND"].iter_rows(values_only=True):
        if r[0] != "0013德安段" or _blank(r[1]):
            continue
        m, s = divmod(int(r[1]), 10000)
        ln = f"{m}" if s == 0 else f"{m}-{s}"
        if ln in ls:
            own[ln].setdefault(r[23], (str(r[33]), str(r[32])))
            area[ln] = float(r[6])
    return place, lots, dict(own), area


def _sign_ok(ws):
    last = max(r for r in range(1, ws.max_row + 1)
               if any(not _blank(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)))
    vals = [ws.cell(last, c).value for c in range(1, ws.max_column + 1) if not _blank(ws.cell(last, c).value)]
    return vals == SIGN, last, vals


def run(a2, a3):
    import openpyxl
    res = []

    def rec(n, ok, msg=""):
        res.append((n, bool(ok), msg))
        print(f"{'PASS' if ok else 'FAIL'} {n} {msg}")

    place, lots, own, area = truth()
    places = {p for pl in place for p in pl if p}
    if len(place) != 1:
        rec("T0", False, f"本資料應恰一組鄉鎮市區／段：{place}")
        return res
    town, sec = place[0]

    # 附表 2
    try:
        wb = openpyxl.load_workbook(a2)
        ws = wb.worksheets[0]
        if len(wb.worksheets) != 1:
            raise ValueError(f"工作表數 {len(wb.worksheets)}≠1")
        merges = {str(m) for m in ws.merged_cells.ranges}
        row2 = [ws.cell(2, c).value for c in range(1, 7)]
        row3 = [ws.cell(3, c).value for c in range(1, 9)]
        t1 = ("地號摘錄簿" in str(ws["A1"].value) and str(ws["A1"].value).startswith("○○縣(市)")
              and "A1:H1" in merges and row2 == ["鄉鎮市區", town, "段", sec, "小段", None]
              and row3 == H2_ROW3 and str(ws.print_title_rows).replace("$", "") == "1:3" and ws.max_column == 8)
        rec("T1", t1, f"題 {ws['A1'].value!r}；列2 {row2}；列3 {row3}；列印標題 {ws.print_title_rows}；欄數 {ws.max_column}")
        sok, last, sv = _sign_ok(ws)
        seq, notes, stray = [], [], []
        for r in range(4, last):
            for j in range(4):
                v, n = ws.cell(r, 1 + 2 * j).value, ws.cell(r, 2 + 2 * j).value
                if not _blank(v):
                    seq.append(str(v))
                if not _blank(n):
                    notes.append((r, n))
                if v in places or n in places:
                    stray.append(r)
        # 由左而右、再由上而下；除末列外每列滿四組
        full = all(sum(not _blank(ws.cell(r, 1 + 2 * j).value) for j in range(4)) == 4
                   for r in range(4, 4 + len(lots) // 4))
        rec("T2", seq == lots and not notes and not stray and full,
            f"地號 {len(seq)}（期望 {len(lots)}）序 {'合' if seq == lots else '不合'}；備註非空 {notes[:3]}；資料列重出鄉鎮市區／段 {stray[:3]}；每列四組 {full}")
        rec("T3", sok, f"末列 {sv}")
    except Exception as e:  # noqa: BLE001
        rec("T1", False, f"附表2 讀取失敗 {type(e).__name__}: {e}")

    # 附表 3
    try:
        wb = openpyxl.load_workbook(a3)
        ws = wb.worksheets[0]
        if len(wb.worksheets) != 1:
            raise ValueError(f"工作表數 {len(wb.worksheets)}≠1")
        merges = {str(m) for m in ws.merged_cells.ranges}
        row3 = [ws.cell(3, c).value for c in range(1, 7)]
        r4 = {c: ws.cell(4, c).value for c in H3_ROW4}
        r5 = [ws.cell(5, c).value for c in range(1, 8)]
        t4 = (str(ws["A1"].value) == "○○縣(市)第○○期○○市地重劃區" and ws["A2"].value == "重劃前土地清冊"
              and row3 == ["鄉鎮市區", town, "段", sec, "小段", None] and r4 == H3_ROW4 and r5 == H3_ROW5
              and merges == H3_MERGES and str(ws.print_title_rows).replace("$", "") == "1:5" and ws.max_column == 15)
        rec("T4", t4, f"併格差 {sorted(merges ^ H3_MERGES)}；列3 {row3}；列5 {r5}；列印標題 {ws.print_title_rows}；欄數 {ws.max_column}")
        sok, last, sv = _sign_ok(ws)
        body = [[ws.cell(r, c).value for c in range(1, 16)] for r in range(6, last)
                if any(not _blank(ws.cell(r, c).value) for c in range(1, 16))]
        per = defaultdict(list)
        for b in body:
            per[str(b[0])].append(b)
        bad = []
        for ln in lots:
            got = sorted((str(b[4]), str(b[5])) for b in per.get(ln, []))
            exp = sorted(own[ln].values())
            if got != exp:
                bad.append((ln, "人"))
                continue
            s = sum(float(b[3]) for b in per[ln])
            if abs(s - area[ln]) > 0.005 * len(per[ln]) + 1e-9 or any(abs(float(b[1]) - area[ln]) > 1e-9 for b in per[ln]):
                bad.append((ln, f"面積 Σ{s:.2f}≠{area[ln]}"))
        stray = [i for i, b in enumerate(body) if any(v in places for v in b)]
        exp_n = sum(len(v) for v in own.values())
        order_ok = [str(b[0]) for b in body] == sorted([str(b[0]) for b in body], key=lkey)
        rec("T5", len(body) == exp_n == 95 and set(per) == set(lots) and not bad and not stray and order_ok,
            f"列 {len(body)}（期望 {exp_n}）；不符 {bad[:4]}；資料列重出鄉鎮市區／段 {stray[:3]}；地號序 {order_ok}")
        rec("T6", sok, f"末列 {sv}")
    except Exception as e:  # noqa: BLE001
        rec("T4", False, f"附表3 讀取失敗 {type(e).__name__}: {e}")
    return res


N_T = 6

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    r = run(*sys.argv[1:3])
    sys.exit(0 if len(r) == N_T and all(x[1] for x in r) else 1)
