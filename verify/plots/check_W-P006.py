# -*- coding: utf-8 -*-
"""
W-P.006 之量測（單 §三 3）：附表 13 重劃前後土地分配清冊、附表 24 土地分配計算表。
獨立於 plots/：⛔ import 之（V6 唯以子行程呼叫 plots/post_tables.py 之命令列）；期望取自
配地結果_退縮*.json（plots/alloc_state.py）、地號對照表與重劃後顏色對照表（W-P.005·KL 定版）、重劃前顏色對照表
（W-P.001／002·KL 定版·歸戶號與區段單價）、data/地籍資料來源_匿名版.xlsx 之 U_LAND（逐欄自讀）、
verify/case_params_UC9898.json、手冊第 3990〜4045 列之範例、KL 2026-10-28 23:31 之答。

命令列：python verify/plots/check_W-P006.py <成果目錄> <重劃前成果目錄> <dict-revised.json> <退縮…> [--prog <post_tables.py>]
  <成果目錄>     ＝ 含 配地結果_退縮<t>m.json、地號對照表_退縮<t>m.xlsx、顏色對照表_重劃後_退縮<t>m.xlsx、
                   附表13_…_退縮<t>m.xlsx、附表24_…_退縮<t>m.xlsx。
  <重劃前成果目錄> ＝ plots/prelim_maps.py 之輸出（顏色對照表_重劃前.xlsx）。
各退縮印 V0〜V6 之 PASS／FAIL；全 PASS 者結束碼 0。
W-P.006_補令一（KL 2026-10-31 08:14）：縣有依管理機關分表（抵費地之管理人 花蓮縣政府、公共設施之管理人留空）；
取配面積／分配面積／Σa／Σa' 之可填與否（戶之 a 之和 ＝ 重劃前登記面積之和）；重劃前單價唯列所用之值；
表下之加註；欄名依手冊之附表 24 範例。
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from decimal import Decimal, ROUND_HALF_UP

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CASE = os.path.join(REPO, "verify", "case_params_UC9898.json")
XLSX = os.path.join(REPO, "data", "地籍資料來源_匿名版.xlsx")
PROG = os.path.join(REPO, "plots", "post_tables.py")
SECTION = "0013德安段"
U = {"sec": 0, "code": 1, "area": 6, "cur": 9, "seq": 23, "den": 28, "num": 29, "uid": 32, "name": 33,
     "addr": 34, "muid": 39, "mname": 40}
FMT = {"state": "配地結果_退縮{}m.json", "map": "地號對照表_退縮{}m.xlsx", "color": "顏色對照表_重劃後_退縮{}m.xlsx",
       "a13": "附表13_重劃前後土地分配清冊_退縮{}m.xlsx", "a24": "附表24_土地分配計算表_退縮{}m.xlsx"}
PRE_COLOR = "顏色對照表_重劃前.xlsx"
COUNTY = "花蓮縣"
T_A13 = "○○縣(市)第○○期○○○○市地重劃區重劃前後土地分配清冊"
T_A24 = "○○縣（市）第○○期○○○○市地重劃區土地分配計算表"
N_PUB = "依平均地權條例第 60 條規定，由土地所有權人共同負擔取得土地"
N_OFF = "依平均地權條例第 60 條規定，由土地所有權人以未建築土地折價抵付之土地"
N_A13 = "附註：重劃後分配面積與實地指界交接土地面積如有不符，應以地政機關實際登記面積為準。"
N_A24 = ("③暫以匯入之地籍資料所載之公告土地現值計（暫值）；依手冊，以重劃土地分配結果公告期滿時之"
         "當期公告土地現值為準")
N_ROUND = "計算式各數為取位後之顯示值，與重劃後應分配權利面積之差係取位所致"
N_PRE = ("評定之重劃前單價、原位置重劃前地價：配地計算所用之重劃前地價區段之區段地價（未取整、顯示至小數二位）；"
         "附表 18 之評定單價為其取至元之值")
N_SPLIT = "本宗之 a 含同戶他筆土地或本筆土地分屬他宗之部分，其逐筆之來源待配地程式提供，分配面積、Σa、Σa'暫空"
MGR_OFF = "花蓮縣政府"
CNTY = "縣有（{}）"
L_A24 = "①＝G×C÷(1−C)；②＝(Σa＋Σa')−G−①；③＝平均公告現值"
H6 = ("鄉鎮市區", "段", "小段", "地號", "登記面積(㎡)", "取配面積(㎡)", "所有權人", None, "持分比率", "評定單價(元／㎡)",
      None, "使用分區", "段", "小段", "地號", "面積(㎡)", "評定單價(元／㎡)", "預計繳領差額地價", None, "所有權人", None,
      "持分比率", "建號", None)
H5 = ("重劃前土地情形",) + (None,) * 9 + ("扣除各項負擔後應分配面積(㎡)", "重劃後土地情形") + (None,) * 11 + ("備註",)
H7 = (None,) * 6 + ("姓名", "國民身分證統一編號") + (None,) * 9 + ("繳（元）", "領（元）", "姓名", "國民身分證統一編號") + (None,) * 3
FIRST = 8
TOL_G = Decimal("0.02")        # 計算式以表上之值重算與 G 之差（G 為幾何二分之解·S 等取二位）
TOL_AREA = 0.05


def dq(x, nd=2):
    return Decimal(repr(float(x))).quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP)


def qd(x, nd):
    return x.quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP)


def nkey(s):
    p = str(s).split("-")
    return (int(p[0]), int(p[1]) if len(p) > 1 else 0)


def landno(code):
    m, s = divmod(int(code), 10000)
    return f"{m}" if s == 0 else f"{m}-{s}"


def blank(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", "null"))


# ---------------------------------------------------------------- 手冊之式（V0 以範例自證）
def burden(G, tot, C, post, cur):
    b1 = qd(G * C / (1 - C), 2)
    b2 = tot - G - b1
    y1 = qd(b1 * post, 0)
    y2 = None if cur is None else qd(b2 * cur, 0)
    return b1, b2, y1, y2


def eval_formula(s):
    """「G=〔Σ×(1−A×B)−F×l1×Rw−S×l2〕×(1−C)」⇒ (G, 右式之值)。"""
    m = re.fullmatch(r"([\d.]+)=〔([\d.]+)×\(1−([\d.]+)×([\d.]+)\)−([\d.]+)×([\d.]+)×([\d.]+)−([\d.]+)×([\d.]+)〕×\(1−([\d.]+)\)", s)
    if not m:
        return None
    g, t, a, b, f, l1, rw, s_, l2, c = (Decimal(v) for v in m.groups())
    return g, (t * (1 - a * b) - f * l1 * rw - s_ * l2) * (1 - c)


# ---------------------------------------------------------------- 讀取
class Ctx:
    pass


def read_uland():
    import openpyxl
    own, area, cur = {}, {}, {}
    for r in openpyxl.load_workbook(XLSX, read_only=True)["U_LAND"].iter_rows(values_only=True):
        if r[U["sec"]] != SECTION or blank(r[U["code"]]):
            continue
        try:
            ln = landno(r[U["code"]])
        except (TypeError, ValueError):
            continue
        area[ln] = float(r[U["area"]])
        cur.setdefault(ln, set()).add(r[U["cur"]])
        own.setdefault(ln, {}).setdefault(int(r[U["seq"]]), r)
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    reg = list(wb["重劃區地籍"].iter_rows(values_only=True))
    h = reg[0]
    place = {str(r[h.index("地號")]).strip(): (r[h.index("鄉鎮市")], r[h.index("地段")]) for r in reg[1:]
             if not blank(r[h.index("地號")])}
    rows = {}
    for ln, d in own.items():
        rows[ln] = [(str(r[U["name"]]), str(r[U["uid"]]), f"{int(r[U['num']])}/{int(r[U['den']])}", r[U["addr"]],
                     None if blank(r[U["mname"]]) else r[U["mname"]], None if blank(r[U["muid"]]) else r[U["muid"]])
                    for _, r in sorted(d.items())]
    return rows, area, cur, place


def load(out, pre, tag):
    import openpyxl
    c = Ctx()
    c.tag = tag
    c.p = {k: os.path.join(out, v.format(tag)) for k, v in FMT.items()}
    c.state = json.load(open(c.p["state"], encoding="utf-8"))
    ws = openpyxl.load_workbook(c.p["map"], data_only=True)["地號對照表"]
    c.no = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        if r and r[0] is not None and r[1] is not None:
            c.no[r[3] if r[3] is not None else r[0]] = str(r[5])
    pw = openpyxl.load_workbook(os.path.join(pre, PRE_COLOR), data_only=True)
    c.hh_of = {}
    for r in list(pw["重劃前歸戶圖"].iter_rows(values_only=True))[1:]:
        for ln in str(r[4]).split("、"):
            c.hh_of[ln] = int(r[1])
    c.price_of = {}
    for r in list(pw["重劃前地價圖"].iter_rows(values_only=True))[1:]:
        for ln in str(r[4]).split("、"):
            c.price_of[ln] = int(r[3])
    cw = openpyxl.load_workbook(c.p["color"], data_only=True)["分配後歸戶圖"]
    c.post_hh = {}
    for r in list(cw.iter_rows(values_only=True))[1:]:
        if isinstance(r[1], int):
            for no in str(r[3]).split("、"):
                c.post_hh[no] = r[1]
    c.a13 = openpyxl.load_workbook(c.p["a13"], data_only=True)
    c.a24 = openpyxl.load_workbook(c.p["a24"], data_only=True)
    c.own, c.area, c.cur, c.place = read_uland()
    cp = json.load(open(CASE, encoding="utf-8"))
    c.zoning = cp["global"]["分區"]
    c.lots = [x for x in c.state["宗"] if x["類別"] == "分配地"]
    c.others = [x for x in c.state["宗"] if x["類別"] != "分配地"]
    c.public = sorted(b for b, v in c.state["街廓"].items() if "宗地分配線方向" not in v)
    c.fin = c.state["財務"]
    tot = {}
    for x in c.lots:
        n = c.hh_of[x["原地號"]]
        k, s = tot.get(n, (0, Decimal(0)))
        tot[n] = (k + 1, s + dq(x["a"]))
    c.full = {n for n, (k, s) in tot.items()
              if abs(s - sum(dq(c.area[ln]) for ln, h in c.hh_of.items() if h == n)) <= Decimal("0.01") * k}
    c.single = {n for n in c.full if sum(1 for h in c.hh_of.values() if h == n) == 1}
    cats = {}
    for x in c.others:
        cats.setdefault("抵費地", []).append(c.no[x["配地暫編"]])
    for b in c.public:
        cats.setdefault(c.state["街廓"][b]["分類"], []).append(c.no[b])
    c.county = [CNTY.format(k) for k in sorted(cats, key=lambda k: (k != "抵費地", min(nkey(v) for v in cats[k])))]
    c.cat_of = {no: k for k, v in cats.items() for no in v}
    return c


def rows_of(ws, first=FIRST):
    out = []
    r = first
    while any(ws.cell(r, k).value is not None for k in range(1, 25)):
        out.append(tuple(ws.cell(r, k).value for k in range(1, 25)))
        r += 1
    return out, r


def money(c, x):
    post = Decimal(repr(c.fin["重劃後單價"][x["街廓"]]))
    return qd((dq(x["面積"]) - dq(x["G"])) * post, 0)


# ---------------------------------------------------------------- 量測
def v0(c):
    bad = []
    b1, b2, y1, y2 = burden(Decimal("202.83"), Decimal("340.64"), Decimal("0.07790626"), Decimal(162000),
                            Decimal(69600))
    if (b1, b2, y1, y2, y1 + y2) != (Decimal("17.14"), Decimal("120.67"), 2776680, 8398632, 11175312):
        bad.append(f"手冊之範例 ①②與負擔額不合：{b1} {b2} {y1} {y2}")
    g, val = eval_formula("202.83=〔340.64×(1−2.32758621×0.13579793)−0.00×0.00×0.0000−6.50×2.00〕×(1−0.07790626)")
    if dq(val) != g:
        bad.append(f"手冊之範例 G 式 {val} ≠ 202.83")
    return bad


def v1(c):
    """附表 13：一戶一表（歸戶號 ＝ 重劃前顏色對照表）＋縣有；表頭、題、附註、簽欄、首列登記人。"""
    bad = []
    hh = sorted(set(c.hh_of.values()))
    want = [f"歸戶號{n}" for n in hh] + c.county
    if c.a13.sheetnames != want:
        return [f"附表 13 之表 {c.a13.sheetnames}"]
    for ws in c.a13:
        if ws.cell(1, 1).value != T_A13:
            bad.append(f"{ws.title} 之題")
        for rr, h in ((5, H5), (6, H6), (7, H7)):
            if tuple(ws.cell(rr, k).value for k in range(1, 25)) != h:
                bad.append(f"{ws.title} 列 {rr} 之表頭")
        rows, r = rows_of(ws)
        if ws.cell(r + 1, 1).value != N_A13 or [ws.cell(r + 3, k).value for k in (1, 5, 9, 13, 17)] != \
                ["日期：", "承辦：", "股長：", "科長：", "局(處)長："]:
            bad.append(f"{ws.title} 之附註或簽欄")
        if ws.title.startswith("縣有"):
            mgr = MGR_OFF if ws.title == CNTY.format("抵費地") else None
            if ws.cell(3, 4).value != COUNTY or ws.cell(4, 4).value != mgr:
                bad.append(f"{ws.title} 之所有權人或管理人 {ws.cell(3, 4).value}、{ws.cell(4, 4).value}")
            continue
        n = int(ws.title[3:])
        lns = sorted((ln for ln, h in c.hh_of.items() if h == n), key=nkey)
        f = c.own[lns[0]][0]
        got = (ws.cell(3, 4).value, ws.cell(3, 9).value, ws.cell(3, 13).value, ws.cell(4, 4).value, ws.cell(4, 9).value)
        if got != (f[0], f[1], f[3], f[4], f[5]):
            bad.append(f"{ws.title} 之首列登記人 {got}")
    return bad


def v2(c):
    """附表 13 之重劃前：每戶之各宗 × 各所有權人一列（地號序·登記次序）；取配面積：戶之 a 之和 ＝ 登記面積之和者
    ＝ 登記面積，否則留空。"""
    bad = []
    for n in sorted(set(c.hh_of.values())):
        ws = c.a13[f"歸戶號{n}"]
        rows, _ = rows_of(ws)
        got = [r[:10] for r in rows if r[3] is not None]
        want = []
        for ln in sorted((ln for ln, h in c.hh_of.items() if h == n), key=nkey):
            for nm, uid, sh, *_ in c.own[ln]:
                want.append((c.place[ln][0], c.place[ln][1], None, ln, c.area[ln], c.area[ln] if n in c.full else None,
                             nm, uid, sh, c.price_of[ln]))
        if got != want:
            bad.append(f"歸戶號{n} 之重劃前 {len(got)}／{len(want)} 列不合")
        if any(r[:10] != (None,) * 10 for r in rows[len(got):]):
            bad.append(f"歸戶號{n} 之重劃前有多出之列")
    return bad


def v3(c):
    """附表 13 之重劃後：各分配地入其戶（歸戶號 ＝ 重劃後顏色對照表）、各所有權人一列、G／面積／單價／差額；
    縣有之表 ＝ 抵費地與公設（備註照手冊）；每一正式地號恰一次；面積之和 ＝ 全部街廓之和。"""
    bad = []
    seen = []
    tot = 0.0
    for x in c.lots:
        no = c.no[x["配地暫編"]]
        n = c.post_hh.get(no)
        if n is None or n != c.hh_of[x["原地號"]]:
            bad.append(f"{no} 之歸戶號 {n}")
            continue
        rows, _ = rows_of(c.a13[f"歸戶號{n}"])
        got = [r[10:23] for r in rows if r[14] == no]
        d = money(c, x)
        want = [(float(dq(x["G"])), c.zoning, None, None, no, float(dq(x["面積"])),
                 c.fin["重劃後單價"][x["街廓"]], float(d) if d >= 0 else None, float(-d) if d < 0 else None,
                 nm, uid, sh, None) for nm, uid, sh, *_ in c.own[x["原地號"]]]
        if got != want:
            bad.append(f"{no} 之重劃後之列不合（{len(got)}／{len(want)}）")
        seen.append(no)
        tot += float(dq(x["面積"]))
    rows = []
    for t in c.county:
        rr = rows_of(c.a13[t])[0]
        if any(c.cat_of.get(r[14]) != t[3:-1] for r in rr):
            bad.append(f"{t} 有他類之列")
        rows += rr
    want = []
    for x in c.others:
        want.append((c.no[x["配地暫編"]], c.zoning, c.fin["重劃後單價"][x["街廓"]], N_OFF, float(dq(x["面積"]))))
    for b in c.public:
        from shapely.geometry import Polygon
        want.append((c.no[b], c.state["街廓"][b]["分類"], None, N_PUB, float(dq(Polygon(c.state["街廓"][b]["多邊形"]).area))))
    want.sort(key=lambda t: nkey(t[0]))
    got = sorted(((r[14], r[11], r[16], r[23], r[15]) for r in rows), key=lambda t: nkey(t[0]))
    if got != want:
        bad.append(f"縣有之列不合：{[g for g, w in zip(got, want) if g != w][:3]}")
    if any(r[19] != COUNTY or r[21] != "1/1" or any(r[k] is not None for k in range(0, 11)) for r in rows):
        bad.append("縣有之所有權人、持分或重劃前欄不合")
    seen += [r[14] for r in rows]
    tot += sum(r[15] for r in rows)
    allno = []
    for ws in c.a13:
        allno += sorted({r[14] for r in rows_of(ws)[0] if r[14] is not None}, key=nkey)
    if sorted(allno, key=nkey) != sorted(seen, key=nkey) or len(set(allno)) != len(allno) \
            or sorted(allno, key=nkey) != sorted(c.no.values(), key=nkey):
        bad.append("正式地號非恰一次")
    from shapely.geometry import Polygon
    blk = sum(Polygon(v["多邊形"]).area for v in c.state["街廓"].values())
    if abs(tot - blk) > TOL_AREA:
        bad.append(f"面積之和 {tot:.2f} ≠ 街廓 {blk:.2f}")
    for ws in c.a13:
        if not ws.title.startswith("縣有") and any(r[23] is not None for r in rows_of(ws)[0]):
            bad.append(f"{ws.title} 有備註（本版無現金補償）")
    return bad


SECTIONS = ("重劃前", "計算分配有關資料", "重劃後")


def a24_vals(ws):
    """⇒ {(節, 標籤): [B, C, D 之值]}；節 ＝ 其上最近之「重劃前／計算分配有關資料／重劃後」（題之前為 ""）。"""
    out, sec = {}, ""
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(r, 1).value
        if lab is None:
            continue
        if lab in SECTIONS and all(ws.cell(r, k).value is None for k in range(2, 5)):
            sec = lab
            continue
        key = (sec, lab)
        if key in out:
            raise AssertionError(f"{ws.title} 之 {key} 重出")
        out[key] = [ws.cell(r, k).value for k in range(2, 5)]
    return out


def same(a, b):
    """二列之值相等（浮點以相對 1e-12 為度·Excel 之存值或少末位）。"""
    if a is None or b is None or len(a) != len(b):
        return a == b
    for x, y in zip(a, b):
        if isinstance(x, float) and isinstance(y, (int, float)) and not isinstance(y, bool):
            if abs(x - y) > 1e-12 * max(1.0, abs(y)):
                return False
        elif x != y:
            return False
    return True


def v4(c):
    """附表 24：一分配地一張（正式地號之序）；歸戶號、代表人、重劃前、計算分配有關資料、計算式（與 G 之差 ≤ 0.02）。"""
    bad = []
    lots = sorted(c.lots, key=lambda x: nkey(c.no[x["配地暫編"]]))
    want = [f"地號{c.no[x['配地暫編']]}" for x in lots]
    if c.a24.sheetnames != want:
        return [f"附表 24 之表 {c.a24.sheetnames}"]
    B, C = Decimal(repr(c.fin["B"])), Decimal(repr(c.fin["C"]))
    for x in lots:
        ws = c.a24[f"地號{c.no[x['配地暫編']]}"]
        v = a24_vals(ws)
        if ("", T_A24) not in v:
            bad.append(f"{ws.title} 之題")
        ln = x["原地號"]
        own = c.own[ln]
        rep = own[0][0] if len(own) == 1 else f"{own[0][0]}等{len(own)}人"
        post = Decimal(repr(c.fin["重劃後單價"][x["街廓"]]))
        pre = Decimal(repr(c.fin["重劃前單價"][x["重劃前區段"]]))
        A = post / pre
        tot = dq(x["a"])
        single = c.hh_of[ln] in c.single
        a0 = float(tot) if single else None
        P, K = "重劃前", "計算分配有關資料"
        exp = {
            ("", "歸戶號"): [c.hh_of[ln], "所有權(代表)人姓名", rep], (P, "段"): [c.place[ln][1], None, None],
            (P, "小段"): [None, None, None], (P, "地號"): [ln, None, None],
            (P, "登記面積(㎡)"): [c.area[ln], None, None], (P, "分配面積(㎡)"): [a0, None, None],
            (P, "共有人申請單獨分配持分比例"): [None, None, None], (P, "持分面積(㎡)"): [None, None, None],
            (P, "法定空地面積(㎡)"): [None, None, None],
            (P, "當期公告現值(元/㎡)"): [float(next(iter(c.cur[ln]))), None, None],
            (P, "評定之重劃前單價(元/㎡)"): [float(pre), None, None],
            "Σa(㎡)": [a0, None, None], "Σa'(㎡)": [0.0 if single else None, None, None],
            "Σ(a＋a')(㎡)": [float(tot), None, None], "原位置重劃前地價(元/㎡)": [float(pre), None, None],
            "預配位置重劃前原地平均單價(元/㎡)": [None, None, None], "重劃後宗地地價(元/㎡)": [float(post), None, None],
            "A＝重劃後宗地地價／預配位置重劃前地價": [float(A), None, None], "B": [float(B), None, None],
            "C": [float(C), None, None], "1−C": [float(1 - C), None, None], "S(m)": [float(dq(x["S"])), None, None],
            "W(m)": [float(dq(x["W"])), None, None], "F(m)": [float(dq(x["F"])), None, None],
            "l1": [float(dq(x["l側"])), None, None], "l2": [float(dq(x["l正"])), None, None],
            "Rw": [float(dq(x["Rw"] / 100, 4)), None, None],
        }
        exp = {(k if isinstance(k, tuple) else (K, k)): e for k, e in exp.items()}
        for k, e in exp.items():
            if not same(v.get(k), e):
                bad.append(f"{ws.title} 之「{k}」{v.get(k)} ≠ {e}")
        if c.price_of[ln] != int(qd(pre, 0)):
            bad.append(f"{ws.title} 之重劃前單價 {pre} 取整 ≠ 附表 18 之 {c.price_of[ln]}")
        f8 = lambda d: f"{qd(d, 8):.8f}"  # noqa: E731
        s = (f"{dq(x['G']):.2f}=〔{tot:.2f}×(1−{f8(A)}×{f8(B)})−{dq(x['F']):.2f}×{dq(x['l側']):.2f}×"
             f"{dq(x['Rw'] / 100, 4):.4f}−{dq(x['S']):.2f}×{dq(x['l正']):.2f}〕×(1−{f8(C)})")
        got = (v.get((K, "計算式")) or [None])[0]
        if got != s:
            bad.append(f"{ws.title} 之計算式 {got} ≠ {s}")
        else:
            g, val = eval_formula(got)
            if abs(val - g) > TOL_G:
                bad.append(f"{ws.title} 之計算式重算 {val:.4f} 與 G {g} 差逾 {TOL_G}")
    return bad


def v5(c):
    """附表 24 之重劃後與重劃費用負擔：①②③、負擔額、合計、差額、分配率、圖例與加註（KL 23:31 問三）。"""
    bad = []
    C = Decimal(repr(c.fin["C"]))
    for x in c.lots:
        no = c.no[x["配地暫編"]]
        v = a24_vals(c.a24[f"地號{no}"])
        post = Decimal(repr(c.fin["重劃後單價"][x["街廓"]]))
        G, tot, area = dq(x["G"]), dq(x["a"]), dq(x["面積"])
        cur = Decimal(str(next(iter(c.cur[x["原地號"]]))))
        hh_cur = {Decimal(str(next(iter(c.cur[ln])))) for ln, h in c.hh_of.items() if h == c.hh_of[x["原地號"]]}
        if dq(x["a"]) - dq(x["a原位次"]) > 0:
            cur = next(iter(hh_cur)) if len(hh_cur) == 1 else None
        b1, b2, y1, y2 = burden(G, tot, C, post, cur)
        d = money(c, x)
        box = ("■繳納□領取" if d > 0 else "□繳納■領取" if d < 0 else "□繳納□領取") + f"新臺幣 {abs(d):,} 元"
        R = "重劃後"
        total = None if y2 is None else float(y1 + y2)
        exp = {
            "段": [None, None, None], "小段": [None, None, None], "地號": [no, None, None],
            "評定之重劃後單價(元/㎡)": [float(post), None, None], "重劃費用負擔總額(元)": [total, None, None],
            "重劃後應分配權利面積(㎡)": [float(G), None, None],
            "實配面積(㎡)": [float(area), None, None], "分配率＝G／a": [float(G / tot), None, None],
            "面積(㎡)": [float(area), None, None], "土地使用分區": [c.zoning, None, None],
            "預計繳納(領取)差額地價總額": [box, None, None],
            "抵費地負擔": ["①面積(㎡)", "重劃後單價(元/㎡)", "負擔額(元)"],
            "①": [float(b1), float(post), float(y1)],
            "公共設施用地共同負擔": ["②面積(㎡)", "③平均公告現值(元/㎡)", "負擔額(元)"],
            "②": [float(b2), None if cur is None else float(cur), None if y2 is None else float(y2)],
            "合計(元)": [None, None, total],
            "預計應繳納差額地價(元)": [None, None, float(d) if d > 0 else 0.0],
            "預計應領取差額地價(元)": [None, None, float(-d) if d < 0 else 0.0],
        }
        exp = {(R, k): e for k, e in exp.items()}
        for k, e in exp.items():
            if not same(v.get(k), e):
                bad.append(f"地號{no} 之「{k}」{v.get(k)} ≠ {e}")
        single = c.hh_of[x["原地號"]] in c.single
        for n_ in (L_A24, N_A24, N_ROUND, N_PRE):
            if (R, n_) not in v:
                bad.append(f"地號{no} 之圖例或加註缺「{n_[:12]}」")
        if ((R, N_SPLIT) in v) == single:
            bad.append(f"地號{no} 之 Σa 暫空之註 {'多' if single else '缺'}")
        if v.get((R, "製表日期：")) != ["承辦：", "股長：", "科長："]:
            bad.append(f"地號{no} 之簽欄")
        if not (b2 > 0 and b1 > 0):
            bad.append(f"地號{no} 之 ① {b1}、② {b2} 非正")
    return bad


def _run(prog, state, mapx, moedict, out):
    cmd = [sys.executable, prog, "--state", state, "--moedict", moedict, "--out", out]
    if mapx:
        cmd += ["--map", mapx]
    return subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8")


def v6(c, moedict, prog):
    """現金補償之列（KL 23:31 問二甲）、對照表之改號帶入（KL 23:37）、缺鍵即停——以子行程試之。"""
    import openpyxl
    bad = []
    tmp = tempfile.mkdtemp(prefix="cW-P006_")
    try:
        have = {c.hh_of[x["原地號"]] for x in c.lots}
        none = sorted(set(c.hh_of.values()) - have)
        if not none:
            return ["示範資料無未分配之戶，無從試現金補償"]
        n = none[0]
        ln = sorted((ln for ln, h in c.hh_of.items() if h == n), key=nkey)[0]
        code = c.state["歸戶表"][ln]
        st = dict(c.state)
        st["現金補償"] = [{"歸戶": code, "金額": 1234567}]
        sp = os.path.join(tmp, os.path.basename(c.p["state"]))
        json.dump(st, open(sp, "w", encoding="utf-8"), ensure_ascii=False)
        r = _run(prog, sp, None, moedict, os.path.join(tmp, "a"))
        if r.returncode != 0:
            bad.append(f"現金補償之試跑失敗：{r.stderr[-300:]}")
        else:
            wb = openpyxl.load_workbook(os.path.join(tmp, "a", os.path.basename(c.p["a13"])), data_only=True)
            hit = [(ws.title, rr[23]) for ws in wb for rr in rows_of(ws)[0] if rr[23] and "現金補償" in str(rr[23])]
            ws = wb[f"歸戶號{n}"]
            post_blank = all(rr[10:23] == (None,) * 13 for rr in rows_of(ws)[0])
            if hit != [(f"歸戶號{n}", "未達最小分配面積標準二分之一，發給現金補償 1,234,567 元")] or not post_blank:
                bad.append(f"現金補償之備註 {hit}、重劃後留空 {post_blank}")
        lot_code = c.state["歸戶表"][c.lots[0]["原地號"]]
        st["現金補償"] = [{"歸戶": lot_code, "金額": 100}]
        json.dump(st, open(sp, "w", encoding="utf-8"), ensure_ascii=False)
        if _run(prog, sp, None, moedict, os.path.join(tmp, "b")).returncode == 0:
            bad.append("已有分配地之戶列現金補償而未停")
        st.pop("現金補償")
        json.dump(st, open(sp, "w", encoding="utf-8"), ensure_ascii=False)
        if _run(prog, sp, None, moedict, os.path.join(tmp, "c")).returncode == 0:
            bad.append("配地結果缺「現金補償」而未停")
        x1, x2 = sorted(c.lots, key=lambda x: nkey(c.no[x["配地暫編"]]))[:2]
        n1, n2 = c.no[x1["配地暫編"]], c.no[x2["配地暫編"]]
        wb = openpyxl.load_workbook(c.p["map"])
        ws = wb["地號對照表"]
        for rr in range(2, ws.max_row + 1):
            if ws.cell(rr, 4).value == x1["配地暫編"]:
                ws.cell(rr, 6, n2)
            elif ws.cell(rr, 4).value == x2["配地暫編"]:
                ws.cell(rr, 6, n1)
        mp = os.path.join(tmp, "map.xlsx")
        wb.save(mp)
        r = _run(prog, c.p["state"], mp, moedict, os.path.join(tmp, "d"))
        if r.returncode != 0:
            bad.append(f"改號之試跑失敗：{r.stderr[-300:]}")
        else:
            wb = openpyxl.load_workbook(os.path.join(tmp, "d", os.path.basename(c.p["a24"])), data_only=True)
            gk = ("重劃後", "重劃後應分配權利面積(㎡)")
            g1 = a24_vals(wb[f"地號{n2}"]).get(gk, [None])[0]
            g2 = a24_vals(wb[f"地號{n1}"]).get(gk, [None])[0]
            if (g1, g2) != (float(dq(x1["G"])), float(dq(x2["G"]))):
                bad.append(f"改號未帶入附表 24：{n2}→{g1}、{n1}→{g2}")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return bad


def run(fn, *a):
    """量測之一項 ⇒ 不合之列；量測中之例外（例：缺表）亦為不合。"""
    try:
        return fn(*a)
    except Exception as ex:  # noqa: BLE001
        return [f"量測中止：{type(ex).__name__}: {ex}"]


def main(argv):
    args = list(argv)
    prog = PROG
    if "--prog" in args:
        i = args.index("--prog")
        prog = args[i + 1]
        del args[i:i + 2]
    out, pre, moedict, tags = args[0], args[1], args[2], args[3:]
    ok = True
    for tag in tags:
        tag = f"{float(tag):g}"
        c = load(out, pre, tag)
        for name, fn in (("V0", v0), ("V1", v1), ("V2", v2), ("V3", v3), ("V4", v4), ("V5", v5)):
            bad = run(fn, c)
            ok &= not bad
            print(f"退縮 {tag} m {name} {'PASS' if not bad else 'FAIL'}" + ("" if not bad else "：" + "；".join(bad[:4])))
        bad = run(v6, c, moedict, prog)
        ok &= not bad
        print(f"退縮 {tag} m V6 {'PASS' if not bad else 'FAIL'}" + ("" if not bad else "：" + "；".join(bad[:4])))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
