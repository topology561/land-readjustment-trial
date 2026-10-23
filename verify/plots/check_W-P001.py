# -*- coding: utf-8 -*-
"""
W-P.001 之出圖量測（獨立於 plots/prelim_maps.py：⛔ import 之；期望值取自倉內之原始資料）。

命令列：python verify/plots/check_W-P001.py <重劃前圖.dxf> <附表03_重劃前土地清冊.xlsx>
  附表 2 取同目錄之「附表02_重劃前土地地號摘錄簿.xlsx」。
任一量測不過 ⇒ 結束碼 1。

量測（單 §三 3）：M0 讀回；M1 宗數與地號集合；M2 註記點與宗一一相符；M3 面積錨；
M3b 頂點距原圖界址線 ≤ 0.01 m；M4/M4b/M4c 歸戶圖；M5/M5b 地價圖；M6 圖層；M7 附表 2、3。
M6 之讀法：系統圖層「0」（DXF 規格強制）與「Defpoints」（ezdxf 存檔時必補·AutoCAD 之系統圖層）不計，
  惟其上須無任何圖元；其餘圖層表與圖元所用之圖層皆恰為三層。
"""
import functools
import json
import os
import sys
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DXF_SRC = os.path.join(REPO, "data", "V6_1.dxf")
XLSX_SRC = os.path.join(REPO, "data", "地籍資料來源_匿名版.xlsx")
CASE_PARAMS = os.path.join(REPO, "verify", "case_params_UC9898.json")

L_CAD, L_OWN, L_PRICE = "重劃前地籍線", "重劃前歸戶圖", "重劃前地價圖"
AREA_ANCHORS = {"628-1": 1183.62, "628-3": 13.15, "628-41": 562.14, "628-44": 3.80}
EXPECTED_HH = 33
EXPECTED_APPX3_ROWS = 95


def _blank(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", "null"))


def lkey(s):
    p = str(s).split("-")
    return (int(p[0]), int(p[1]) if len(p) > 1 else 0)


@functools.lru_cache(maxsize=None)
def source_truth():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX_SRC, read_only=True)
    rows = list(wb["重劃區地籍"].iter_rows(values_only=True))
    i = rows[0].index("地號")
    lots = [str(r[i]).strip() for r in rows[1:] if not _blank(r[i])]
    ls = set(lots)
    owners = defaultdict(dict)
    for r in wb["U_LAND"].iter_rows(values_only=True):
        if r[0] != "0013德安段" or _blank(r[1]):
            continue
        m, s = divmod(int(r[1]), 10000)
        ln = f"{m}" if s == 0 else f"{m}-{s}"
        if ln in ls:
            owners[ln].setdefault(r[23], (str(r[32]), str(r[33])))
    hh = {ln: frozenset(v.values()) for ln, v in owners.items()}
    v = json.load(open(CASE_PARAMS, encoding="utf-8"))["財務接線_v3"]
    price = {z: int(Decimal(str(d["單價_元每m2"])).quantize(Decimal(1), ROUND_HALF_UP))
             for z, d in v["重劃前區段_面積單價"].items()}
    zone = dict(v["原地號_區段"])
    return lots, owners, hh, zone, price


@functools.lru_cache(maxsize=None)
def src_lines():
    import ezdxf
    from shapely.geometry import LineString
    d = ezdxf.readfile(DXF_SRC)
    return [LineString([tuple(e.dxf.start)[:2], tuple(e.dxf.end)[:2]])
            for e in d.modelspace() if e.dxftype() == "LINE" and e.dxf.layer == "CADASTRAL_BOUND"]


def hatch_poly(h):
    from shapely.geometry import Polygon
    rings = []
    for p in h.paths:
        if hasattr(p, "vertices"):
            rings.append([(v[0], v[1]) for v in p.vertices])
    if len(rings) != 1:
        return None
    return Polygon(rings[0])


def hatch_rgb(h):
    return tuple(h.rgb) if h.rgb is not None else None


def run(dxf_path, appx3_path):
    results = []

    def rec(name, ok, msg=""):
        results.append((name, bool(ok), msg))
        print(f"{'PASS' if ok else 'FAIL'} {name} {msg}")

    lots, owners, hh, zone, price = source_truth()
    lotset = set(lots)

    # M0
    try:
        import ezdxf
        doc = ezdxf.readfile(dxf_path)
        msp = doc.modelspace()
        rec("M0", True, "ezdxf.readfile 讀回")
    except Exception as e:  # noqa: BLE001 — 讀不回即量測之失敗，照記
        rec("M0", False, f"讀回失敗：{type(e).__name__}: {e}")
        return results
    from shapely.geometry import Point, Polygon

    ents = list(msp)
    cad_polys = [Polygon([(p[0], p[1]) for p in e.get_points("xy")])
                 for e in ents if e.dxftype() == "LWPOLYLINE" and e.dxf.layer == L_CAD and e.closed]
    cad_txt = [(e.dxf.text, Point(e.dxf.align_point[0], e.dxf.align_point[1]))
               for e in ents if e.dxftype() == "TEXT" and e.dxf.layer == L_CAD]
    txt_set = [t for t, _ in cad_txt]
    rec("M1", len(cad_polys) == len(lots) and len(cad_txt) == len(lots) and set(txt_set) == lotset
        and len(set(txt_set)) == len(txt_set),
        f"多邊形 {len(cad_polys)}、地號 {len(cad_txt)}、期望 {len(lots)}；差 {sorted(set(txt_set) ^ lotset)}")

    # M2：註記點恰落於一宗、一一相符
    poly_of = {}
    bad = []
    used = defaultdict(list)
    for t, pt in cad_txt:
        hit = [i for i, pg in enumerate(cad_polys) if pg.contains(pt)]
        if len(hit) != 1:
            bad.append((t, len(hit)))
            continue
        poly_of[t] = hit[0]
        used[hit[0]].append(t)
    dup = {i: v for i, v in used.items() if len(v) > 1}
    rec("M2", not bad and not dup and len(poly_of) == len(cad_polys), f"不恰一宗 {bad}；一宗多號 {dup}")

    # M3 面積錨
    m3 = {}
    for ln, exp in AREA_ANCHORS.items():
        a = cad_polys[poly_of[ln]].area if ln in poly_of else None
        m3[ln] = (round(a, 3) if a is not None else None, exp)
    rec("M3", all(a is not None and abs(a - e) <= 0.05 for a, e in m3.values()), str(m3))

    # M3b 頂點距原圖界址線
    from shapely.strtree import STRtree
    lines = src_lines()
    tree = STRtree(lines)
    worst = 0.0
    for pg in cad_polys:
        for x, y in list(pg.exterior.coords):
            p = Point(x, y)
            j = tree.nearest(p)
            worst = max(worst, lines[j].distance(p))
    rec("M3b", worst <= 0.01 + 1e-9 and len(cad_polys) > 0, f"最大距離 {worst:.6f} m")

    def match_hatches(layer):
        out = {}
        miss = []
        for e in ents:
            if e.dxftype() != "HATCH" or e.dxf.layer != layer:
                continue
            hp = hatch_poly(e)
            if hp is None or not hp.is_valid:
                miss.append("壞邊界")
                continue
            rp = hp.representative_point()
            hit = [t for t, i in poly_of.items() if cad_polys[i].contains(rp)]
            if len(hit) != 1 or abs(cad_polys[poly_of[hit[0]]].area - hp.area) > 0.01:
                miss.append(hit)
                continue
            if not e.dxf.solid_fill:
                miss.append(("非實心", hit[0]))
                continue
            if hit[0] in out:
                miss.append(("重複", hit[0]))
            out[hit[0]] = hatch_rgb(e)
        return out, miss

    # M4 歸戶圖
    own, miss = match_hatches(L_OWN)
    rec("M4", set(own) == lotset and not miss and all(c is not None for c in own.values()),
        f"填色宗 {len(own)}；異常 {miss[:5]}")
    col_of_hh = defaultdict(set)
    for ln, c in own.items():
        col_of_hh[hh.get(ln)].add(c)
    hh_n = len(set(hh[ln] for ln in lots))
    same = all(len(v) == 1 for v in col_of_hh.values())
    rec("M4b", same and hh_n == EXPECTED_HH and len(col_of_hh) == hh_n,
        f"戶數 {hh_n}（期望 {EXPECTED_HH}）；同戶多色 {[k for k, v in col_of_hh.items() if len(v) > 1][:3]}")
    viol = []
    names = list(poly_of)
    for a_i in range(len(names)):
        for b_i in range(a_i + 1, len(names)):
            a, b = names[a_i], names[b_i]
            if hh.get(a) == hh.get(b):
                continue
            inter = cad_polys[poly_of[a]].boundary.intersection(cad_polys[poly_of[b]].boundary)
            if inter.length > 0.01 and own.get(a) == own.get(b):
                viol.append((a, b))
    rec("M4c", not viol and bool(own), f"共邊異戶同色 {viol[:5]}")

    # M5 地價圖
    pr, miss = match_hatches(L_PRICE)
    zc = defaultdict(set)
    for ln, c in pr.items():
        zc[zone.get(ln)].add(c)
    rec("M5", set(pr) == lotset and not miss and all(len(v) == 1 for v in zc.values())
        and len({next(iter(v)) for v in zc.values()}) == len(zc),
        f"填色宗 {len(pr)}；區段色 {dict((k, sorted(v)) for k, v in zc.items())}；異常 {miss[:5]}")
    ptxt = [(e.dxf.text, Point(e.dxf.align_point[0], e.dxf.align_point[1]))
            for e in ents if e.dxftype() == "TEXT" and e.dxf.layer == L_PRICE]
    got = {}
    for t, pt in ptxt:
        hit = [ln for ln, i in poly_of.items() if cad_polys[i].contains(pt)]
        if len(hit) == 1:
            got.setdefault(hit[0], []).append(t)
    wrong = [(ln, got.get(ln), price[zone[ln]]) for ln in lots
             if got.get(ln) != [str(price[zone[ln]])]]
    rec("M5b", not wrong and len(ptxt) == len(lots), f"註記 {len(ptxt)}；不符 {wrong[:5]}")

    # M6 圖層
    tbl = {l.dxf.name for l in doc.layers} - {"0", "Defpoints"}
    usedl = {e.dxf.layer for e in ents}
    rec("M6", tbl == {L_CAD, L_OWN, L_PRICE} and usedl == {L_CAD, L_OWN, L_PRICE},
        f"圖層表 {sorted(tbl)}；圖元所用 {sorted(usedl)}")

    # M7 附表 2、3
    import openpyxl
    ok7, msg7 = True, []
    try:
        r3 = list(openpyxl.load_workbook(appx3_path, read_only=True).active.iter_rows(values_only=True))
        h3 = r3[0]
        i_ln = h3.index("地號")
        body3 = [r for r in r3[1:] if any(c is not None for c in r)]
        exp_rows = sum(len(v) for v in owners.values())
        lots3 = {str(r[i_ln]) for r in body3}
        per = defaultdict(int)
        for r in body3:
            per[str(r[i_ln])] += 1
        bad_per = [ln for ln in lots if per[ln] != len(owners[ln])]
        if not (len(body3) == exp_rows == EXPECTED_APPX3_ROWS and lots3 == lotset and not bad_per):
            ok7 = False
        msg7.append(f"附表3 {len(body3)} 列（期望 {exp_rows}）；每宗人數不符 {bad_per[:5]}")
    except Exception as e:  # noqa: BLE001
        ok7 = False
        msg7.append(f"附表3 讀取失敗 {type(e).__name__}: {e}")
    p2 = os.path.join(os.path.dirname(os.path.abspath(appx3_path)), "附表02_重劃前土地地號摘錄簿.xlsx")
    try:
        r2 = list(openpyxl.load_workbook(p2, read_only=True).active.iter_rows(values_only=True))
        i2 = r2[0].index("地號")
        seq = [str(r[i2]) for r in r2[1:] if any(c is not None for c in r)]
        exp_seq = sorted(lots, key=lkey)
        if seq != exp_seq:
            ok7 = False
        msg7.append(f"附表2 {len(seq)} 列；序 {'合' if seq == exp_seq else '不合'}")
    except Exception as e:  # noqa: BLE001
        ok7 = False
        msg7.append(f"附表2 讀取失敗 {type(e).__name__}: {e}")
    rec("M7", ok7, "；".join(msg7))
    return results


def main(argv):
    if len(argv) != 2:
        print("用法：check_W-P001.py <dxf> <附表3.xlsx>")
        return 2
    res = run(*argv)
    n_ok = sum(1 for _, ok, _ in res if ok)
    print(f"結果：{n_ok}/{len(res)} 通過")
    return 0 if res and all(ok for _, ok, _ in res) and len(res) == 12 else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
