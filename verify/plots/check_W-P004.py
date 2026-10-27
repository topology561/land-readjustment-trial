# -*- coding: utf-8 -*-
"""
W-P.004：附表 20、22（23 同式）之量測。
獨立於 plots/：⛔ import 之；期望取自倉內原始資料（data/V6.dxf、data/地籍資料來源_匿名版.xlsx、
verify/case_params_UC9898.json）、手冊（參考/市地重劃作業手冊.txt 第 3077〜3200、3292〜3304、3336〜3450 列）、
KL 之錨（case_params 之 斷言錨 與 _斷言錨_身分）與 KL 2026-10-27 23:15、23:20 之答。

命令列：python verify/plots/check_W-P004.py <成果目錄>
  成果目錄 ＝ plots/burden_tables.py --out 之目錄。
任一不過 ⇒ 結束碼 1。
S0 二檔皆在、可讀；S1 附表 20 之表樣；S2 附表 20 之列（區塊、合計面積、合計＝特別＋一般）；
S3 附表 20 之臨街地特別負擔（歸屬、各列乘 1−C 後取位、小計＝錨）；S4 附表 22 第一〜五項；
S5 附表 22 第六〜九項；S6 附表 22 第十〜十三項；S7 附表 22 之備註、簽欄與表下之註。
"""
import io
import json
import os
import re
import sys
from decimal import Decimal, ROUND_HALF_UP

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DXF = os.path.join(REPO, "data", "V6.dxf")
XLSX_SRC = os.path.join(REPO, "data", "地籍資料來源_匿名版.xlsx")
CASE = os.path.join(REPO, "verify", "case_params_UC9898.json")
SECTION_NAME = "0013德安段"
FILES = {20: "附表20_公共設施用地負擔面積統計表.xlsx", 22: "附表22_計算負擔總計表.xlsx"}

# KL 之錨（case_params 之 _斷言錨_身分：KL Tab7 C 值明細／Step H 截圖；斷言錨）
ANCHOR_A20 = (Decimal("12146.56"), Decimal("957.58"), Decimal("11188.98"))
ANCHOR_B, ANCHOR_C, ANCHOR_RATE = Decimal("0.171043"), Decimal("0.091319"), Decimal("0.40712387")
# KL 2026-10-27 23:15 問五乙（監看窗之讀第 5 點）、23:20 甲；問四甲（私有 61、公有 2）
ANCHOR_PRE = (Decimal("1336204754"), Decimal("38500.68"))
ANCHOR_RISE = Decimal("1.871619")
ANCHOR_OWNERS = (61, 2)
ANCHOR_MAP_PRE = (Decimal("1344405805"), Decimal("38498.68"))   # 圖面面積計之重劃前總地價與平均地價

A20_TITLE = "○○縣(市)第○○期○○○○市地重劃區公共設施用地負擔面積統計表"
A20_HEAD = ["編號", "名稱", "合計面積（平方公尺）", "臨街地特別負擔面積（平方公尺）", "一般負擔面積（平方公尺）", "備註"]
A20_SIGN = ["製表日期:　年　月　日", "承辦:", "股長:", "科長:"]
A22_TITLE = "○○縣(市)第○期○○○○市地重劃區計算負擔總計表"
A22_SIGN = ["承辦", "股長", "科長", "局(處)長", "製表日期:　年　月　日"]
NOTE_KL = "一般負擔係數所用之重劃前地價，係以各宗地之圖面面積計"
FACES = ("正面", "左側", "右側")


def _blank(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", "null"))


def D(x):
    return Decimal(repr(x)) if isinstance(x, float) else Decimal(str(x))


def rhu(x, nd):
    return D(x).quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP)


def _ws(out, n):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(out, FILES[n]))
    if len(wb.worksheets) != 1:
        raise AssertionError(f"附表 {n} 之工作表數 {len(wb.worksheets)}≠1")
    return wb.worksheets[0]


def _last(ws):
    last = 0
    for row in ws.iter_rows():
        for c in row:
            if not _blank(c.value):
                last = max(last, c.row)
    return last


def _nonblank(ws, r):
    return [ws.cell(r, c).value for c in range(1, ws.max_column + 1) if not _blank(ws.cell(r, c).value)]


# ---------------------------------------------------------------- 獨立之讀
def src_dxf():
    """⇒ 公設區塊 {編號: (名稱, Polygon)}、街廓 {編號: Polygon}、{圖層: [LineString]}。"""
    import ezdxf
    from shapely.geometry import LineString, Point, Polygon
    with open(DXF, "rb") as f:
        msp = ezdxf.read(io.TextIOWrapper(f, encoding="cp950", newline=None)).modelspace()
    polys = [Polygon([(v.dxf.location.x, v.dxf.location.y) for v in e.vertices])
             for e in msp.query('POLYLINE[layer=="BLOCK"]')]
    polys += [Polygon(list(e.get_points("xy"))) for e in msp.query('LWPOLYLINE[layer=="BLOCK"]')]
    texts = [(e.dxf.text.strip(), Point(e.dxf.insert.x, e.dxf.insert.y)) for e in msp.query('TEXT[layer=="BLOCK"]')]
    pub, lot = {}, {}
    for p in polys:
        labs = [t for t, pt in texts if p.contains(pt)]
        if len(labs) != 1:
            continue
        no, name = labs[0].split("-", 1)
        if re.match(r"^R\d+$", no):
            lot[no] = p
        else:
            pub[no] = (name, p)
    lines = {lay: [LineString([(e.dxf.start.x, e.dxf.start.y), (e.dxf.end.x, e.dxf.end.y)])
                   for e in msp.query(f'LINE[layer=="{lay}"]')] for lay in ("FRONT_LINE", "SIDE_LINE")}
    return pub, lot, lines


def src_case():
    d = json.load(open(CASE, encoding="utf-8"))
    return d["blocks"], d["財務接線_v3"]


def src_ledger():
    """⇒ {地號: (面積, 公私別, {統一編號})}（U_LAND·德安段·「重劃區地籍」之 53 宗）。"""
    import openpyxl
    wb = openpyxl.load_workbook(XLSX_SRC, read_only=True)
    rows = list(wb["重劃區地籍"].iter_rows(values_only=True))
    il = rows[0].index("地號")
    lots = {str(r[il]).strip() for r in rows[1:] if not _blank(r[il])}
    out, head = {}, None
    for r in wb["U_LAND"].iter_rows(values_only=True):
        if head is None:
            head = list(r)
            ia, ipp, iu = head.index("面積"), head.index("公私別"), head.index("所有權統一編號")
            continue
        if r[0] != SECTION_NAME or _blank(r[1]):
            continue
        m, s = divmod(int(r[1]), 10000)
        ln = f"{m}" if s == 0 else f"{m}-{s}"
        if ln not in lots:
            continue
        a, pp, u = out.get(ln, (D(r[ia]), r[ipp], set()))
        if a != D(r[ia]) or pp != r[ipp]:
            raise AssertionError(f"{ln} 之面積或公私別不一")
        u.add(str(r[iu]).strip())
        out[ln] = (a, pp, u)
    if set(out) != lots:
        raise AssertionError(f"U_LAND 缺宗：{sorted(lots - set(out))}")
    return out


def _post_total(v):
    return sum(rhu(D(x["面積_m2"]) * D(x["單價_元每m2"]), 2) for x in v["後街廓_面積單價"].values())


def _cost(v):
    m = v["中繼錨"]
    return Decimal(int(m["工程費_元"]) + int(m["重劃費_元"]) + int(v["貸款利息_元"]))


def expect_a20(pub, lot, lines, blocks, v):
    """⇒ [(編號, 名稱, 合計, 特別, 一般)]（期望之序：道路在前、各依編號）、{編號: 未乘之特別負擔}、公設圖面面積之和。
    特別負擔之歸屬（與 plots/ 異法）：街廓邊之 1 m 緩衝區扣除該街廓後，與各公設區塊之交集面積最大者。"""
    raw = {k: Decimal(0) for k in pub}
    for b in blocks:
        for face in FACES:
            f = blocks[b][face]
            if D(f["期望負擔面積_m2"]) == 0:
                continue
            lay = "FRONT_LINE" if face == "正面" else "SIDE_LINE"
            hit = [ln for ln in lines[lay] if abs(ln.length - f["期望長度_m"]) < 5e-4]
            if len(hit) != 1:
                raise AssertionError(f"{b}.{face} 之線 {len(hit)} 條")
            zone = hit[0].buffer(1.0).difference(lot[b])
            road = max(pub, key=lambda k: zone.intersection(pub[k][1]).area)
            raw[road] += D(f["期望負擔面積_m2"])
    omc = 1 - _cost(v) / _post_total(v)

    def order(k):
        m = re.match(r"^(\D*)(\d*)$", k)
        return (0 if "道路" in pub[k][0] else 1, m.group(1), int(m.group(2) or 0))
    rows = []
    for k in sorted(pub, key=order):
        t = rhu(pub[k][1].area, 2)
        sp = rhu(raw[k] * omc, 2)
        rows.append((k, pub[k][0], t, sp, t - sp))
    return rows, raw, sum(D(pub[k][1].area) for k in pub)


# ---------------------------------------------------------------- 量測
def s0(out):
    import openpyxl
    for f in FILES.values():
        p = os.path.join(out, f)
        if not os.path.isfile(p) or os.path.getsize(p) == 0:
            return False, f"缺或空：{f}"
        if _last(openpyxl.load_workbook(p).worksheets[0]) == 0:
            return False, f"無內容：{f}"
    return True, "二檔皆在、可讀、非空"


def _a20_body(ws):
    r, rows = 3, []
    while ws.cell(r, 2).value != "小　計":
        if r > 200:
            raise AssertionError("無小計之列")
        rows.append([ws.cell(r, c).value for c in range(1, 7)])
        r += 1
    return rows, r


def s1(out):
    ws = _ws(out, 20)
    if ws.cell(1, 1).value != A20_TITLE or "A1:F1" not in {str(m) for m in ws.merged_cells.ranges}:
        return False, f"題：{ws.cell(1, 1).value!r}"
    if [ws.cell(2, c).value for c in range(1, 7)] != A20_HEAD:
        return False, f"表頭：{[ws.cell(2, c).value for c in range(1, 7)]}"
    rows, rs = _a20_body(ws)
    if _nonblank(ws, rs + 1) != A20_SIGN or _last(ws) != rs + 1:
        return False, f"簽欄（列 {rs + 1}）：{_nonblank(ws, rs + 1)}"
    for r in range(3, rs + 1):
        for c in (3, 4, 5):
            if ws.cell(r, c).number_format != "#,##0.00":
                return False, f"列 {r} 欄 {c} 之格式 {ws.cell(r, c).number_format}"
    return True, f"題、表頭 6 欄、{len(rows)} 列、小計列 {rs}、簽欄、面積二位"


def s2(out, exp):
    ws = _ws(out, 20)
    rows, _ = _a20_body(ws)
    got = [(r[0], r[1], rhu(r[2], 2)) for r in rows]
    want = [(e[0], e[1], e[2]) for e in exp]
    if got != want:
        return False, f"列：{got} ≠ {want}"
    for r in rows:
        if rhu(r[2], 2) != rhu(r[3], 2) + rhu(r[4], 2) or not _blank(r[5]):
            return False, f"{r[0]}：合計≠特別＋一般 或備註有值 {r}"
    return True, f"{len(rows)} 列＝圖上之公設區塊（{'、'.join(g[0] for g in got)}）、合計面積＝圖面面積二位、合計＝特別＋一般"


def s3(out, exp, raw):
    ws = _ws(out, 20)
    rows, rs = _a20_body(ws)
    got = [rhu(r[3], 2) for r in rows]
    want = [e[3] for e in exp]
    if got != want:
        return False, f"特別負擔：{got} ≠ {want}（未乘 {[raw[e[0]] for e in exp]}）"
    sub = tuple(rhu(ws.cell(rs, c).value, 2) for c in (3, 4, 5))
    if sub != ANCHOR_A20 or sub != tuple(sum(rhu(r[c], 2) for r in rows) for c in (2, 3, 4)):
        return False, f"小計 {sub} ≠ 錨 {ANCHOR_A20} 或 ≠ 各列之和"
    if sum(raw.values()) != Decimal("1053.80"):
        return False, f"未乘之和 {sum(raw.values())}"
    return True, f"特別負擔 {'、'.join(str(x) for x in got)}（未乘之和 1053.80）；小計 {sub[0]}／{sub[1]}／{sub[2]}＝錨"


def _find(ws, lab, col):
    rs = [r for r in range(1, ws.max_row + 1) if ws.cell(r, col).value == lab]
    if len(rs) != 1:
        raise AssertionError(f"「{lab}」於欄 {col} 之列數 {len(rs)}")
    return rs[0]


def _val(ws, lab, col, nf=None):
    r = _find(ws, lab, col)
    c = ws.cell(r, col + 1)
    if nf is not None and c.number_format != nf:
        raise AssertionError(f"「{lab}」之格式 {c.number_format}≠{nf}")
    return c.value


def s4(out, ledger, v):
    ws = _ws(out, 22)
    if ws.cell(1, 1).value != A22_TITLE:
        return False, f"題：{ws.cell(1, 1).value!r}"
    priv = sum(a for a, pp, _ in ledger.values() if pp == "私")
    pub = sum(a for a, pp, _ in ledger.values() if pp == "公")
    exact = D(v["重劃區總面積_m2_精確"])
    A = "#,##0.00"
    want = [("重劃區實測面積", rhu(exact, 2)), ("1.私有土地面積", rhu(priv, 2)), ("2.公有土地面積", rhu(pub, 2)),
            ("3.未登記土地面積", Decimal(0)), ("4.實測誤差面積", Decimal(0)),
            ("原公有道路溝渠河川及未登記等土地抵充面積", Decimal(0))]
    for lab, w in want:
        g = _val(ws, lab, 3, A)
        if rhu(g, 2) != w:
            return False, f"「{lab}」{g} ≠ {w}"
    r = _find(ws, "原公有道路溝渠河川及未登記等土地抵充面積", 3)
    if [rhu(ws.cell(r + i, 4).value, 2) for i in (1, 2)] != [0, 0]:
        return False, "抵充之二目非 0"
    up = {u for _, pp, us in ledger.values() if pp == "私" for u in us}
    uq = {u for _, pp, us in ledger.values() if pp == "公" for u in us}
    n = (len(up), len(uq))
    if n != ANCHOR_OWNERS or up & uq:
        return False, f"人數 {n} ≠ 錨 {ANCHOR_OWNERS}"
    if [_val(ws, "重劃區土地所有權人數", 3), _val(ws, "1.私有：", 3), _val(ws, "2.公有：", 3)] != [sum(n), n[0], n[1]]:
        return False, "人數之格"
    r4 = _find(ws, "私有土地所有權人申請（同意）辦理重劃情形　人數：", 3)
    if not _blank(ws.cell(r4, 4).value) or not _blank(ws.cell(r4 + 1, 4).value):
        return False, "第四項有值（資料所無·須留空）"
    zone = v["原地號_區段"]
    price = {z: rhu(d["單價_元每m2"], 0) for z, d in v["重劃前區段_面積單價"].items()}
    tot = sum(a * price[zone[ln]] for ln, (a, _, _) in ledger.items())
    avg = rhu(tot / (priv + pub), 2)
    g_tot = _val(ws, "評定之重劃前平均地價及總地價　總地價：", 3, "#,##0")
    g_avg = _val(ws, "每平方公尺地價：", 3, "#,##0.00")
    if (rhu(tot, 0), avg) != ANCHOR_PRE or (rhu(g_tot, 0), rhu(g_avg, 2)) != ANCHOR_PRE:
        return False, f"第五項 {g_tot}／{g_avg}；期 {rhu(tot, 0)}／{avg}；錨 {ANCHOR_PRE}"
    return True, (f"實測 {rhu(exact, 2)}、私 {rhu(priv, 2)}、公 {rhu(pub, 2)}、未登記與誤差 0、抵充 0；"
                  f"人數 {sum(n)}（{n[0]}＋{n[1]}）；第四項空；第五項 {ANCHOR_PRE[0]}／{ANCHOR_PRE[1]}")


def s5(out, exp, v):
    ws = _ws(out, 22)
    sub = tuple(sum(e[i] for e in exp) for i in (2, 3, 4))
    A = "#,##0.00"
    got = (rhu(_val(ws, "公共設施用地扣除原公有道路溝渠河川及未登記土地等抵充土地後共同負擔總面積", 3, A), 2),
           rhu(_val(ws, "1.臨街地特別負擔×(1-C)", 3, A), 2), rhu(_val(ws, "2.一般負擔", 3, A), 2))
    if got != sub or got != ANCHOR_A20:
        return False, f"第六項 {got} ≠ 附表 20 之小計 {sub}／錨"
    cost = _cost(v)
    if D(_val(ws, "費用負擔總額：1.費用總額", 3, "#,##0")) != cost or cost != Decimal(150053099):
        return False, f"費用 {cost}"
    if not _blank(ws.cell(_find(ws, "2.補助金額", 3), 4).value):
        return False, "補助金額有值（資料所無·須留空）"
    pt = _post_total(v)
    g = (rhu(_val(ws, "評定之重劃後平均地價及總地價　總地價：", 8, "#,##0"), 0), rhu(_val(ws, "每平方公尺地價：", 8, A), 2))
    if g != (rhu(pt, 0), rhu(v["重劃後平均地價_元每m2"], 2)) or pt != Decimal("1643176255.06"):
        return False, f"第八項 {g}（Σ {pt}）"
    alloc = rhu(D(v["重劃區總面積_m2_精確"]) - sub[0], 2)
    if rhu(_val(ws, "重劃後可分配土地面積", 8, A), 2) != alloc or alloc != Decimal("22803.33"):
        return False, f"第九項 ≠ {alloc}"
    return True, f"第六項＝附表 20 小計＝錨；費用 {cost}、補助空；第八項 {g[0]}／{g[1]}；第九項 {alloc}"


def s6(out, v, exp, pub_dxf):
    ws = _ws(out, 22)
    K = "0.000000"
    pre = rhu(_val(ws, "每平方公尺地價：", 3), 2)
    post = rhu(_val(ws, "每平方公尺地價：", 8), 2)
    rise = rhu(_val(ws, "重劃前後每平方公尺平均地價上漲率：", 8, K), 6)
    if rise != rhu(post / pre, 6) or rise != ANCHOR_RISE:
        return False, f"上漲率 {rise}（期 {rhu(post / pre, 6)}·錨 {ANCHOR_RISE}）"
    exact = D(v["重劃區總面積_m2_精確"])
    gen = sum(e[4] for e in exp)
    post_full = D(v["重劃後平均地價_元每m2"])
    b = rhu(gen * D(v["重劃前平均地價_元每m2"]) / (post_full * exact), 6)
    c = rhu(_cost(v) / _post_total(v), 6)
    gb = rhu(_val(ws, "公共設施用地一般負擔係數：", 8, K), 6)
    gc = rhu(_val(ws, "費用負擔係數：", 8, K), 6)
    if (gb, gc) != (b, c) or (b, c) != (ANCHOR_B, ANCHOR_C):
        return False, f"B／C {gb}／{gc}（期 {b}／{c}·錨 {ANCHOR_B}／{ANCHOR_C}）"
    r1, r2 = pub_dxf / exact, _cost(v) / (post_full * exact)
    want = (rhu(ANCHOR_RATE, 4), rhu(r1, 4), rhu(r2, 4))
    P = "0.00%"
    got = (rhu(_val(ws, "平均負擔比率：", 8, P), 4), rhu(_val(ws, "1.公共設施用地負擔比率：", 8, P), 4),
           rhu(_val(ws, "2.費用負擔比率：", 8, P), 4))
    if got != want or rhu(r1 + r2, 8) != ANCHOR_RATE:
        return False, f"比率 {got} ≠ {want}（r1+r2 {rhu(r1 + r2, 8)}）"
    return True, f"上漲率 {rise}＝{post}÷{pre}；B {gb}、C {gc}＝錨；比率 {got[0]}＝{got[1]}＋{got[2]}；格式六位與 0.00%"


def s7(out, v, ledger):
    ws = _ws(out, 22)
    texts = [str(ws.cell(r, c).value) for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)
             if not _blank(ws.cell(r, c).value)]
    diff = rhu(D(v["重劃區總面積_m2_精確"]) - sum(a for a, _, _ in ledger.values()), 2)
    need = [f"實測面積較登記面積多 {diff:,.2f} 平方公尺，係登記與實測之差，非未登記地",
            "1.公共設施用地負擔比率", "2.費用負擔比率", "備註"]
    miss = [t for t in need if not any(t in x for x in texts)]
    if miss or diff != Decimal("243.89"):
        return False, f"備註缺：{miss}（差 {diff}）"
    last = _last(ws)
    note = ws.cell(last, 1).value or ""
    if not note.startswith("註：" + NOTE_KL) or f"{ANCHOR_MAP_PRE[0]:,}" not in note or f"{ANCHOR_MAP_PRE[1]:,}" not in note:
        return False, f"表下之註：{note!r}"
    if _nonblank(ws, last - 1) != A22_SIGN:
        return False, f"簽欄：{_nonblank(ws, last - 1)}"
    zt = sum(D(z["面積_m2"]) * D(z["單價_元每m2"]) for z in v["重劃前區段_面積單價"].values())
    za = sum(D(z["面積_m2"]) for z in v["重劃前區段_面積單價"].values())
    if (rhu(zt, 0), rhu(zt / za, 2)) != ANCHOR_MAP_PRE:
        return False, f"圖面之重劃前地價 {rhu(zt, 0)}／{rhu(zt / za, 2)}"
    return True, f"備註（二比率之式、差額 {diff}）、簽欄、表下之註（KL 逐字·{ANCHOR_MAP_PRE[0]:,}／{ANCHOR_MAP_PRE[1]:,}）"


def run(out, src=None, quiet=False):
    """⇒ 不過之項名之列。src ＝ load_src() 之值（省略者即讀之）。"""
    exp, raw, pub_dxf, ledger, v = load_src() if src is None else src
    tests = [
        ("S0", lambda: s0(out)),
        ("S1", lambda: s1(out)),
        ("S2", lambda: s2(out, exp)),
        ("S3", lambda: s3(out, exp, raw)),
        ("S4", lambda: s4(out, ledger, v)),
        ("S5", lambda: s5(out, exp, v)),
        ("S6", lambda: s6(out, v, exp, pub_dxf)),
        ("S7", lambda: s7(out, v, ledger)),
    ]
    fails = []
    for name, fn in tests:
        try:
            ok, msg = fn()
        except Exception as e:  # noqa: BLE001  量測之例外即不過
            ok, msg = False, f"{type(e).__name__}: {e}"
        if not quiet:
            print(f"{'PASS' if ok else 'FAIL'} {name} {msg}")
        if not ok:
            fails.append(name)
    if not quiet:
        print(f"{len(tests) - len(fails)}/{len(tests)}" + (f"；不過：{fails}" if fails else ""))
    return fails


def load_src():
    pub, lot, lines = src_dxf()
    blocks, v = src_case()
    exp, raw, pub_dxf = expect_a20(pub, lot, lines, blocks, v)
    return exp, raw, pub_dxf, src_ledger(), v


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 1:
        print(__doc__)
        return 2
    return 1 if run(argv[0]) else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
