# -*- coding: utf-8 -*-
"""
W-P.003：附表 4、14、15、18、19、21、25 之量測。
獨立於 plots/：⛔ import 之；期望取自倉內原始資料（data/地籍資料來源_匿名版.xlsx、verify/case_params_UC9898.json）、
手冊（參考/市地重劃作業手冊.txt 第 879〜899、2771〜2773、2851〜2853、3312〜3330、4048〜4078 列）與 KL 之錨。

命令列：python verify/plots/check_W-P003.py <成果目錄>
  成果目錄 ＝ plots/appx_tables.py --out 之目錄。
任一不過 ⇒ 結束碼 1。
P0 七檔皆在、可讀；P1 附表 4；P2 附表 14；P3 附表 15；P4 附表 18；P5 附表 19；P6 附表 21；P7 附表 25。
"""
import hashlib
import json
import os
import sys
from decimal import Decimal, ROUND_HALF_UP

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
XLSX_SRC = os.path.join(REPO, "data", "地籍資料來源_匿名版.xlsx")
CASE = os.path.join(REPO, "verify", "case_params_UC9898.json")
SECTION_NAME = "0013德安段"
FILES = {
    4: "附表04_抵充土地清冊.xlsx",
    14: "附表14_他項權利清冊.xlsx",
    15: "附表15_限制登記清冊.xlsx",
    18: "附表18_重劃前地價評議表.xlsx",
    19: "附表19_重劃後地價評議表.xlsx",
    21: "附表21_重劃後可分配土地街廓面積地價計算表.xlsx",
    25: "附表25_街角地側面道路負擔百分率累積表.xlsx",
}
# KL 之錨（verify/case_params_UC9898.json 之 財務接線_v3：C分母形式 Σ(後街廓面積×單價) ＝ 1643176255.06；
#   重劃後平均地價之來源注 Σ可建築面積 22803.33）
ANCHOR_SUM_TOTAL = Decimal("1643176255.06")
ANCHOR_SUM_AREA = Decimal("22803.33")
ANCHOR_ZONE = {"a": (43909, 5), "b": (37323, 48)}     # W-P.001 之區段地價（個位）與宗數
# 手冊附表 25 之 171 格（「W.十分位:值」以「;」接）之 sha256（發單端自手冊第 4053〜4076 列解析）
RW_SHA = "3ddefae4f4ab41a16eaedb7c057050aa921580027950d3a0efbf7e34abf4d8fa"
RW_SPOT = {(1, 0): 17.36, (1, 1): 18.08, (5, 3): 45.85, (12, 0): 80.00, (17, 9): 99.75, (18, 0): 100.00}

A4_HEAD = ["鄉鎮市區", "段", "小段", "地號", "登記面積（㎡）", "抵充面積（㎡）", "權屬", "管理機關", "代管機關", "備註"]
A4_FOOT = ["合計道路、溝渠、河川等土地抵充面積：", "未登記土地面積：", "總計抵充土地面積："]
A4_SIGN = ["日期：", "編造：", "校對：", "承辦人：", "科長："]
A14_SIGN = ["日期：", "承辦：", "股長：", "科長：", "局(處)長："]
A14_NOTE = "附註：本表未列事項詳原土地登記簿記載"
A14_CELLS = {
    (1, 1): "○○縣(市)第○○期○○○○市地重劃區他項權利清冊",
    (3, 1): "重劃前原載情形", (3, 15): "重劃後變更情形", (3, 21): "備註",
    (4, 1): "土地標示", (4, 7): "權利人", (4, 9): "主登記次序", (4, 10): "權利種類", (4, 11): "權利範圍",
    (4, 12): "權利價值", (4, 13): "存續期間", (4, 14): "其他登記事項", (4, 15): "土地標示", (4, 20): "協議內容",
    (5, 1): "鄉鎮市區", (5, 2): "段", (5, 3): "小段", (5, 4): "地號", (5, 5): "面積（㎡）", (5, 6): "持分比率",
    (5, 7): "姓名", (5, 8): "住址", (5, 15): "段", (5, 16): "小段", (5, 17): "地號", (5, 18): "面積（㎡）",
    (5, 19): "持分比率",
}
A14_MERGES = {"A1:U1", "A3:N3", "O3:T3", "U3:U5", "A4:F4", "G4:H4", "I4:I5", "J4:J5", "K4:K5", "L4:L5",
              "M4:M5", "N4:N5", "O4:S4", "T4:T5"}
A15_CELLS = {
    (1, 1): "○○縣(市)第○○期○○○○市地重劃區限制登記清冊",
    (3, 1): "重劃前原載情形", (3, 13): "重劃後轉載情形", (3, 23): "備註",
    (4, 1): "土地標示", (4, 7): "權利人", (4, 9): "登記種類", (4, 10): "權利範圍", (4, 11): "內容",
    (4, 12): "限制登記文號及囑託機關", (4, 13): "土地標示", (4, 18): "權利人", (4, 20): "登記種類",
    (4, 21): "權利範圍", (4, 22): "協議內容",
    (5, 1): "鄉鎮市區", (5, 2): "段", (5, 3): "小段", (5, 4): "地號", (5, 5): "面積（㎡）", (5, 6): "持分比率",
    (5, 7): "姓名", (5, 8): "國民身分證統一編號", (5, 13): "段", (5, 14): "小段", (5, 15): "地號",
    (5, 16): "面積（㎡）", (5, 17): "持分比率", (5, 18): "姓名", (5, 19): "國民身分證統一編號",
}
A15_MERGES = {"A1:W1", "A3:L3", "M3:V3", "W3:W5", "A4:F4", "G4:H4", "I4:I5", "J4:J5", "K4:K5", "L4:L5",
              "M4:Q4", "R4:S4", "T4:T5", "U4:U5", "V4:V5"}
OWNER_BLOCK = ["土地所有權人姓名", "住址", "國民身分證統一編號", "分配清冊　第　冊　第　頁　號"]
A18_HEAD3 = {1: "區段編號", 2: "查估之重劃前每平方公尺單價", 3: "土地坐落",
             7: "地價及標準地價評議委員會修正意見、具體事實及理由", 8: "評議結果（元/平方公尺）", 9: "備註"}
A18_HEAD4 = {3: "鄉鎮市區", 4: "段", 5: "小段", 6: "地號"}
A19_HEAD = ["區段編號", "街路或地段名稱", "起訖點或範圍", "查估之路線價或區段價（元/平方公尺）",
            "地價及標準地價評議委員會修正意見、具體事實及理由", "評議結果（元/平方公尺）", "備註"]
A21_HEAD4 = {1: "街廓編號", 2: "都市計畫使用分區", 3: "規劃測量人員計算之面積", 6: "單價 × 面積", 7: "總價"}
A21_HEAD5 = {3: "編號", 4: "計算式", 5: "平方公尺"}
A21_SIGN = ["製表日期:　年　月　日", "承辦:", "股長:", "科長:"]
A25_COLHEAD = ["0.0", "1", "2", "3", "4", "5", "6", "7", "8", "9"]
REVIEW_DATE = "評議日期：　　年　　月　　日"


def _blank(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", "null"))


def rhu(x, nd):
    q = Decimal(1).scaleb(-nd)
    return Decimal(repr(float(x))).quantize(q, rounding=ROUND_HALF_UP)


def landno_key(s):
    p = str(s).split("-")
    return (int(p[0]), int(p[1]) if len(p) > 1 else 0)


def _ws(out, n):
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(out, FILES[n]))
    if len(wb.worksheets) != 1:
        raise AssertionError(f"附表 {n} 之工作表數 {len(wb.worksheets)}≠1")
    return wb.worksheets[0]


def _merges(ws):
    return {str(m) for m in ws.merged_cells.ranges}


def _last(ws):
    last = 0
    for row in ws.iter_rows():
        for c in row:
            if not _blank(c.value):
                last = max(last, c.row)
    return last


def _row_vals(ws, r, ncol):
    return [ws.cell(r, c).value for c in range(1, ncol + 1)]


def _nonblank(ws, r):
    return [v for v in _row_vals(ws, r, ws.max_column) if not _blank(v)]


def _cells(ws, want):
    bad = [(rc, ws.cell(*rc).value, v) for rc, v in want.items() if ws.cell(*rc).value != v]
    return bad


# ---------------------------------------------------------------- 獨立之讀
def src_lots():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX_SRC, read_only=True)
    rows = list(wb["重劃區地籍"].iter_rows(values_only=True))
    h = rows[0]
    it, isec, il = h.index("鄉鎮市"), h.index("地段"), h.index("地號")
    lots = {str(r[il]).strip(): (r[it], r[isec]) for r in rows[1:] if not _blank(r[il])}
    head = None
    rights = 0
    for r in wb["U_LAND"].iter_rows(values_only=True):
        if head is None:
            head = r
            ik = [i for i, x in enumerate(head) if x == "權利種類"]
            if len(ik) != 1:
                raise AssertionError(f"U_LAND 之「權利種類」欄 {len(ik)} 個")
            ik = ik[0]
            continue
        if r[0] != SECTION_NAME or _blank(r[1]):
            continue
        m, s = divmod(int(r[1]), 10000)
        ln = f"{m}" if s == 0 else f"{m}-{s}"
        if ln in lots and not _blank(r[ik]):
            rights += 1
    return lots, rights


def src_case():
    d = json.load(open(CASE, encoding="utf-8"))
    return d["財務接線_v3"], d["global"]["分區"]


# ---------------------------------------------------------------- 量測
def p0(out):
    import openpyxl
    for n, f in FILES.items():
        p = os.path.join(out, f)
        if not os.path.isfile(p) or os.path.getsize(p) == 0:
            return False, f"缺或空：{f}"
        wb = openpyxl.load_workbook(p)
        if _last(wb.worksheets[0]) == 0:
            return False, f"無內容：{f}"
    return True, "七檔皆在、可讀、非空"


def p1(out, v):
    ws = _ws(out, 4)
    if ws.cell(1, 1).value != "○○縣(市)第○○期○○市地重劃區抵充土地清冊" or "A1:J1" not in _merges(ws):
        return False, f"題：{ws.cell(1, 1).value!r}"
    if _row_vals(ws, 2, 10) != A4_HEAD:
        return False, f"表頭：{_row_vals(ws, 2, 10)}"
    want_area = v["抵充地_m2"]
    if want_area != 0:
        return False, f"case_params 之抵充地_m2＝{want_area}（本單唯 0）"
    for i, lab in enumerate(A4_FOOT):
        got = _nonblank(ws, 3 + i)
        if got != [lab, 0, "平方公尺"]:
            return False, f"列 {3 + i}：{got}"
    sign = _last(ws)
    if _nonblank(ws, sign) != A4_SIGN or sign != 7 or any(_nonblank(ws, r) for r in (6,)):
        return False, f"簽欄（列 {sign}）：{_nonblank(ws, sign)}"
    return True, "表頭 10 欄、無列、三合計 0 ㎡（case_params 抵充地_m2＝0）、簽欄列 7"


def _rights_form(ws, n, cells, merges, ncol):
    bad = _cells(ws, cells)
    if bad:
        return f"表頭格：{bad[:3]}"
    miss = merges - _merges(ws)
    if miss:
        return f"合併格缺：{sorted(miss)}"
    if _nonblank(ws, 2) != OWNER_BLOCK:
        return f"所有權人欄：{_nonblank(ws, 2)}"
    if any(_nonblank(ws, 6)):
        return f"列 6 有值：{_nonblank(ws, 6)}"
    if _nonblank(ws, 7) != [A14_NOTE] or _nonblank(ws, 8) != A14_SIGN or _last(ws) != 8:
        return f"附註／簽欄：{_nonblank(ws, 7)} {_nonblank(ws, 8)} 末列 {_last(ws)}"
    if ws.max_column < ncol or _blank(ws.cell(3, ncol).value):
        return f"欄數不足 {ncol}"
    return None


def p2(out, rights):
    if rights != 0:
        return False, f"示範資料之現存他項權利 {rights} 列（本單唯 0 列之式）"
    ws = _ws(out, 14)
    e = _rights_form(ws, 14, A14_CELLS, A14_MERGES, 21)
    return (False, e) if e else (True, "表頭三層 21 欄合手冊、無列（現存他項權利 0）、附註與簽欄")


def p3(out):
    ws = _ws(out, 15)
    e = _rights_form(ws, 15, A15_CELLS, A15_MERGES, 23)
    return (False, e) if e else (True, "表頭三層 23 欄合手冊、無列（匯入資料無限制登記之欄）、附註與簽欄")


def p4(out, lots, v):
    ws = _ws(out, 18)
    if ws.cell(1, 1).value != "○○縣（市）第○○期○○○○市地重劃區重劃前地價評議表" or ws.cell(2, 1).value != REVIEW_DATE:
        return False, f"題／評議日期：{ws.cell(1, 1).value!r} {ws.cell(2, 1).value!r}"
    if _cells(ws, {(3, c): t for c, t in A18_HEAD3.items()}) or _cells(ws, {(4, c): t for c, t in A18_HEAD4.items()}):
        return False, "表頭"
    zone_of = v["原地號_區段"]
    unit = {z: d["單價_元每m2"] for z, d in v["重劃前區段_面積單價"].items()}
    price = {z: int(rhu(u, 0)) for z, u in unit.items()}
    for z, (pp, cnt) in ANCHOR_ZONE.items():
        if price.get(z) != pp or sum(1 for ln in lots if zone_of.get(ln) == z) != cnt:
            return False, f"錨：區段 {z} 之地價 {price.get(z)}／宗數"
    want = []
    for z in sorted(set(zone_of[ln] for ln in lots)):
        for ln in sorted([x for x in lots if zone_of[x] == z], key=landno_key):
            want.append((z, price[z], lots[ln][0], lots[ln][1], None, ln))
    merged = {}
    for m in ws.merged_cells.ranges:
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                merged[(r, c)] = ws.cell(m.min_row, m.min_col).value
    got = []
    for r in range(5, _last(ws) + 1):
        vals = [merged.get((r, c), ws.cell(r, c).value) for c in range(1, 10)]
        if any(not _blank(x) for x in vals[6:]):
            return False, f"列 {r} 之評議欄有值：{vals[6:]}"
        got.append(tuple(None if _blank(x) else x for x in vals[:6]))
    if got != want:
        diff = [(i, g, w) for i, (g, w) in enumerate(zip(got, want)) if g != w][:2]
        return False, f"列數 {len(got)}／{len(want)}；差 {diff}"
    for z in ANCHOR_ZONE:
        rs = [5 + i for i, w in enumerate(want) if w[0] == z]
        rng = f"A{rs[0]}:A{rs[-1]}"
        if len(rs) > 1 and rng not in _merges(ws):
            return False, f"區段 {z} 之格未合併（{rng}）"
    return True, f"{len(got)} 列＝{len(lots)} 宗、序（區段→地號）、區段地價 a 43909×5、b 37323×48、評議欄空"


def _post(v):
    b = v["後街廓_面積單價"]
    keys = sorted(b, key=lambda k: (int(k[1:]), k))
    return [(k, Decimal(repr(float(b[k]["面積_m2"]))), float(b[k]["單價_元每m2"])) for k in keys]


def p5(out, v):
    ws = _ws(out, 19)
    if ws.cell(1, 1).value != "○○縣（市）第○○期○○○○市地重劃區重劃後地價評議表" or ws.cell(2, 1).value != REVIEW_DATE:
        return False, "題／評議日期"
    if _row_vals(ws, 3, 7) != A19_HEAD:
        return False, f"表頭：{_row_vals(ws, 3, 7)}"
    post = _post(v)
    if _last(ws) != 3 + len(post):
        return False, f"末列 {_last(ws)}≠{3 + len(post)}"
    for i, (k, a, u) in enumerate(post):
        r = 4 + i
        row = _row_vals(ws, r, 7)
        if row[0] != k or not _blank(row[1]) or row[2] != f"街廓 {k}" or any(not _blank(x) for x in row[4:]):
            return False, f"列 {r}：{row}"
        if not isinstance(row[3], (int, float)) or abs(row[3] - u) > 1e-6:
            return False, f"列 {r} 單價 {row[3]}≠{u}"
        if ws.cell(r, 4).number_format != "#,##0.00":
            return False, f"列 {r} 單價之格式 {ws.cell(r, 4).number_format}"
    return True, f"{len(post)} 列 R1〜R{len(post)}、單價全精度（|差|≤1e-6）、顯示二位"


def p6(out, v, zoning):
    ws = _ws(out, 21)
    if ws.cell(1, 1).value != "○○縣(市)第○○期○○市地重劃區" or ws.cell(2, 1).value != "重劃後可分配土地街廓面積地價計算表":
        return False, "題"
    if ws.cell(3, 5).value != "第○可分配土地街廓面積地價圖":
        return False, f"圖號格：{ws.cell(3, 5).value!r}"
    if _cells(ws, {(4, c): t for c, t in A21_HEAD4.items()}) or _cells(ws, {(5, c): t for c, t in A21_HEAD5.items()}):
        return False, "表頭"
    post = _post(v)
    sa, st = Decimal(0), Decimal(0)
    for i, (k, a, u) in enumerate(post):
        r = 6 + i
        row = _row_vals(ws, r, 7)
        tot = rhu(float(a) * u, 2)
        want = [k, zoning, None, None, float(a), f"{u:,.2f} × {float(a):,.2f}", float(tot)]
        if row[:4] != want[:4] or row[5] != want[5] or abs(row[4] - want[4]) > 1e-9 or abs(row[6] - want[6]) > 1e-6:
            return False, f"列 {r}：{row} ≠ {want}"
        sa += Decimal(repr(row[4]))
        st += rhu(row[6], 2)
    if sa != ANCHOR_SUM_AREA or st != ANCHOR_SUM_TOTAL:
        return False, f"Σ面積 {sa}／Σ總價 {st} ≠ 錨 {ANCHOR_SUM_AREA}／{ANCHOR_SUM_TOTAL}"
    sign = 6 + len(post) + 1
    if _nonblank(ws, sign) != A21_SIGN or _last(ws) != sign:
        return False, f"簽欄（列 {sign}）：{_nonblank(ws, sign)}"
    return True, f"{len(post)} 列；Σ面積 {sa}、Σ總價 {st}（＝KL 之錨）；簽欄列 {sign}"


def p7(out):
    ws = _ws(out, 25)
    if ws.cell(1, 1).value != "街角地側面道路負擔百分率累積表":
        return False, "題"
    if [ws.cell(3, c).value for c in range(2, 12)] != A25_COLHEAD or ws.cell(2, 2).value != "公寸" \
            or ws.cell(3, 1).value != "RW" or ws.cell(2, 12).value != "備註" or ws.cell(4, 12).value != "公分以下用內插法。":
        return False, "表頭"
    vals = {}
    for r in range(4, 22):
        w = ws.cell(r, 1).value
        if w != r - 3:
            return False, f"列 {r} 之公尺 {w}"
        for j in range(10):
            x = ws.cell(r, 2 + j).value
            if w == 18 and j > 0:
                if not _blank(x):
                    return False, f"18 公尺列之第 {j} 格有值 {x}"
                continue
            if not isinstance(x, (int, float)):
                return False, f"({w},{j}) 非數：{x!r}"
            vals[(w, j)] = float(x)
    s = ";".join(f"{w}.{j}:{vals[(w, j)]:.2f}" for w, j in sorted(vals))
    if len(vals) != 171 or hashlib.sha256(s.encode()).hexdigest() != RW_SHA:
        bad = [(k, vals.get(k), x) for k, x in RW_SPOT.items() if vals.get(k) != x]
        return False, f"171 格之 sha256 不合手冊（格數 {len(vals)}；錨差 {bad}）"
    seq = [vals[k] for k in sorted(vals)]
    if any(b <= a for a, b in zip(seq, seq[1:])):
        return False, "非嚴格遞增"
    if ws.cell(22, 1).value != "註：本表係由前面負擔百分率簡表推演而來。" or _last(ws) != 22:
        return False, f"註／末列 {_last(ws)}"
    return True, "171 格合手冊（sha256 " + RW_SHA[:8] + "…）、嚴格遞增、表頭與註"


def run(out, src=None, quiet=False):
    """⇒ 不過之項名之列。src ＝ (lots, rights, v, zoning)（省略者即讀之）。"""
    if src is None:
        lots, rights = src_lots()
        v, zoning = src_case()
    else:
        lots, rights, v, zoning = src
    tests = [
        ("P0", lambda: p0(out)),
        ("P1", lambda: p1(out, v)),
        ("P2", lambda: p2(out, rights)),
        ("P3", lambda: p3(out)),
        ("P4", lambda: p4(out, lots, v)),
        ("P5", lambda: p5(out, v)),
        ("P6", lambda: p6(out, v, zoning)),
        ("P7", lambda: p7(out)),
    ]
    fails = []
    for name, fn in tests:
        try:
            ok, msg = fn()
        except Exception as e:  # noqa: BLE001  量測之例外即不過
            ok, msg = False, f"{type(e).__name__}: {e}"
        if not quiet:
            print(f"{'PASS' if ok else 'FAIL'} {name} {msg}")
        if not ok:
            fails.append(name)
    if not quiet:
        print(f"{len(tests) - len(fails)}/{len(tests)}" + (f"；不過：{fails}" if fails else ""))
    return fails


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 1:
        print(__doc__)
        return 2
    return 1 if run(argv[0]) else 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
