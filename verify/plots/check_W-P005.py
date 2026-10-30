# -*- coding: utf-8 -*-
"""
W-P.005：重劃後三圖層與地號對照表之量測。
獨立於 plots/：⛔ import 之（V7 唯以子行程呼叫 plots/post_maps.py 之命令列）；期望取自 配地結果_退縮*.json
（plots/alloc_state.py 之輸出·即配地流程之 g_rows）、重劃前圖.dxf（W-P.001／002 定版之三圖層）、
verify/case_params_UC9898.json、KL 2026-10-28 23:31、23:37 之答。

命令列：python verify/plots/check_W-P005.py <成果目錄> <重劃前圖.dxf> <dict-revised.json> <退縮…>
  成果目錄 ＝ plots/alloc_state.py 與 plots/post_maps.py 之 --out（同一目錄）。
任一不過 ⇒ 結束碼 1。
V0 各檔皆在、可讀；V1 圖層（恰六＋系統二）、重劃前三圖層之圖元與重劃前圖逐一同；
V2 公告分配地籍線之宗（數、形、面積、各街廓密接而無重疊、略之列皆 0）；
V3 地號與註記（每宗恰一地號、無重號、＝ 對照表之正式地號；面積、寬、深之值）；
V4 S 型之預設編定（異法重導：形心·區·街廓之序·交替·公設續編於後；1〜N 連號）；
V5 分配後歸戶圖（原地主之地 ＝ 其重劃前之歸戶色與歸戶號；抵費地一色並註「抵費地」；公設不填色、只畫界線）；
V6 重劃後地價圖（街廓之價·二位；一價一色；公設不填色）；
V7 對照表之用（改號即帶入圖；重號、空白 ⇒ 停止而無輸出）與顏色對照表；
V8 街角第 1 宗（W-P.005_補令一·KL 2026-10-30 08:00）：每一側面道路線恰一宗；其 W 為截角前之寬（異法重導）；
   道路截角（異法：該宗與理論角之凸包減該宗、取含理論角之一塊）面積 > 0、落於道路、⛔ 與任一宗重疊。
V3 之寬 ＝ S（G 值計算之臨正面道路寬度）、深 ＝ 面積 ÷ W，街角第 1 宗之深 ＝（面積 ＋ 道路截角面積）÷ W
   （截角以 V8 之異法算之）。
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from decimal import Decimal, ROUND_HALF_UP

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CASE = os.path.join(REPO, "verify", "case_params_UC9898.json")
POST_MAPS = os.path.join(REPO, "plots", "post_maps.py")
PRE_LAYERS = ("重劃前地籍線", "重劃前歸戶圖", "重劃前地價圖")
L_CAD, L_OWN, L_PRICE = "公告分配地籍線", "分配後歸戶圖", "重劃後地價圖"
SYS_LAYERS = {"0", "Defpoints"}
OFFSET_WORD = "抵費地"
MAP_SHEET = "地號對照表"
MAP_HEAD = ("暫編", "街廓", "類別", "配地之暫編地號", "預設地號（S 型）", "正式地號", "面積（㎡）")
NO_RE = re.compile(r"^\d+(-\d+)?$")
TOL_SHAPE = 1e-4      # 多邊形之對稱差（㎡）
TOL_TILE = 0.05       # 一街廓之宗之和與街廓面積（㎡）
FMT = {"dxf": "公告圖_退縮{}m.dxf", "map": "地號對照表_退縮{}m.xlsx",
       "color": "顏色對照表_重劃後_退縮{}m.xlsx", "state": "配地結果_退縮{}m.json"}


def rhu(x, nd=2):
    return Decimal(repr(float(x))).quantize(Decimal(1).scaleb(-nd), rounding=ROUND_HALF_UP)


def s2(x):
    return f"{rhu(x):.2f}"


def nkey(s):
    a = str(s).split("-")
    return (int(a[0]), int(a[1]) if len(a) > 1 else 0)


class Ctx:
    pass


def load(out, pre_dxf, tag):
    import ezdxf
    import openpyxl
    from shapely.geometry import Polygon
    c = Ctx()
    c.tag = tag
    c.p = {k: os.path.join(out, v.format(tag)) for k, v in FMT.items()}
    c.state = json.load(open(c.p["state"], encoding="utf-8"))
    c.doc = ezdxf.readfile(c.p["dxf"])
    c.pre = ezdxf.readfile(pre_dxf)
    ws = openpyxl.load_workbook(c.p["map"], data_only=True)[MAP_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    c.map_head = tuple(rows[0][:len(MAP_HEAD)])
    c.map = [r for r in rows[1:] if r and r[0] is not None and r[1] is not None]
    c.color = openpyxl.load_workbook(c.p["color"], data_only=True)
    c.msp = c.doc.modelspace()

    def polys(msp, layer):
        return [(Polygon([(p[0], p[1]) for p in e.get_points("xy")]), e)
                for e in msp.query(f'LWPOLYLINE[layer=="{layer}"]')]

    def texts(msp, layer):
        from shapely.geometry import Point
        return [(Point(e.dxf.align_point[0], e.dxf.align_point[1]), e.dxf.text)
                for e in msp.query(f'TEXT[layer=="{layer}"]')]

    def hatches(msp, layer):
        out_ = []
        for e in msp.query(f'HATCH[layer=="{layer}"]'):
            if len(e.paths) != 1:
                raise AssertionError("HATCH 之邊界非一")
            out_.append((Polygon([(v[0], v[1]) for v in e.paths[0].vertices]), tuple(e.rgb)))
        return out_
    c.polys, c.texts, c.hatches = polys, texts, hatches
    blocks = c.state["街廓"]
    c.alloc = {k for k, v in blocks.items() if "宗地分配線方向" in v}
    c.public = sorted(k for k in blocks if k not in c.alloc)
    c.lots = c.state["宗"]
    c.post_price = {b: d["單價_元每m2"] for b, d in
                    json.load(open(CASE, encoding="utf-8"))["財務接線_v3"]["後街廓_面積單價"].items()}
    return c


def _sig(e):
    t = e.dxftype()
    if t == "LWPOLYLINE":
        return (t, e.dxf.layer, tuple((round(p[0], 6), round(p[1], 6)) for p in e.get_points("xy")))
    if t == "TEXT":
        return (t, e.dxf.layer, e.dxf.text, round(e.dxf.align_point[0], 6), round(e.dxf.align_point[1], 6))
    if t == "HATCH":
        return (t, e.dxf.layer, tuple(e.rgb),
                tuple((round(v[0], 6), round(v[1], 6)) for v in e.paths[0].vertices))
    return (t, e.dxf.layer)


def _match(c):
    """配地結果之每宗（與公設之區塊）⇒ 公告分配地籍線之多邊形（對稱差 ≤ TOL_SHAPE）。"""
    from shapely.geometry import Polygon
    ps = c.polys(c.msp, L_CAD)
    want = [(x["配地暫編"], x, Polygon(x["多邊形"])) for x in c.lots]
    want += [(b, None, Polygon(c.state["街廓"][b]["多邊形"])) for b in c.public]
    got = {}
    used = set()
    for key, x, wp in want:
        hit = [i for i, (p, _) in enumerate(ps) if p.symmetric_difference(wp).area <= TOL_SHAPE]
        if len(hit) != 1 or hit[0] in used:
            return None, f"{key} 於 {L_CAD} 之多邊形 {len(hit)} 個"
        used.add(hit[0])
        got[key] = (x, ps[hit[0]][0])
    if len(used) != len(ps):
        return None, f"{L_CAD} 有多出之多邊形 {len(ps) - len(used)} 個"
    return got, None


def v0(c):
    return []


def v1(c):
    bad = []
    names = {ly.dxf.name for ly in c.doc.layers}
    want = SYS_LAYERS | set(PRE_LAYERS) | {L_CAD, L_OWN, L_PRICE}
    if names != want:
        bad.append(f"圖層 {sorted(names ^ want)}")
    for ly in PRE_LAYERS:
        a = sorted(_sig(e) for e in c.msp if e.dxf.layer == ly)
        b = sorted(_sig(e) for e in c.pre.modelspace() if e.dxf.layer == ly)
        if a != b or not a:
            bad.append(f"{ly} 之圖元與重劃前圖不同（{len(a)}／{len(b)}）")
    for e in c.msp:
        if e.dxf.layer in SYS_LAYERS:
            bad.append(f"系統圖層 {e.dxf.layer} 上有圖元")
            break
    return bad


def v2(c):
    from shapely.ops import unary_union
    from shapely.geometry import Polygon
    bad = []
    got, err = _match(c)
    if err:
        return [err]
    n = len(c.polys(c.msp, L_CAD))
    if n != len(c.lots) + len(c.public):
        bad.append(f"宗數 {n} ≠ {len(c.lots)}＋{len(c.public)}")
    for x in c.lots:
        if s2(Polygon(x["多邊形"]).area) != s2(x["面積"]):
            bad.append(f"{x['配地暫編']} 之面積 {x['面積']} ≠ 圖 {Polygon(x['多邊形']).area:.4f}")
    for b in sorted(c.alloc):
        ps = [got[x["配地暫編"]][1] for x in c.lots if x["街廓"] == b]
        bp = Polygon(c.state["街廓"][b]["多邊形"])
        s = sum(p.area for p in ps)
        u = unary_union(ps)
        if abs(s - bp.area) > TOL_TILE or u.symmetric_difference(bp).area > TOL_TILE:
            bad.append(f"{b}：宗之和 {s:.3f}、聯集之差 {u.symmetric_difference(bp).area:.4f}、街廓 {bp.area:.3f}")
    for d in c.state["略"]:
        if s2(d["幾何面積"]) != "0.00":
            bad.append(f"略之列 {d['配地暫編']} 之面積非 0")
    return bad


def _map_rows(c):
    if c.map_head != MAP_HEAD:
        raise AssertionError(f"對照表之表頭 {c.map_head}")
    return {r[3] if r[3] is not None else r[0]: r for r in c.map}


def v3(c):
    bad = []
    got, err = _match(c)
    if err:
        return [err]
    try:
        mr = _map_rows(c)
    except AssertionError as ex:
        return [str(ex)]
    if len(mr) != len(got):
        bad.append(f"對照表 {len(mr)} 列 ≠ {len(got)} 宗")
    tx = c.texts(c.msp, L_CAD)
    seen = {}
    for key, (x, poly) in got.items():
        inside = [s for pt, s in tx if poly.contains(pt)]
        nos = [s for s in inside if NO_RE.match(s)]
        if len(nos) != 1:
            bad.append(f"{key} 之地號之註記 {nos}")
            continue
        no = nos[0]
        if no in seen:
            bad.append(f"地號 {no} 重號：{seen[no]}、{key}")
        seen[no] = key
        r = mr.get(key)
        if r is None or str(r[5]) != no:
            bad.append(f"{key} 之地號 {no} ≠ 對照表 {None if r is None else r[5]}")
        area = x["面積"] if x else rhu(poly.area)
        want = [no, s2(area)]
        if x and x["類別"] == "分配地":
            w = rhu(x["宗地寬度"])
            a = Decimal(repr(float(x["面積"])))
            if "街角第1宗側別" in x:
                a += rhu(_cut_alt(c, x).area)
            d = (a / w).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            want += [f"寬{rhu(x['S']):.2f}", f"深{d:.2f}"]
        if sorted(inside) != sorted(want):
            bad.append(f"{key} 之註記 {inside} ≠ {want}")
    if len(tx) != sum(4 if (x and x["類別"] == "分配地") else 2 for x, _ in got.values()):
        bad.append(f"{L_CAD} 之文字數不合（{len(tx)}）")
    return bad


def _theory_corner(c, x):
    """異法：以向量之叉積解正面臨路線與側面道路線（皆延長）之交點。"""
    b = c.state["街廓"][x["街廓"]]
    (ax, ay), (bx, by) = b["正面臨路線"]
    (px, py), (qx, qy) = b["側面道路線"][x["街角第1宗側別"]]
    rx, ry, sx, sy = bx - ax, by - ay, qx - px, qy - py
    den = rx * sy - ry * sx
    t = ((px - ax) * sy - (py - ay) * sx) / den
    return ax + t * rx, ay + t * ry


def _cut_alt(c, x):
    """異法之道路截角：（該宗 ∪ 理論角）之凸包 − 該宗，取與理論角相接之一塊。"""
    from shapely.geometry import MultiPoint, Point, Polygon
    lot = Polygon(x["多邊形"])
    cp = Point(_theory_corner(c, x))
    rest = MultiPoint(list(lot.exterior.coords) + [cp.coords[0]]).convex_hull.difference(lot)
    parts = [g for g in getattr(rest, "geoms", [rest]) if g.area > 1e-9 and g.distance(cp) < 1e-6]
    if len(parts) != 1:
        raise AssertionError(f"{x['配地暫編']} 之道路截角 {len(parts)} 塊")
    return parts[0]


def v8(c):
    import math
    from shapely.geometry import LineString, Point, Polygon
    from shapely.ops import unary_union
    bad = []
    blocks = c.state["街廓"]
    lots = [x for x in c.lots if x["類別"] == "分配地"]
    roads = unary_union([Polygon(blocks[b]["多邊形"]) for b in c.public if blocks[b]["分類"] != "公園"]
                        or [Polygon()])
    every = unary_union([Polygon(x["多邊形"]) for x in c.lots])
    n_c = 0
    for b in sorted(c.alloc):
        for side in sorted(blocks[b].get("側面道路線", {})):
            hit = [x for x in lots if x["街廓"] == b and x.get("街角第1宗側別") == side]
            sl = LineString(blocks[b]["側面道路線"][side])
            touch = [x for x in lots if x["街廓"] == b and Polygon(x["多邊形"]).intersection(sl.buffer(0.02)).area > 0.01]
            if len(hit) != 1:
                bad.append(f"{b} {side}側 之街角第 1 宗 {len(hit)} 宗")
                continue
            x = hit[0]
            if [y["配地暫編"] for y in touch] != [x["配地暫編"]]:
                bad.append(f"{b} {side}側 臨側面道路線之宗 {[y['配地暫編'] for y in touch]} ≠ {x['配地暫編']}")
            n_c += 1
            (ax, ay), (bx, by) = blocks[b]["正面臨路線"]
            L = math.hypot(bx - ax, by - ay)
            ux, uy = (bx - ax) / L, (by - ay) / L
            dx, dy = blocks[b]["宗地分配線方向"]
            dl = math.hypot(dx, dy)
            cosv = abs(ux * dx / dl + uy * dy / dl)
            sinv = math.sqrt(max(0.0, 1 - cosv * cosv))   # 正面臨路線與「垂直於宗地分配線」之夾角之餘弦
            cx, cy = _theory_corner(c, x)
            fl = LineString(blocks[b]["正面臨路線"])
            ts = [abs((px - cx) * ux + (py - cy) * uy) for px, py in Polygon(x["多邊形"]).exterior.coords
                  if fl.distance(Point(px, py)) <= 0.02]
            w_pre, w_post = max(ts) * sinv, (max(ts) - min(ts)) * sinv
            if abs(rhu(w_pre) - rhu(x["宗地寬度"])) > Decimal("0.01"):
                bad.append(f"{x['配地暫編']} 之 W {x['宗地寬度']} ≠ 截角前之寬 {w_pre:.3f}（截角後 {w_post:.3f}）")
            if not w_post < w_pre - 1:
                bad.append(f"{x['配地暫編']} 之截角前後之寬不分（{w_pre:.3f}／{w_post:.3f}）")
            try:
                cut = _cut_alt(c, x)
            except AssertionError as ex:
                bad.append(str(ex))
                continue
            if cut.area <= 0.01:
                bad.append(f"{x['配地暫編']} 之道路截角之面積 {cut.area:.4f}")
            if cut.intersection(roads).area < 0.99 * cut.area:
                bad.append(f"{x['配地暫編']} 之道路截角不在道路（{cut.intersection(roads).area:.4f}／{cut.area:.4f}）")
            if cut.intersection(every).area > 1e-4:
                bad.append(f"{x['配地暫編']} 之道路截角與宗重疊")
    if n_c == 0:
        bad.append("無街角第 1 宗")
    if any("街角第1宗側別" in x and x.get("街角第1宗側別") not in blocks[x["街廓"]].get("側面道路線", {})
           for x in lots):
        bad.append("街角第 1 宗之側別無側面道路線")
    return bad


def _expected_order(c):
    """異法：形心；街廓之序 ＝ 形心之 y 降、x 降；第奇個街廓自東而西（形心 x 降）、第偶個自西而東；公設續後。"""
    from shapely.geometry import Polygon
    from shapely.ops import unary_union
    blocks = c.state["街廓"]
    P = {k: Polygon(v["多邊形"]) for k, v in blocks.items()}
    u = unary_union([p.buffer(0.01) for p in P.values()])
    regs = list(getattr(u, "geoms", [u]))
    regs.sort(key=lambda g: (-g.centroid.y, -g.centroid.x))
    seq, pubs = [], []
    for g in regs:
        mem = [k for k in P if g.contains(P[k].centroid) or g.contains(P[k].representative_point())]
        al = sorted([k for k in mem if k in c.alloc], key=lambda k: (-P[k].centroid.y, -P[k].centroid.x))
        for i, k in enumerate(al):
            lots = [x for x in c.lots if x["街廓"] == k]
            lots.sort(key=lambda x: Polygon(x["多邊形"]).centroid.x, reverse=(i % 2 == 0))
            seq += [x["配地暫編"] for x in lots]
        pubs += sorted([k for k in mem if k not in c.alloc], key=lambda k: (-P[k].centroid.y, -P[k].centroid.x))
    return seq + pubs


def v4(c):
    bad = []
    try:
        mr = _map_rows(c)
    except AssertionError as ex:
        return [str(ex)]
    exp = _expected_order(c)
    dflt = {k: r[4] for k, r in mr.items()}
    want = {k: str(i) for i, k in enumerate(exp, start=1)}
    if {k: str(v) for k, v in dflt.items()} != want:
        diff = [(k, dflt.get(k), want.get(k)) for k in exp if str(dflt.get(k)) != want[k]][:6]
        bad.append(f"預設地號 ≠ S 型之序：{diff}")
    got, err = _match(c)
    if err:
        return bad + [err]
    tx = c.texts(c.msp, L_CAD)
    on_map = {}
    for key, (_x, poly) in got.items():
        nos = [s for pt, s in tx if poly.contains(pt) and NO_RE.match(s)]
        on_map[key] = nos[0] if len(nos) == 1 else None
    if c.with_default and on_map != want:
        bad.append("圖上之地號 ≠ S 型之預設（未給對照表者）")
    return bad


def v5(c):
    from shapely.geometry import Polygon
    bad = []
    got, err = _match(c)
    if err:
        return [err]
    pre = c.pre.modelspace()
    pre_poly = {}
    for p, _e in c.polys(pre, PRE_LAYERS[0]):
        labs = [s for pt, s in c.texts(pre, PRE_LAYERS[0]) if p.contains(pt)]
        if len(labs) == 1:
            pre_poly[labs[0]] = p
    pre_h = c.hatches(pre, PRE_LAYERS[1])
    pre_t = c.texts(pre, PRE_LAYERS[1])
    hh = c.hatches(c.msp, L_OWN)
    ht = c.texts(c.msp, L_OWN)
    outl = c.polys(c.msp, L_OWN)
    owner_colors = set()
    offs = set()
    n_h = 0
    for key, (x, poly) in got.items():
        mine = [col for hp, col in hh if hp.symmetric_difference(poly).area <= TOL_SHAPE]
        labs = [s for pt, s in ht if poly.contains(pt)]
        if x is None:
            if mine or labs:
                bad.append(f"公設 {key} 有填色或註記")
            if len([1 for p, _ in outl if p.symmetric_difference(poly).area <= TOL_SHAPE]) != 1:
                bad.append(f"公設 {key} 於 {L_OWN} 無界線")
            continue
        n_h += 1
        if len(mine) != 1 or len(labs) != 1:
            bad.append(f"{key} 之填色 {len(mine)}、註記 {labs}")
            continue
        if x["類別"] == "抵費地":
            offs.add(mine[0])
            if labs[0] != OFFSET_WORD:
                bad.append(f"{key} 之註記 {labs[0]}")
            continue
        pp = pre_poly.get(x["原地號"])
        if pp is None:
            bad.append(f"{key} 之原地號 {x['原地號']} 不在重劃前圖")
            continue
        pc = [col for hp, col in pre_h if hp.symmetric_difference(pp).area <= TOL_SHAPE]
        pl = [s for pt, s in pre_t if pp.contains(pt)]
        if pc != [mine[0]] or pl != [labs[0]]:
            bad.append(f"{key}（原 {x['原地號']}）：色 {mine[0]}／歸戶號 {labs[0]} ≠ 重劃前 {pc}／{pl}")
        owner_colors.add(mine[0])
    if len(offs) > 1:
        bad.append(f"抵費地之色不一：{offs}")
    if offs & owner_colors or offs & {col for _, col in pre_h}:
        bad.append("抵費地之色與某戶之色同")
    if len(hh) != n_h or len(ht) != n_h:
        bad.append(f"{L_OWN} 之填色 {len(hh)}、註記 {len(ht)} ≠ {n_h}")
    # 相鄰之不同戶異色
    lots = [(x, poly) for x, poly in got.values() if x and x["類別"] == "分配地"]
    for i, (a, pa) in enumerate(lots):
        for b, pb in lots[i + 1:]:
            if a["歸戶"] != b["歸戶"] and pa.buffer(0.05).intersects(pb):
                ca = [col for hp, col in hh if hp.symmetric_difference(pa).area <= TOL_SHAPE]
                cb = [col for hp, col in hh if hp.symmetric_difference(pb).area <= TOL_SHAPE]
                if ca == cb:
                    bad.append(f"相鄰之異戶 {a['配地暫編']}、{b['配地暫編']} 同色")
    return bad


def v6(c):
    bad = []
    got, err = _match(c)
    if err:
        return [err]
    hh = c.hatches(c.msp, L_PRICE)
    ht = c.texts(c.msp, L_PRICE)
    col_of = {}
    n = 0
    for key, (x, poly) in got.items():
        mine = [col for hp, col in hh if hp.symmetric_difference(poly).area <= TOL_SHAPE]
        labs = [s for pt, s in ht if poly.contains(pt)]
        if x is None:
            if mine or labs:
                bad.append(f"公設 {key} 有填色或註記")
            continue
        n += 1
        want = s2(c.post_price[x["街廓"]])
        if len(mine) != 1 or labs != [want]:
            bad.append(f"{key} 之填色 {len(mine)}、註記 {labs} ≠ {want}")
            continue
        col_of.setdefault(want, set()).add(mine[0])
    if any(len(v) != 1 for v in col_of.values()):
        bad.append(f"同價異色：{col_of}")
    cols = [next(iter(v)) for v in col_of.values()]
    if len(set(cols)) != len(cols):
        bad.append("異價同色")
    if len(hh) != n or len(ht) != n:
        bad.append(f"{L_PRICE} 之填色 {len(hh)}、註記 {len(ht)} ≠ {n}")
    return bad


def v7(c):
    bad = []
    # 顏色對照表 ＝ 圖
    ws = c.color[L_OWN]
    rows = list(ws.iter_rows(values_only=True))[1:]
    got, err = _match(c)
    if err:
        return [err]
    hh = c.hatches(c.msp, L_OWN)
    ht = c.texts(c.msp, L_OWN)
    tx = c.texts(c.msp, L_CAD)
    tab = {}
    for _ly, no, col, nos in rows:
        for s in str(nos).split("、"):
            tab[s] = (str(no), col)
    for key, (x, poly) in got.items():
        if x is None:
            continue
        no = [s for pt, s in tx if poly.contains(pt) and NO_RE.match(s)][0]
        col = [cc for hp, cc in hh if hp.symmetric_difference(poly).area <= TOL_SHAPE][0]
        lab = [s for pt, s in ht if poly.contains(pt)][0]
        if tab.get(no) != (lab, "{},{},{}".format(*col)):
            bad.append(f"顏色對照表之 {no} {tab.get(no)} ≠ 圖 {lab}／{col}")
    if c.run_map:
        bad += _map_runs(c)
    return bad


def _map_runs(c):
    """以子行程跑 plots/post_maps.py：①二宗之正式地號互換 ⇒ 圖隨之；②重號、③空白 ⇒ 結束碼非 0 且無輸出。"""
    import openpyxl
    bad = []
    tmp = tempfile.mkdtemp(prefix="wp005_map_")
    try:
        base = [sys.executable, POST_MAPS, "--state", c.p["state"], "--pre-dxf", c.pre_path,
                "--moedict", c.moedict]
        lots = [r for r in c.map if r[2] == "分配地"]
        a, b = lots[0], lots[-1]

        def edited(name, fn):
            wb = openpyxl.load_workbook(c.p["map"])
            ws = wb[MAP_SHEET]
            for row in ws.iter_rows(min_row=2):
                fn(row)
            p = os.path.join(tmp, name)
            wb.save(p)
            return p

        def swap(row):
            if row[0].value == a[0]:
                row[5].value = b[5]
            elif row[0].value == b[0]:
                row[5].value = a[5]

        def dup(row):
            if row[0].value == a[0]:
                row[5].value = b[5]

        def blank(row):
            if row[0].value == a[0]:
                row[5].value = None
        for name, fn, ok_expected in (("swap.xlsx", swap, True), ("dup.xlsx", dup, False),
                                      ("blank.xlsx", blank, False)):
            o = os.path.join(tmp, name + ".out")
            r = subprocess.run(base + ["--out", o, "--map", edited(name, fn)], capture_output=True, text=True)
            made = os.path.exists(os.path.join(o, FMT["dxf"].format(c.tag)))
            if ok_expected:
                if r.returncode != 0 or not made:
                    bad.append(f"改號之對照表未能產出（{r.returncode}）：{r.stderr[-300:]}")
                    continue
                import ezdxf
                from shapely.geometry import Polygon
                d2 = ezdxf.readfile(os.path.join(o, FMT["dxf"].format(c.tag)))
                m2 = d2.modelspace()
                tx = c.texts(m2, L_CAD)
                for row, want in ((a, b[5]), (b, a[5])):
                    x = next(x for x in c.lots if x["配地暫編"] == row[3])
                    poly = Polygon(x["多邊形"])
                    nos = [s for pt, s in tx if poly.contains(pt) and NO_RE.match(s)]
                    if nos != [str(want)]:
                        bad.append(f"改號後 {row[0]} 之地號 {nos} ≠ {want}")
            elif r.returncode == 0 or made:
                bad.append(f"{name} 未停止（結束碼 {r.returncode}、輸出 {made}）")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    return bad


CHECKS = [("V0", v0), ("V1", v1), ("V2", v2), ("V3", v3), ("V4", v4), ("V5", v5), ("V6", v6), ("V7", v7), ("V8", v8)]


def run(out, pre_dxf, moedict, tags, quiet=False, run_map=True, with_default=True):
    fails = []
    for tag in tags:
        try:
            c = load(out, pre_dxf, tag)
        except Exception as ex:   # noqa: BLE001  （V0：缺檔、不可讀）
            if not quiet:
                print(f"FAIL V0 退縮{tag}m：{ex}")
            fails.append(f"V0@{tag}")
            continue
        c.pre_path, c.moedict = pre_dxf, moedict
        c.run_map, c.with_default = run_map, with_default
        for name, fn in CHECKS:
            try:
                bad = fn(c)
            except Exception as ex:   # noqa: BLE001
                bad = [f"{type(ex).__name__}: {ex}"]
            if not quiet:
                print(f"{'PASS' if not bad else 'FAIL'} {name} 退縮{tag}m" + ("" if not bad else "：" + "；".join(bad[:5])))
            if bad:
                fails.append(f"{name}@{tag}")
    return fails


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) < 4:
        print(__doc__)
        return 2
    out, pre_dxf, moedict, *tags = argv
    fails = run(out, pre_dxf, moedict, tags)
    print(f"{'全過' if not fails else '不過 ' + str(fails)}（{len(CHECKS) * len(tags) - len(fails)}/{len(CHECKS) * len(tags)}）")
    return 0 if not fails else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
