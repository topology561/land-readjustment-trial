# -*- coding: utf-8 -*-
"""
W-P.006_補令一 之預覽 PNG（補令 §三 6）：附表 13 之一戶與「縣有（抵費地）」、附表 24 之二張之截取。

命令列：python verify/plots/preview_W-P006a.py <成果目錄> <PNG 之輸出目錄> <退縮…>
輸出：W-P.006_補令一_附表13-24_退縮<t>m.png（每一退縮一檔）。
所取：附表 24 中「分配面積」已填之宗之所有權人最多者（同者地號小者）一張、「分配面積」暫空之宗之地號最小者一張；
附表 13 中前者所屬之戶之一表、及「縣有（抵費地）」之表。
預覽之標題唯供 KL 一眼辨表，⛔ 屬附表之內容；格之合併與框線以 Excel 開啟為準。
"""
import importlib.util
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location("preview_W_P006", os.path.join(HERE, "preview_W-P006.py"))
pv = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pv)
plt = pv.plt
OUT = "W-P.006_補令一_附表13-24_退縮{}m.png"
OFFSET_SHEET = "縣有（抵費地）"


def _alloc(ws):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == "分配面積(㎡)":
            return ws.cell(r, 2).value
    raise ValueError(f"{ws.title} 無「分配面積(㎡)」")


def _pick(wb24):
    filled, empty = None, None
    for ws in wb24:
        rep = str(ws.cell(2, 4).value or "")
        n = int(rep.rsplit("等", 1)[1][:-1]) if "等" in rep else 1
        key = tuple(int(p) for p in ws.title[2:].split("-"))
        if _alloc(ws) is not None:
            if filled is None or (-n, key) < filled[0]:
                filled = ((-n, key), ws.title)
        elif empty is None or key < empty[0]:
            empty = (key, ws.title)
    if filled is None or empty is None:
        raise ValueError("附表 24 無分配面積已填或暫空之宗")
    return filled[1], empty[1]


def _a13(ax, ws, t):
    last = max(r for r in range(1, ws.max_row + 1) if any(ws.cell(r, c).value is not None for c in range(1, 25)))
    cells = [[pv._fmt(ws.cell(r, c)) for c in range(1, 25)] for r in range(3, last + 1)]
    ax.axis("off")
    ax.set_title(f"附表 13　{ws.title}（退縮 {t} m）", fontsize=13, loc="left")
    tb = ax.table(cellText=cells, loc="upper left", cellLoc="center")
    tb.auto_set_font_size(False)
    tb.set_fontsize(6)
    tb.scale(1, 1.4)


def _a24(ax, ws, t):
    cells = [[pv._fmt(ws.cell(r, c)) for c in range(1, 5)] for r in range(2, ws.max_row + 1)
             if any(ws.cell(r, c).value is not None for c in range(1, 5))]
    ax.axis("off")
    ax.set_title(f"附表 24　{ws.title}（退縮 {t} m）", fontsize=13, loc="left")
    tb = ax.table(cellText=cells, loc="upper left", cellLoc="left", colWidths=[0.42, 0.34, 0.12, 0.12])
    tb.auto_set_font_size(False)
    tb.set_fontsize(7)
    tb.scale(1, 1.2)


def main(argv):
    import openpyxl
    out, png, tags = argv[0], argv[1], argv[2:]
    os.makedirs(png, exist_ok=True)
    for t in tags:
        t = f"{float(t):g}"
        wb24 = openpyxl.load_workbook(os.path.join(out, pv.A24.format(t)))
        wb13 = openpyxl.load_workbook(os.path.join(out, pv.A13.format(t)))
        n1, n2 = _pick(wb24)
        no = n1[2:]
        ws13 = next(ws for ws in wb13 if any(ws.cell(r, 15).value == no for r in range(8, ws.max_row + 1)))
        fig = plt.figure(figsize=(26, 30))
        _a13(fig.add_axes([0.01, 0.80, 0.98, 0.18]), ws13, t)
        _a13(fig.add_axes([0.01, 0.62, 0.98, 0.16]), wb13[OFFSET_SHEET], t)
        _a24(fig.add_axes([0.01, 0.01, 0.48, 0.58]), wb24[n1], t)
        _a24(fig.add_axes([0.51, 0.01, 0.48, 0.58]), wb24[n2], t)
        p = os.path.join(png, OUT.format(t))
        fig.savefig(p, dpi=80)
        plt.close(fig)
        print(f"輸出 {p}（{ws13.title}、{OFFSET_SHEET}、{n1}、{n2}）")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
