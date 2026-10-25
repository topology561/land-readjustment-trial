# -*- coding: utf-8 -*-
"""
W-P.003 之故意錯誤之測（單 §三 4）：成果須綠、Ⓐ〜Ⓖ 共 7 種變異體須各紅（量測 ＝ check_W-P003.py 之 8 項）。

命令列：python verify/plots/mutate_W-P003.py <成果目錄> <拋棄式之工作目錄>
  成果目錄 ＝ plots/appx_tables.py --out 之目錄（⛔ 改之；變異體皆寫於工作目錄）。
  Ⓑ、Ⓓ 掛換 plots/ 之函式或常數後重出（於本行程之記憶體中·⛔ 改檔）；他者改成果之複本。
結束碼：成果綠且 7 種皆紅 ⇒ 0；否則 1。
"""
import importlib.util
import os
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
_spec = importlib.util.spec_from_file_location("check_W_P003", os.path.join(HERE, "check_W-P003.py"))
chk = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(chk)

# 程式現行之 Rw（app.py rw_from_width：整數公尺之一位小數·線性內插）——與手冊之表差 ≤ 0.05
APP_RW = [0.0, 17.4, 24.4, 31.3, 37.8, 44.0, 50.0, 55.7, 61.1, 66.3,
          71.1, 75.7, 80.0, 84.0, 87.8, 91.3, 94.4, 97.4, 100.0]


def _copy(src, dst):
    if os.path.exists(dst):
        shutil.rmtree(dst)
    shutil.copytree(src, dst)


def _book(d, n):
    import openpyxl
    p = os.path.join(d, chk.FILES[n])
    return p, openpyxl.load_workbook(p)


def _plots():
    sys.path.insert(0, os.path.join(REPO, "plots"))
    import appx_tables
    return appx_tables


def _regen(d, obj, attr, val):
    orig = getattr(obj, attr)
    setattr(obj, attr, val)
    try:
        if os.path.exists(d):
            shutil.rmtree(d)
        _plots().main(["--out", d])
    finally:
        setattr(obj, attr, orig)


def mA_drop_lot(d):
    p, wb = _book(d, 18)
    ws = wb.worksheets[0]
    ws.delete_rows(chk._last(ws))
    wb.save(p)


def mB_unrounded_price(d):
    at = _plots()
    pm = at.PM
    orig = pm.load_price_zones

    def load(path=pm.CASE_PARAMS):
        z, u, _ = orig(path)
        return z, u, dict(u)
    _regen(d, pm, "load_price_zones", load)


def mC_rounded_unit_total(d):
    p, wb = _book(d, 21)
    ws = wb.worksheets[0]
    for r in range(6, 12):
        a = ws.cell(r, 5).value
        u = float(ws.cell(r, 6).value.split(" × ")[0].replace(",", ""))
        ws.cell(r, 7, float(chk.rhu(round(u) * a, 2)))
    wb.save(p)


def mD_app_rw(d):
    at = _plots()
    tab = {w: tuple(round(APP_RW[w] + (APP_RW[w + 1] - APP_RW[w]) * j / 10, 2) for j in range(10))
           for w in range(1, 18)}
    tab[18] = (100.0,)
    _regen(d, at, "RW_TABLE", tab)


def mE_swap_post(d):
    p, wb = _book(d, 19)
    ws = wb.worksheets[0]
    a, b = ws.cell(5, 4).value, ws.cell(6, 4).value
    ws.cell(5, 4, b)
    ws.cell(6, 4, a)
    wb.save(p)


def mF_drop_header(d):
    p, wb = _book(d, 14)
    ws = wb.worksheets[0]
    ws.cell(4, 12).value = None
    wb.save(p)


def mG_empty_file(d):
    open(os.path.join(d, chk.FILES[4]), "wb").close()


MUTANTS = [
    ("Ⓐ 附表 18 漏一宗", mA_drop_lot, True),
    ("Ⓑ 區段地價未取個位", mB_unrounded_price, False),
    ("Ⓒ 附表 21 總價以取整之單價計", mC_rounded_unit_total, True),
    ("Ⓓ 附表 25 用程式現行之 Rw", mD_app_rw, False),
    ("Ⓔ 附表 19 R2、R3 單價互換", mE_swap_post, True),
    ("Ⓕ 附表 14 漏「權利價值」", mF_drop_header, True),
    ("Ⓖ 附表 4 為空檔", mG_empty_file, True),
]


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 2:
        print(__doc__)
        return 2
    out, work = argv
    lots, rights = chk.src_lots()
    v, zoning = chk.src_case()
    src = (lots, rights, v, zoning)
    base = chk.run(out, src, quiet=True)
    print(f"成果：{'綠' if not base else '紅 ' + str(base)}")
    ok = not base
    for i, (name, fn, from_copy) in enumerate(MUTANTS):
        d = os.path.join(work, f"m{i}")
        if from_copy:
            _copy(out, d)
            fn(d)
        else:
            fn(d)
        fails = chk.run(d, src, quiet=True)
        print(f"{name}：{'紅 ' + str(fails) if fails else '綠（未捕）'}")
        ok = ok and bool(fails)
    print("合" if ok else "不合")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
