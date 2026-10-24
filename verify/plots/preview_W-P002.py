# -*- coding: utf-8 -*-
"""
W-P.002 之預覽 PNG（單 §三 6）：左「重劃前歸戶圖」（自輸出之 DXF 讀回而繪：填色＋歸戶號），右附表 26 全表。

命令列：python verify/plots/preview_W-P002.py <成果目錄> <PNG 之輸出目錄>
輸出：W-P.002_歸戶號與附表26.png
預覽之標題唯供 KL 一眼辨圖，⛔ 屬 DXF／附表之內容。
"""
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import font_manager  # noqa: E402

for _f in ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",):
    if os.path.exists(_f):
        font_manager.fontManager.addfont(_f)
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=_f).get_name()

L_CAD, L_OWN = "重劃前地籍線", "重劃前歸戶圖"
OUT_NAME = "W-P.002_歸戶號與附表26.png"


def draw_owner_map(ax, dxf_path):
    import ezdxf
    msp = ezdxf.readfile(dxf_path).modelspace()
    for e in msp:
        if e.dxftype() == "HATCH" and e.dxf.layer == L_OWN:
            for p in e.paths:
                ax.fill([v[0] for v in p.vertices], [v[1] for v in p.vertices],
                        color=tuple(c / 255 for c in e.rgb), linewidth=0)
    for e in msp:
        if e.dxftype() == "LWPOLYLINE" and e.dxf.layer == L_CAD:
            pts = [(p[0], p[1]) for p in e.get_points("xy")]
            ax.plot([p[0] for p in pts] + [pts[0][0]], [p[1] for p in pts] + [pts[0][1]],
                    color="black", linewidth=0.5)
    for e in msp:
        if e.dxftype() == "TEXT" and e.dxf.layer == L_OWN:
            ax.text(e.dxf.align_point[0], e.dxf.align_point[1], e.dxf.text,
                    ha="center", va="center", fontsize=9, fontweight="bold")
    ax.set_title("重劃前歸戶圖（填色＋歸戶號）", fontsize=18)
    ax.set_aspect("equal")
    ax.axis("off")


def draw_appx26(ax, xlsx_path):
    import openpyxl
    ws = openpyxl.load_workbook(xlsx_path).worksheets[0]
    cells = [["" if v is None else str(v) for v in r] for r in ws.iter_rows(min_row=2, values_only=True)]
    ncol = max(len(r) for r in cells)
    cells = [r + [""] * (ncol - len(r)) for r in cells]
    ax.axis("off")
    ax.set_title(f"附表 26　{ws['A1'].value}（全表）", fontsize=18, loc="left")
    t = ax.table(cellText=cells, loc="upper left", cellLoc="center")
    t.auto_set_font_size(False)
    t.set_fontsize(8)
    t.scale(1, 1.2)


def main(argv):
    src, out = argv
    os.makedirs(out, exist_ok=True)
    p = os.path.join(out, OUT_NAME)
    fig, axes = plt.subplots(1, 2, figsize=(26, 14), gridspec_kw={"width_ratios": [1.15, 1]})
    draw_owner_map(axes[0], os.path.join(src, "重劃前圖.dxf"))
    draw_appx26(axes[1], os.path.join(src, "附表26_土地所有權人索引表.xlsx"))
    fig.tight_layout()
    fig.savefig(p, dpi=100)
    plt.close(fig)
    sz = os.path.getsize(p)
    print(f"{p}\t{sz} bytes")
    return 0 if sz > 0 else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
