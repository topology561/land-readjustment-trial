# -*- coding: utf-8 -*-
"""
W-P.004 之預覽 PNG（單 §三 6）：附表 20、22 之截取。

命令列：python verify/plots/preview_W-P004.py <成果目錄> <PNG 之輸出目錄>
輸出：W-P.004_附表20-22.png
預覽之標題唯供 KL 一眼辨表，⛔ 屬附表之內容；格之合併與框線以 Excel 開啟為準。
"""
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import font_manager  # noqa: E402

for _f in ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",):
    if os.path.exists(_f):
        font_manager.fontManager.addfont(_f)
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=_f).get_name()

OUT_NAME = "W-P.004_附表20-22.png"
PANELS = [
    ("附表20_公共設施用地負擔面積統計表.xlsx", "附表 20　公共設施用地負擔面積統計表"),
    ("附表22_計算負擔總計表.xlsx", "附表 22（附表 23 同式）　計算負擔總計表"),
]


def _fmt(cell):
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        nf = cell.number_format
        if nf in ("#,##0.00",):
            return f"{v:,.2f}"
        if nf == "#,##0":
            return f"{v:,.0f}"
        if nf == "0.000000":
            return f"{v:.6f}"
        if nf == "0.00%":
            return f"{v * 100:.2f}%"
    return str(v)


def draw(ax, path, title, widths, note=False):
    import openpyxl
    ws = openpyxl.load_workbook(path).worksheets[0]
    last = ws.max_row - 1 if note else ws.max_row
    cells = [[_fmt(ws.cell(r, c)) for c in range(1, ws.max_column + 1)] for r in range(2, last + 1)]
    ax.axis("off")
    ax.set_title(f"{title}　｜　{ws['A1'].value}", fontsize=13, loc="left")
    t = ax.table(cellText=cells, loc="upper left", cellLoc="left", colWidths=widths)
    t.auto_set_font_size(False)
    t.set_fontsize(8)
    t.scale(1, 1.4)
    if note:
        ax.text(0, 0.02, str(ws.cell(ws.max_row, 1).value), fontsize=10, transform=ax.transAxes)


def main(argv):
    src, out = argv
    os.makedirs(out, exist_ok=True)
    p = os.path.join(out, OUT_NAME)
    fig = plt.figure(figsize=(26, 18))
    gs = fig.add_gridspec(2, 1, height_ratios=[0.45, 1.3])
    draw(fig.add_subplot(gs[0]), os.path.join(src, PANELS[0][0]), PANELS[0][1], [0.08, 0.08, 0.16, 0.2, 0.16, 0.12])
    draw(fig.add_subplot(gs[1]), os.path.join(src, PANELS[1][0]), PANELS[1][1],
         [0.05, 0.03, 0.27, 0.07, 0.07, 0.04, 0.03, 0.3, 0.07, 0.04], note=True)
    fig.tight_layout()
    fig.savefig(p, dpi=80)
    plt.close(fig)
    sz = os.path.getsize(p)
    print(f"{p}\t{sz} bytes")
    return 0 if sz > 0 else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
