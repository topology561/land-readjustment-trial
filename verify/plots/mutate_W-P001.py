# -*- coding: utf-8 -*-
"""
W-P.001 之故意錯誤之測（單 §三 4；補令一 §三 4 擴之）：成果須綠、①〜⑦、⑧′、⓪、⑨〜⑬ 共 14 種變異體須各紅。

命令列：python verify/plots/mutate_W-P001.py <成果目錄> <拋棄式之工作目錄>
  成果目錄 ＝ plots/prelim_maps.py --out 之目錄（⛔ 改之；變異體皆寫於工作目錄之複本）。
結束碼：成果綠且 14 種皆紅 ⇒ 0；否則 1。
"""
import importlib.util
import os
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location("check_W_P001", os.path.join(HERE, "check_W-P001.py"))
chk = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(chk)

DXF = "重劃前圖.dxf"
A2 = "附表02_重劃前土地地號摘錄簿.xlsx"
A3 = "附表03_重劃前土地清冊.xlsx"


def _copy(src, dst):
    if os.path.exists(dst):
        shutil.rmtree(dst)
    os.makedirs(dst)
    for f in (DXF, A2, A3):
        shutil.copy2(os.path.join(src, f), os.path.join(dst, f))


def _lot_polys(doc):
    from shapely.geometry import Point, Polygon
    msp = doc.modelspace()
    polys = [Polygon([(p[0], p[1]) for p in e.get_points("xy")])
             for e in msp if e.dxftype() == "LWPOLYLINE" and e.dxf.layer == chk.L_CAD]
    out = {}
    for e in msp:
        if e.dxftype() == "TEXT" and e.dxf.layer == chk.L_CAD:
            pt = Point(e.dxf.align_point[0], e.dxf.align_point[1])
            for pg in polys:
                if pg.contains(pt):
                    out[e.dxf.text] = pg
    return out


def _edit_dxf(d, fn):
    import ezdxf
    p = os.path.join(d, DXF)
    doc = ezdxf.readfile(p)
    fn(doc)
    doc.saveas(p)


def m1_drop_parcel(d):
    def f(doc):
        from shapely.geometry import Point
        pg = _lot_polys(doc)["628-20"]
        msp = doc.modelspace()
        dead = []
        for e in msp:
            if e.dxftype() == "TEXT":
                pt = Point(e.dxf.align_point[0], e.dxf.align_point[1])
            elif e.dxftype() == "LWPOLYLINE":
                from shapely.geometry import Polygon
                pt = Polygon([(p[0], p[1]) for p in e.get_points("xy")]).representative_point()
            elif e.dxftype() == "HATCH":
                pt = chk.hatch_poly(e).representative_point()
            else:
                continue
            if pg.contains(pt):
                dead.append(e)
        for e in dead:
            msp.delete_entity(e)
    _edit_dxf(d, f)


def m2_swap_labels(d):
    def f(doc):
        t = {e.dxf.text: e for e in doc.modelspace()
             if e.dxftype() == "TEXT" and e.dxf.layer == chk.L_CAD}
        t["628-1"].dxf.text, t["628-3"].dxf.text = "628-3", "628-1"
    _edit_dxf(d, f)


def m3_layer_name(d):
    bad = "重劃前歸户圖"

    def f(doc):
        doc.layers.add(bad)
        for e in doc.modelspace():
            if e.dxf.layer == chk.L_OWN:
                e.dxf.layer = bad
        doc.layers.remove(chk.L_OWN)
    _edit_dxf(d, f)


def m4_shift(d):
    from ezdxf.math import Matrix44

    def f(doc):
        m = Matrix44.translate(0.5, 0, 0)
        for e in doc.modelspace():
            e.transform(m)
    _edit_dxf(d, f)


def m5_recolor(d):
    def f(doc):
        lots, owners, hh, zone, price = chk.source_truth()
        polys = _lot_polys(doc)
        cnt = {}
        for ln in lots:
            cnt[hh[ln]] = cnt.get(hh[ln], 0) + 1
        target = sorted(ln for ln in lots if cnt[hh[ln]] >= 2)[0]
        pg = polys[target]
        for e in doc.modelspace():
            if e.dxftype() == "HATCH" and e.dxf.layer == chk.L_OWN \
                    and pg.contains(chk.hatch_poly(e).representative_point()):
                e.rgb = (1, 2, 3)
    _edit_dxf(d, f)


def m6_price_off_by_one(d):
    def f(doc):
        for e in doc.modelspace():
            if e.dxftype() == "TEXT" and e.dxf.layer == chk.L_PRICE:
                e.dxf.text = str(int(e.dxf.text) + 1)
                break
    _edit_dxf(d, f)


def m7_empty(d):
    open(os.path.join(d, DXF), "wb").close()


def _edit_xlsx(d, name, fn):
    import openpyxl
    p = os.path.join(d, name)
    wb = openpyxl.load_workbook(p)
    fn(wb.worksheets[0])
    wb.save(p)


def m8p_drop_coowner_row(d):
    """⑧′ 附表 3 於列 6 起之資料中刪一共有人之列（補令一）。"""
    def f(ws):
        lots = [ws.cell(r, 1).value for r in range(6, ws.max_row + 1)]
        for r, ln in enumerate(lots, start=6):
            if ln is not None and lots.count(ln) >= 2:
                ws.delete_rows(r, 1)
                return
        raise RuntimeError("找不到共有之宗")
    _edit_xlsx(d, A3, f)


OLD_SHA = "f00dabae89c3f84263f35dbf35c1bcdd6668f134"


def m0_old_layout(d):
    """⓪ 以 wip/plots＝f00daba 之 plots/prelim_maps.py 所出之附表 2、3 代之（補令一）。"""
    import subprocess
    repo = os.path.dirname(os.path.dirname(HERE))
    src = subprocess.run(["git", "-C", repo, "show", f"{OLD_SHA}:plots/prelim_maps.py"],
                         check=True, capture_output=True).stdout.decode("utf-8")
    ns = {"__name__": "prelim_maps_f00daba", "__file__": os.path.join(repo, "plots", "prelim_maps.py")}
    exec(compile(src, f"{OLD_SHA}:plots/prelim_maps.py", "exec"), ns)
    lots, owners, area = ns["load_cadastre"]()
    ns["write_appx2"](os.path.join(d, A2), lots)
    ns["write_appx3"](os.path.join(d, A3), lots, owners, area)


def m9_three_groups(d):
    """⑨ 附表 2 一列三組（清列 3 起之 G、H 欄）。"""
    def f(ws):
        for r in range(3, ws.max_row + 1):
            ws.cell(r, 7).value = None
            ws.cell(r, 8).value = None
    _edit_xlsx(d, A2, f)


def m10_no_chief(d):
    """⑩ 附表 2 之簽欄缺「科長：」。"""
    def f(ws):
        for row in ws.iter_rows():
            for c in row:
                if c.value == "科長：":
                    c.value = None
                    return
        raise RuntimeError("找不到「科長：」")
    _edit_xlsx(d, A2, f)


def m11_unmerge_owner(d):
    """⑪ 附表 3 拆 E4:G4 之合併。"""
    _edit_xlsx(d, A3, lambda ws: ws.unmerge_cells("E4:G4"))


def m12_no_print_titles(d):
    """⑫ 附表 3 失列印標題（ws.print_title_rows = None 於 openpyxl 無效，故清 _print_rows）。"""
    def f(ws):
        ws._print_rows = None
    _edit_xlsx(d, A3, f)


def m13_place_in_body(d):
    """⑬ 附表 3 之列 6 之 O 欄填「德安段」。"""
    def f(ws):
        ws.cell(6, 15).value = "德安段"
    _edit_xlsx(d, A3, f)


MUTANTS = [
    ("① 漏一宗（628-20）", m1_drop_parcel),
    ("② 地號錯置（628-1↔628-3）", m2_swap_labels),
    ("③ 圖層名錯（歸戶→歸户）", m3_layer_name),
    ("④ 全圖平移 0.5 m", m4_shift),
    ("⑤ 同戶之一宗改色", m5_recolor),
    ("⑥ 一宗之地價差 1 元", m6_price_off_by_one),
    ("⑦ 空檔", m7_empty),
    ("⑧′ 附表 3 列 6 起刪一共有人之列", m8p_drop_coowner_row),
    ("⓪ 舊版式（f00daba 之附表 2、3）", m0_old_layout),
    ("⑨ 附表 2 一列三組", m9_three_groups),
    ("⑩ 附表 2 簽欄缺「科長：」", m10_no_chief),
    ("⑪ 附表 3 拆 E4:G4 之合併", m11_unmerge_owner),
    ("⑫ 附表 3 失列印標題", m12_no_print_titles),
    ("⑬ 附表 3 列 6 之 O 欄填「德安段」", m13_place_in_body),
]


def verdict(d):
    res = chk.run(os.path.join(d, DXF), os.path.join(d, A3))
    ok = bool(res) and len(res) == 12 and all(x[1] for x in res)
    return ok, [x[0] for x in res if not x[1]]


def main(argv):
    src, work = argv
    rows = []
    base = os.path.join(work, "base")
    _copy(src, base)
    print("== 成果 ==")
    ok, fails = verdict(base)
    rows.append(("成果", "綠" if ok else "紅", fails, ok))
    for name, fn in MUTANTS:
        d = os.path.join(work, f"mut{len(rows)}")
        _copy(src, d)
        fn(d)
        print(f"== {name} ==")
        ok, fails = verdict(d)
        rows.append((name, "綠" if ok else "紅", fails, not ok))
    print("\n| 對象 | 判 | 不過之量測 | 合期望 |")
    print("|---|---|---|---|")
    for n, v, f, good in rows:
        print(f"| {n} | {v} | {'、'.join(f) or '—'} | {'是' if good else '否'} |")
    return 0 if all(r[3] for r in rows) else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
