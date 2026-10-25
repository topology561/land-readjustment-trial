# -*- coding: utf-8 -*-
"""
W-P.003 之預覽 PNG（單 §三 7）：附表 4、14、15、18（前 24 列）、19、21、25 之截取。

命令列：python verify/plots/preview_W-P003.py <成果目錄> <PNG 之輸出目錄>
輸出：W-P.003_附表4-14-15-18-19-21-25.png
預覽之標題唯供 KL 一眼辨表，⛔ 屬附表之內容；格之合併與框線以 Excel 開啟為準。
"""
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import font_manager  # noqa: E402

for _f in ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",):
    if os.path.exists(_f):
        font_manager.fontManager.addfont(_f)
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=_f).get_name()

OUT_NAME = "W-P.003_附表4-14-15-18-19-21-25.png"
PANELS = [
    ("附表04_抵充土地清冊.xlsx", "附表 4　抵充土地清冊", None),
    ("附表19_重劃後地價評議表.xlsx", "附表 19　重劃後地價評議表", None),
    ("附表14_他項權利清冊.xlsx", "附表 14　他項權利清冊", None),
    ("附表21_重劃後可分配土地街廓面積地價計算表.xlsx", "附表 21　重劃後可分配土地街廓面積地價計算表", None),
    ("附表15_限制登記清冊.xlsx", "附表 15　限制登記清冊", None),
    ("附表25_街角地側面道路負擔百分率累積表.xlsx", "附表 25　街角地側面道路負擔百分率累積表", None),
    ("附表18_重劃前地價評議表.xlsx", "附表 18　重劃前地價評議表（前 24 列）", 28),
]


def _fmt(cell):
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        nf = cell.number_format
        if nf == "#,##0.00":
            return f"{v:,.2f}"
        if nf == "0.00":
            return f"{v:.2f}"
        if nf == "#,##0":
            return f"{v:,.0f}"
    return str(v)


def draw(ax, path, title, max_row):
    import openpyxl
    ws = openpyxl.load_workbook(path).worksheets[0]
    last = ws.max_row if max_row is None else min(ws.max_row, max_row)
    cells = [[_fmt(ws.cell(r, c)) for c in range(1, ws.max_column + 1)] for r in range(2, last + 1)]
    ax.axis("off")
    ax.set_title(f"{title}　｜　{ws['A1'].value}", fontsize=13, loc="left")
    t = ax.table(cellText=cells, loc="upper left", cellLoc="center")
    t.auto_set_font_size(False)
    t.set_fontsize(6.5)
    t.scale(1, 1.15)


def main(argv):
    src, out = argv
    os.makedirs(out, exist_ok=True)
    p = os.path.join(out, OUT_NAME)
    fig = plt.figure(figsize=(30, 30))
    gs = fig.add_gridspec(4, 2, height_ratios=[0.6, 0.7, 0.9, 1.6])
    for i, (f, title, mr) in enumerate(PANELS):
        ax = fig.add_subplot(gs[3, :]) if i == 6 else fig.add_subplot(gs[i // 2, i % 2])
        draw(ax, os.path.join(src, f), title, mr)
    fig.tight_layout()
    fig.savefig(p, dpi=80)
    plt.close(fig)
    sz = os.path.getsize(p)
    print(f"{p}\t{sz} bytes")
    return 0 if sz > 0 else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
