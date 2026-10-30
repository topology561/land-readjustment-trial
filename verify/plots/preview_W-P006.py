# -*- coding: utf-8 -*-
"""
W-P.006 之預覽 PNG（單 §三 6）：附表 13（一戶）與附表 24（一張）之截取。

命令列：python verify/plots/preview_W-P006.py <成果目錄> <PNG 之輸出目錄> <退縮…>
輸出：W-P.006_附表13-24_退縮<t>m.png（每一退縮一檔）。
所取：附表 24 中所有權人最多之宗（同者取地號小者）之一張，及附表 13 中該宗所屬之戶之一表。
預覽之標題唯供 KL 一眼辨表，⛔ 屬附表之內容；格之合併與框線以 Excel 開啟為準。
"""
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import font_manager  # noqa: E402

for _f in ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",):
    if os.path.exists(_f):
        font_manager.fontManager.addfont(_f)
        plt.rcParams["font.family"] = font_manager.FontProperties(fname=_f).get_name()

A13 = "附表13_重劃前後土地分配清冊_退縮{}m.xlsx"
A24 = "附表24_土地分配計算表_退縮{}m.xlsx"
OUT = "W-P.006_附表13-24_退縮{}m.png"


def _fmt(cell):
    v = cell.value
    if v is None:
        return ""
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        nf = cell.number_format
        if nf in ("#,##0.00",):
            return f"{v:,.2f}"
        if nf == "#,##0":
            return f"{v:,.0f}"
        if nf == "0.00000000":
            return f"{v:.8f}"
        if nf == "0.0000":
            return f"{v:.4f}"
        if nf == "0.00":
            return f"{v:.2f}"
        if nf == "0.00%":
            return f"{v * 100:.2f}%"
    return str(v)


def pick(wb24):
    """⇒ (表名, 地號)：所有權人最多（代表人「等N人」之 N）者，同者地號小者。"""
    best = None
    for ws in wb24:
        rep = str(ws.cell(2, 4).value or "")
        n = int(rep.rsplit("等", 1)[1][:-1]) if "等" in rep else 1
        no = ws.title[2:]
        key = (-n, tuple(int(p) for p in no.split("-")))
        if best is None or key < best[0]:
            best = (key, ws.title, no)
    return best[1], best[2]


def main(argv):
    import openpyxl
    out, png, tags = argv[0], argv[1], argv[2:]
    os.makedirs(png, exist_ok=True)
    for t in tags:
        t = f"{float(t):g}"
        wb24 = openpyxl.load_workbook(os.path.join(out, A24.format(t)))
        wb13 = openpyxl.load_workbook(os.path.join(out, A13.format(t)))
        name24, no = pick(wb24)
        ws24 = wb24[name24]
        ws13 = next(ws for ws in wb13 if any(ws.cell(r, 15).value == no for r in range(8, ws.max_row + 1)))
        fig = plt.figure(figsize=(26, 20))
        ax1 = fig.add_axes([0.01, 0.52, 0.98, 0.45])
        ax2 = fig.add_axes([0.01, 0.01, 0.6, 0.48])
        last = max(r for r in range(1, ws13.max_row + 1) if any(ws13.cell(r, c).value is not None for c in range(1, 25)))
        cells = [[_fmt(ws13.cell(r, c)) for c in range(1, 25)] for r in range(3, last + 1)]
        ax1.axis("off")
        ax1.set_title(f"附表 13　{ws13.title}（退縮 {t} m）｜　{ws13['A1'].value}", fontsize=14, loc="left")
        tb = ax1.table(cellText=cells, loc="upper left", cellLoc="center")
        tb.auto_set_font_size(False)
        tb.set_fontsize(6.5)
        tb.scale(1, 1.5)
        cells = [[_fmt(ws24.cell(r, c)) for c in range(1, 5)] for r in range(2, ws24.max_row + 1)
                 if any(ws24.cell(r, c).value is not None for c in range(1, 5))]
        ax2.axis("off")
        ax2.set_title(f"附表 24　{name24}（退縮 {t} m）｜　{ws24['A1'].value}", fontsize=14, loc="left")
        tb = ax2.table(cellText=cells, loc="upper left", cellLoc="left", colWidths=[0.28, 0.62, 0.05, 0.05])
        tb.auto_set_font_size(False)
        tb.set_fontsize(8)
        tb.scale(1, 1.25)
        p = os.path.join(png, OUT.format(t))
        fig.savefig(p, dpi=80)
        plt.close(fig)
        print(f"輸出 {p}（{ws13.title}、{name24}）")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
