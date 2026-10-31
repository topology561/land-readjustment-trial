# -*- coding: utf-8 -*-
"""
W-P.007 之預覽 PNG（單 §三 6）：附表 13 之二表依 Excel 之版面（欄寬、列高、合併、直書、分散對齊、粗細框線）繪之，
供 KL 與手冊第 88 頁對看。

命令列：python verify/plots/preview_W-P007.py <成果目錄> <PNG 之輸出目錄> <退縮…>
輸出：W-P.007_附表13版面_退縮<t>m.png（每一退縮一檔）。
所取：有重劃後之列之戶中，資料列最多而 ≤ 8 者（同者歸戶號小者）之一表，及「縣有（抵費地）」之表。
預覽之上方之小標唯供 KL 一眼辨表，⛔ 屬附表之內容；字形為 Noto Sans CJK（Excel 中以使用者之字形為準）。
"""
import importlib.util
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location("preview_W_P006", os.path.join(HERE, "preview_W-P006.py"))
pv = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(pv)
_spec = importlib.util.spec_from_file_location("check_W_P007", os.path.join(HERE, "check_W-P007.py"))
chk = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(chk)
plt = pv.plt
OUT = "W-P.007_附表13版面_退縮{}m.png"
OFFSET_SHEET = "縣有（抵費地）"
PX_PER_W, PX_PER_PT = 7.0, 4.0 / 3.0
LW = {"thin": 0.6, "medium": 1.8}


def _pick(wb):
    best = None
    for ws in wb:
        if not ws.title.startswith("歸戶號"):
            continue
        last = chk.last_row(ws)
        rows = chk.data_rows(ws, chk.FIRST, last)
        if not any(r[14][0] is not None for r in rows) or len(rows) > 8:
            continue
        key = (-len(rows), int(ws.title[3:]))
        if best is None or key < best[0]:
            best = (key, ws.title)
    if best is None:
        raise ValueError("無可取之戶")
    return best[1]


def _draw(ax, ws, raw, x0, y0):
    """自 (x0, y0)（像素·y 向下）繪 ws 之第 1 列至簽欄；⇒ 底之 y。"""
    from openpyxl.utils import get_column_letter
    ncol = chk.NCOL
    last = chk.last_row(ws) + 2
    xs = [x0]
    for c in range(1, ncol + 1):
        xs.append(xs[-1] + ws.column_dimensions[get_column_letter(c)].width * PX_PER_W)
    ys = [y0]
    for r in range(1, last + 1):
        ys.append(ys[-1] + (ws.row_dimensions[r].height or 15) * PX_PER_PT)
    span = {}
    inner = set()
    box = {}
    for m in ws.merged_cells.ranges:
        span[(m.min_row, m.min_col)] = (m.max_row, m.max_col)
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                inner.add((r, c))
                box[(r, c)] = (m.min_row, m.min_col, m.max_row, m.max_col)
        inner.discard((m.min_row, m.min_col))
    for (r, c), (lt, rt, tp, bt) in raw.items():
        if r > last or c > ncol:
            continue
        if (r, c) in box:   # 合併之格之內部之線不繪（Excel 亦不顯示）
            r1, c1, r2, c2 = box[(r, c)]
            lt, rt, tp, bt = (lt if c == c1 else None, rt if c == c2 else None, tp if r == r1 else None,
                              bt if r == r2 else None)
        for st, (xa, ya, xb, yb) in ((lt, (xs[c - 1], ys[r - 1], xs[c - 1], ys[r])), (rt, (xs[c], ys[r - 1], xs[c], ys[r])),
                                     (tp, (xs[c - 1], ys[r - 1], xs[c], ys[r - 1])), (bt, (xs[c - 1], ys[r], xs[c], ys[r]))):
            if st:
                ax.plot([xa, xb], [ya, yb], color="black", lw=LW.get(st, 0.6), solid_capstyle="butt")
    for r in range(1, last + 1):
        for c in range(1, ncol + 1):
            if (r, c) in inner:
                continue
            cell = ws.cell(r, c)
            txt = pv._fmt(cell)
            if not txt:
                continue
            r2, c2 = span.get((r, c), (r, c))
            xa, xb, ya, yb = xs[c - 1], xs[c2], ys[r - 1], ys[r2]
            al = cell.alignment
            fs = 13 if r == 3 else 11 if r == 1 else 6.5
            if al is not None and al.text_rotation == 255:
                n = len(txt)
                step = min((yb - ya - 6) / max(n, 1), 12)
                for i, ch in enumerate(txt):
                    ax.text((xa + xb) / 2, (ya + yb) / 2 + (i - (n - 1) / 2) * step, ch, ha="center", va="center",
                            fontsize=6.5)
            elif al is not None and al.horizontal == "distributed":
                lines = txt.split("\n")
                for j, line in enumerate(lines):
                    yy = (ya + yb) / 2 + (j - (len(lines) - 1) / 2) * 13
                    n = len(line)
                    for i, ch in enumerate(line):
                        xx = xa + 12 + (xb - xa - 24) * (i / max(n - 1, 1))
                        ax.text(xx, yy, ch, ha="center", va="center", fontsize=7)
            else:
                ha = "left" if (al is None or al.horizontal in (None, "general", "left")) and r > 3 and c == 1 \
                    and (r, c) not in span else ("right" if al is not None and al.horizontal == "right" else "center")
                x = xa + 2 if ha == "left" else xb - 2 if ha == "right" else (xa + xb) / 2
                if r == 1:
                    ha, x = "left", xa
                if r > 3:   # 格內之字依格寬折行（Excel 之自動換行）
                    per = max(int((xb - xa - 4) / (fs * 1.45)), 1)
                    txt = "\n".join(txt[i:i + per] for i in range(0, len(txt), per)) if r <= 10 or c == chk.NCOL \
                        else txt
                ax.text(x, (ya + yb) / 2, txt, ha=ha, va="center", fontsize=fs, wrap=False, linespacing=1.15,
                        fontweight="bold" if r == 1 else "normal")
    return ys[-1], xs[-1]


def main(argv):
    import openpyxl
    out, png, tags = argv[0], argv[1], argv[2:]
    os.makedirs(png, exist_ok=True)
    for tag in tags:
        t = f"{float(tag):g}"
        path = os.path.join(out, chk.FMT["a13"].format(t))
        wb = openpyxl.load_workbook(path)
        raw = chk.raw_borders(path)
        names = [_pick(wb), OFFSET_SHEET]
        fig, ax = plt.subplots(figsize=(17, 24))
        y, right = 10, 0
        for nm in names:
            ax.text(0, y + 10, f"【預覽】附表 13　{nm}（退縮 {t} m）", fontsize=10, color="#555555", va="center")
            y, x = _draw(ax, wb[nm], raw[nm], 20, y + 30)
            right = max(right, x)
            y += 40
        ax.set_xlim(0, right + 20)
        ax.set_ylim(y, 0)
        ax.set_aspect("equal")
        ax.axis("off")
        p = os.path.join(png, OUT.format(t))
        fig.savefig(p, dpi=110, bbox_inches="tight")
        plt.close(fig)
        print(f"輸出 {p}（{'、'.join(names)}）")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
