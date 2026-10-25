# -*- coding: utf-8 -*-
"""
W-P.002 之故意錯誤之測（單 §三 4·補令一 §三 4）：成果須綠、Ⓐ〜Ⓗ 共 8 種變異體須各紅（量測 ＝ check_W-P002.py 之 8 項）。

命令列：python verify/plots/mutate_W-P002.py <成果目錄> <拋棄式之工作目錄> <dict-revised.json> <示範多頁之附表26.xlsx>
  成果目錄 ＝ plots/prelim_maps.py --out 之目錄（⛔ 改之；變異體皆寫於工作目錄）。
  示範多頁之附表26.xlsx ＝ verify/plots/synth_W-P002a.py 之輸出（⛔ 改之；複製入各工作目錄）。
  Ⓐ、Ⓓ、Ⓖ 掛換 plots/owner_index 之函式或常數後重出（於本行程之記憶體中·⛔ 改檔）；Ⓑ、Ⓒ、Ⓔ、Ⓕ、Ⓗ 改成果之複本。
結束碼：成果綠且 8 種皆紅 ⇒ 0；否則 1。
"""
import importlib.util
import os
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
_spec = importlib.util.spec_from_file_location("check_W_P002", os.path.join(HERE, "check_W-P002.py"))
chk = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(chk)

FILES = (chk.FILE_DXF, chk.FILE_COLOR, chk.FILE_A26)
SYNTH_NAME = "示範多頁之附表26.xlsx"
SYNTH_SRC = None   # main 設之
MAINLAND = {"草": 9, "蓮": 13, "葉": 12, "邱": 7, "阮": 6}   # 大陸之筆劃（單 §三 4 Ⓐ）


def _copy(src, dst):
    if os.path.exists(dst):
        shutil.rmtree(dst)
    os.makedirs(dst)
    for f in FILES:
        shutil.copy2(os.path.join(src, f), os.path.join(dst, f))
    shutil.copy2(SYNTH_SRC, os.path.join(dst, SYNTH_NAME))


def _fresh(dst):
    if os.path.exists(dst):
        shutil.rmtree(dst)
    os.makedirs(dst)
    shutil.copy2(SYNTH_SRC, os.path.join(dst, SYNTH_NAME))


def _prelim():
    sys.path.insert(0, os.path.join(REPO, "plots"))
    sys.path.insert(0, os.path.join(REPO, "verify"))
    import prelim_maps
    return prelim_maps


def _regen(d, moe_path, attr, repl):
    pm = _prelim()
    oi = pm.owner_index
    orig = getattr(oi, attr)
    setattr(oi, attr, repl(orig))
    try:
        _fresh(d)
        pm.main(["--out", d, "--moedict", moe_path])
    finally:
        setattr(oi, attr, orig)


def mA_mainland_strokes(d, moe_path):
    def repl(orig):
        def load(path):
            tab = dict(orig(path))
            for c, s in MAINLAND.items():
                tab[c] = (s, tab[c][1])
            return tab
        return load
    _regen(d, moe_path, "load_strokes", repl)


def mB_wrong_number(d, moe_path):
    import ezdxf
    from shapely.geometry import Point, Polygon
    p = os.path.join(d, chk.FILE_DXF)
    doc = ezdxf.readfile(p)
    msp = doc.modelspace()
    polys = [Polygon([(q[0], q[1]) for q in e.get_points("xy")])
             for e in msp if e.dxftype() == "LWPOLYLINE" and e.dxf.layer == chk.L_CAD]
    target = None
    for e in msp:
        if e.dxftype() == "TEXT" and e.dxf.layer == chk.L_CAD and e.dxf.text == "628-26":
            pt = Point(e.dxf.align_point[0], e.dxf.align_point[1])
            target = [pg for pg in polys if pg.contains(pt)][0]
    hit = [e for e in msp if e.dxftype() == "TEXT" and e.dxf.layer == chk.L_OWN
           and target.contains(Point(e.dxf.align_point[0], e.dxf.align_point[1]))]
    if len(hit) != 1 or hit[0].dxf.text != "1":
        raise RuntimeError(f"628-26 之歸戶號之註記非恰一個「1」：{[e.dxf.text for e in hit]}")
    hit[0].dxf.text = "2"
    doc.saveas(p)


def _read_seq(ws):
    """依補令一之排列（每頁先左組、再右組）讀回附表 26 之序。"""
    last = chk._last_row(ws)
    seq, i = [], 0
    while True:
        r, j = chk.a26_slot(i)
        if r >= last:
            break
        nm = ws.cell(r, 1 + 4 * j).value
        if not chk._blank(nm):
            seq.append((nm, ws.cell(r, 4 + 4 * j).value))
        i += 1
    return seq, last


def _rewrite(ws, seq, last, slot):
    for r in range(4, last):
        for j in range(2):
            ws.cell(r, 1 + 4 * j).value = None
            ws.cell(r, 4 + 4 * j).value = None
    for i, (nm, no) in enumerate(seq):
        r, j = slot(i)
        ws.cell(r, 1 + 4 * j, nm)
        ws.cell(r, 4 + 4 * j, no)


def mC_drop_index_row(d, moe_path):
    import openpyxl
    p = os.path.join(d, chk.FILE_A26)
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    seq, last = _read_seq(ws)
    hits = [i for i, (nm, _) in enumerate(seq) if nm == "沈昌穎"]
    if len(hits) < 2:
        raise RuntimeError(f"沈昌穎 之列不足二：{hits}")
    del seq[hits[1]]
    _rewrite(ws, seq, last, chk.a26_slot)
    wb.save(p)


def mF_row_major(d, moe_path):
    """舊式：由左而右、再由上而下（KL 10:22 前之排列）。"""
    import openpyxl
    p = os.path.join(d, chk.FILE_A26)
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    seq, last = _read_seq(ws)
    if len(seq) != chk.N_INDEX:
        raise RuntimeError(f"讀回 {len(seq)} 筆≠{chk.N_INDEX}")
    _rewrite(ws, seq, last, lambda i: (4 + i // 2, i % 2))
    wb.save(p)


def mG_per_page_30(d, moe_path):
    """每頁每欄之列數錯（30）：掛換常數後重出示範多頁之附表 26。"""
    sys.path.insert(0, os.path.join(REPO, "plots"))
    import owner_index as oi
    orig = oi.APPX26_PER_PAGE
    oi.APPX26_PER_PAGE = 30
    try:
        oi.write_appx26(os.path.join(d, SYNTH_NAME), list(chk.SYNTH))
    finally:
        oi.APPX26_PER_PAGE = orig


def mH_no_breaks(d, moe_path):
    """示範多頁之附表 26 刪去手動分頁符。"""
    import openpyxl
    p = os.path.join(d, SYNTH_NAME)
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    if not ws.row_breaks.brk:
        raise RuntimeError("示範之附表 26 本無分頁符")
    ws.row_breaks.brk = []
    wb.save(p)


def mD_no_groups(d, moe_path):
    _regen(d, moe_path, "group_of", lambda orig: (lambda row: 0))


def mE_no_section_chief(d, moe_path):
    import openpyxl
    p = os.path.join(d, chk.FILE_A26)
    wb = openpyxl.load_workbook(p)
    ws = wb.worksheets[0]
    for row in ws.iter_rows():
        for c in row:
            if c.value == "股長:":
                c.value = None
                wb.save(p)
                return
    raise RuntimeError("找不到「股長:」")


MUTANTS = [
    ("Ⓐ 大陸筆劃（草 9、蓮 13、葉 12、邱 7、阮 6）", mA_mainland_strokes, False),
    ("Ⓑ 歸戶圖 628-26 之歸戶號 1→2", mB_wrong_number, True),
    ("Ⓒ 附表 26 漏沈昌穎之第二戶", mC_drop_index_row, True),
    ("Ⓓ 不分群（group_of 恆 0）", mD_no_groups, False),
    ("Ⓔ 附表 26 簽欄缺「股長:」", mE_no_section_chief, True),
    ("Ⓕ 附表 26 回舊式（逐列由左而右）", mF_row_major, True),
    ("Ⓖ 每頁每欄 30 列（示範多頁）", mG_per_page_30, True),
    ("Ⓗ 示範多頁之附表 26 無分頁符", mH_no_breaks, True),
]


def verdict(d, moe_path):
    res = chk.run(d, moe_path, os.path.join(d, SYNTH_NAME))
    ok = len(res) == chk.N_Q and all(x[1] for x in res)
    return ok, [x[0] for x in res if not x[1]]


def main(argv):
    global SYNTH_SRC
    src, work, moe_path, SYNTH_SRC = argv
    rows = []
    base = os.path.join(work, "base")
    _copy(src, base)
    print("== 成果 ==")
    ok, fails = verdict(base, moe_path)
    rows.append(("成果", "綠" if ok else "紅", fails, ok))
    for name, fn, from_copy in MUTANTS:
        d = os.path.join(work, f"mut{len(rows)}")
        if from_copy:
            _copy(src, d)
        fn(d, moe_path)
        print(f"== {name} ==")
        ok, fails = verdict(d, moe_path)
        rows.append((name, "綠" if ok else "紅", fails, not ok))
    print("\n| 對象 | 判 | 不過之量測 | 合期望 |")
    print("|---|---|---|---|")
    for n, v, f, good in rows:
        print(f"| {n} | {v} | {'、'.join(f) or '—'} | {'是' if good else '否'} |")
    return 0 if all(r[3] for r in rows) else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main(sys.argv[1:]))
