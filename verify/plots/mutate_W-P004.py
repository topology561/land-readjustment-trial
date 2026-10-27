# -*- coding: utf-8 -*-
"""
W-P.004 之故意錯誤之測（單 §三 4）：成果須綠、Ⓐ〜Ⓕ 共 6 種變異體須各紅（量測 ＝ check_W-P004.py 之 8 項）。

命令列：python verify/plots/mutate_W-P004.py <成果目錄> <拋棄式之工作目錄>
  成果目錄 ＝ plots/burden_tables.py --out 之目錄（⛔ 改之；變異體皆寫於工作目錄）。
  Ⓑ、Ⓓ 掛換 plots/ 之函式後重出（於本行程之記憶體中·⛔ 改檔）；他者改成果之複本。
結束碼：成果綠且 6 種皆紅 ⇒ 0；否則 1。
"""
import importlib.util
import os
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
_spec = importlib.util.spec_from_file_location("check_W_P004", os.path.join(HERE, "check_W-P004.py"))
chk = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(chk)


def _copy(src, dst):
    if os.path.exists(dst):
        shutil.rmtree(dst)
    shutil.copytree(src, dst)


def _book(d, n):
    import openpyxl
    p = os.path.join(d, chk.FILES[n])
    return p, openpyxl.load_workbook(p)


def _plots():
    sys.path.insert(0, os.path.join(REPO, "plots"))
    import burden_tables
    return burden_tables


def _regen(d, obj, attr, val):
    orig = getattr(obj, attr)
    setattr(obj, attr, val)
    try:
        if os.path.exists(d):
            shutil.rmtree(d)
        _plots().main(["--out", d])
    finally:
        setattr(obj, attr, orig)


def mA_drop_block(d):
    p, wb = _book(d, 20)
    ws = wb.worksheets[0]
    r = chk._find(ws, "G1", 1)
    ws.delete_rows(r)
    wb.save(p)


def mB_no_one_minus_c(d):
    bt = _plots()
    orig = bt.a20_rows

    def rows(pub, raw, one_minus_c):
        return orig(pub, raw, 1)
    _regen(d, bt, "a20_rows", rows)


def mC_round_after_sum(d):
    """小計之特別負擔改為「未乘之和 ×（1−C）再取位」（957.57），一般負擔隨之（11188.99）。"""
    p, wb = _book(d, 20)
    ws = wb.worksheets[0]
    r = chk._find(ws, "小　計", 2)
    ws.cell(r, 4, 957.57)
    ws.cell(r, 5, 11188.99)
    wb.save(p)


def mD_map_area_pre_price(d):
    """第五項改用圖面面積計之總地價與平均地價（問五 甲）。"""
    bt = _plots()
    orig = bt.a22_values

    def vals(v, reg_area, n_owner, pre_total, sub20):
        x = orig(v, reg_area, n_owner, pre_total, sub20)
        x["前總"], x["前均"] = x["_前總圖面"], x["_前均圖面"]
        return x
    _regen(d, bt, "a22_values", vals)


def mE_core_rise(d):
    """上漲率改用核心線之重劃前平均地價（÷ 38,498.68 ＝ 1.871716·KL 23:20 之乙）。"""
    p, wb = _book(d, 22)
    ws = wb.worksheets[0]
    r = chk._find(ws, "重劃前後每平方公尺平均地價上漲率：", 8)
    ws.cell(r, 9, 1.871716)
    wb.save(p)


def mF_drop_note(d):
    p, wb = _book(d, 22)
    ws = wb.worksheets[0]
    ws.cell(chk._last(ws), 1).value = None
    wb.save(p)


MUTANTS = [
    ("Ⓐ 附表 20 漏一區塊", mA_drop_block, True),
    ("Ⓑ 臨街地特別負擔未乘（1−C）", mB_no_one_minus_c, False),
    ("Ⓒ 特別負擔先加總後取位", mC_round_after_sum, True),
    ("Ⓓ 第五項用圖面面積之總地價", mD_map_area_pre_price, False),
    ("Ⓔ 上漲率用核心線之重劃前平均地價", mE_core_rise, True),
    ("Ⓕ 漏表下之註", mF_drop_note, True),
]


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 2:
        print(__doc__)
        return 2
    out, work = argv
    src = chk.load_src()
    base = chk.run(out, src, quiet=True)
    print(f"成果：{'綠' if not base else '紅 ' + str(base)}")
    ok = not base
    for i, (name, fn, from_copy) in enumerate(MUTANTS):
        d = os.path.join(work, f"m{i}")
        if from_copy:
            _copy(out, d)
        fn(d)
        fails = chk.run(d, src, quiet=True)
        print(f"{name}：{'紅 ' + str(fails) if fails else '綠（未捕）'}")
        ok = ok and bool(fails)
    print("合" if ok else "不合")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
