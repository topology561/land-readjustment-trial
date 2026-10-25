# -*- coding: utf-8 -*-
"""
W-P.002（含補令一）：歸戶號（重劃前歸戶圖之註記）、顏色對照表之歸戶號、附表 26 土地所有權人索引表之量測。
獨立於 plots/：⛔ import 之；期望取自倉內原始資料、教育部筆劃之資料（dict-revised.json）、手冊與 KL 之答。

命令列：python verify/plots/check_W-P002.py <成果目錄> <dict-revised.json> <示範多頁之附表26.xlsx>
  成果目錄 ＝ plots/prelim_maps.py --out 之目錄（讀 重劃前圖.dxf、顏色對照表_重劃前.xlsx、附表26_土地所有權人索引表.xlsx）。
  示範多頁之附表26.xlsx ＝ verify/plots/synth_W-P002a.py 之輸出（170 筆之示範名·驗跨頁之排列）。
任一不過 ⇒ 結束碼 1。Q0 筆劃之資料；Q1/Q1b 歸戶圖之歸戶號；Q2〜Q4 附表 26；Q5 顏色對照表；Q6 跨頁之排列。
排序之所據：KL 2026-10-23 21:18（單獨在前）、2026-10-25 02:30（一甲 二乙 三甲 四可·甲）、02:33（共有戶以自然人為主）。
排列之所據（W-P.002_補令一）：KL 2026-10-25 10:22「是，索引表改為先左欄由上而下、再右欄」——每頁先左組由上而下排滿
  A26_PER 列、再右組由上而下；次頁接續；每頁之末一列之後置手動分頁符。
"""
import functools
import hashlib
import json
import os
import sys
import unicodedata
from collections import defaultdict

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
XLSX_SRC = os.path.join(REPO, "data", "地籍資料來源_匿名版.xlsx")
MOE_SHA = "f68fc3df140a64d8de6eb905ff6576bfd9d260eed0bbf5cc9c41aba5a131a7ad"
# 教育部標準字體之筆劃（KL 2026-10-24 11:43「甲」；大陸之值：草 9、蓮 13、葉 12、蔡 14、鄭 14、邱 7、阮 6）
STROKE_ANCHORS = {"草": 10, "花": 8, "蓮": 15, "連": 11, "進": 12, "葉": 13, "蔡": 15, "鄭": 15,
                  "邱": 8, "郭": 11, "阮": 7, "余": 7, "沈": 7}
# 發單端依 KL 之答手排之錨（地號 → 歸戶號）
HH_ANCHORS = {"628-26": 1, "628": 2, "628-2": 3, "628-51": 19, "713": 20, "628-3": 21,
              "628-35": 22, "628-39": 23, "628-40": 24, "628-15": 33}
FILE_DXF = "重劃前圖.dxf"
FILE_COLOR = "顏色對照表_重劃前.xlsx"
FILE_A26 = "附表26_土地所有權人索引表.xlsx"
L_CAD, L_OWN = "重劃前地籍線", "重劃前歸戶圖"
A26_MERGES = {"A1:H1", "A2:A3", "B2:C2", "D2:D3", "E2:E3", "F2:G2", "H2:H3"}
A26_ROW2 = ["所有權人姓名", "分配清冊", None, "附註"] * 2
A26_ROW3 = [None, "第 冊", "第 頁", None] * 2
A26_SIGN = ["承辦:", "股長:", "科長:", "製表日期: 年 月 日"]
COLOR_HEAD = ["圖層", "歸戶號", "顏色 R,G,B", "所有權人", "地號"]
N_HH, N_INDEX = 33, 67
A26_PER = 40                       # 每頁每欄之列數（W-P.002_補令一 §二 1）
SYNTH = [(f"示範{i:03d}", i) for i in range(1, 171)]   # synth_W-P002a.py 之輸入（170 筆 ⇒ 3 頁：80、80、10）


def a26_slot(i, per=A26_PER):
    p, k = divmod(i, 2 * per)
    return 4 + p * per + k % per, k // per


def a26_layout(ws, seq, per=A26_PER):
    """附表 26 之資料區（列 4 起）之逐格比對、簽欄之列、手動分頁符。⇒ (ok, 訊息)"""
    want = {}
    for i, (nm, no) in enumerate(seq):
        want[a26_slot(i, per)] = (nm, no)
    body_last = max(r for r, _ in want)
    last = _last_row(ws)
    bad = []
    for r in range(4, last):
        for j in range(2):
            got = tuple(ws.cell(r, 1 + 4 * j + d).value for d in range(4))
            if (r, j) in want:
                nm, no = want[(r, j)]
                ok = str(got[0]) == nm and got[3] == no and _blank(got[1]) and _blank(got[2])
            else:
                ok = all(_blank(v) for v in got)
            if not ok:
                bad.append(((r, j), got[0], got[3], want.get((r, j))))
    n_pages = (len(seq) + 2 * per - 1) // (2 * per)
    brk_want = [3 + p * per for p in range(1, n_pages)]
    brk_got = sorted(b.id for b in ws.row_breaks.brk)
    ok = not bad and last == body_last + 2 and brk_got == brk_want
    return ok, (f"筆 {len(seq)}；不符 {len(bad)} {bad[:3]}；資料末列 {body_last}、簽欄列 {last}（期望 {body_last + 2}）；"
                f"分頁符 {brk_got}（期望 {brk_want}）")


def _blank(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", "null"))


def lkey(s):
    p = str(s).split("-")
    return (int(p[0]), int(p[1]) if len(p) > 1 else 0)


@functools.lru_cache(maxsize=None)
def moe(path):
    raw = open(path, "rb").read()
    sha = hashlib.sha256(raw).hexdigest()
    rad_no = {}
    for i in range(214):
        rad_no[unicodedata.normalize("NFKC", chr(0x2F00 + i))] = i + 1
    rad_no["青"] = 174
    tab = {}
    for e in json.loads(raw):
        t = e.get("title", "")
        if len(t) == 1:
            tab[t] = (int(e["stroke_count"]), rad_no[e["radical"].strip()])
    return sha, tab


def ckey(c, tab):
    if unicodedata.name(c, "").startswith(("CJK UNIFIED IDEOGRAPH", "CJK COMPATIBILITY IDEOGRAPH")):
        s, r = tab[c]          # 缺 ⇒ KeyError ⇒ 量測不過（⛔ 補值）
        return (0, s, r, ord(c))
    return (1, ord(c), 0, 0)


def nkey(name, tab):
    return tuple(ckey(c, tab) for c in name)


def grp(kind, pub):
    kind, pub = str(kind).strip(), str(pub).strip()
    if pub == "公":
        return 2
    if pub == "私":
        return {"1": 0, "7": 1}[kind]
    raise KeyError((kind, pub))


@functools.lru_cache(maxsize=None)
def truth(moe_path):
    import openpyxl
    _, tab = moe(moe_path)
    wb = openpyxl.load_workbook(XLSX_SRC, read_only=True)
    rows = list(wb["重劃區地籍"].iter_rows(values_only=True))
    il = rows[0].index("地號")
    lots = sorted((str(r[il]).strip() for r in rows[1:] if not _blank(r[il])), key=lkey)
    ls = set(lots)
    own = defaultdict(dict)       # 地號 → 登記次序 → (統編, 姓名, 群)
    for r in wb["U_LAND"].iter_rows(values_only=True):
        if r[0] != "0013德安段" or _blank(r[1]):
            continue
        m, s = divmod(int(r[1]), 10000)
        ln = f"{m}" if s == 0 else f"{m}-{s}"
        if ln in ls:
            own[ln].setdefault(int(r[23]), (str(r[32]), str(r[33]), grp(r[37], r[38])))
    hh_lots = defaultdict(list)
    hh_mem = {}
    for ln in lots:
        k = frozenset(own[ln].values())
        hh_lots[k].append(ln)
        hh_mem[k] = sorted(k)
    singles, joints = [], []
    for k, mem in hh_mem.items():
        first = min(lkey(x) for x in hh_lots[k])
        if len(mem) == 1:
            _, nm, g = mem[0]
            singles.append(((g, nkey(nm, tab), first), k))
        else:
            heads = [nkey(nm, tab) for _, nm, g in mem if g == 0]
            if not heads:
                raise ValueError(f"共有戶無自然人 {mem}")
            joints.append(((len(mem), min(heads), first), k))
    seq = [k for _, k in sorted(singles)] + [k for _, k in sorted(joints)]
    no_of_hh = {k: i + 1 for i, k in enumerate(seq)}
    no_of_lot = {ln: no_of_hh[k] for k, v in hh_lots.items() for ln in v}
    idx = sorted(((g, nkey(nm, tab), no_of_hh[k]), nm) for k, mem in hh_mem.items() for _, nm, g in mem)
    index = [(nm, key[2]) for key, nm in idx]
    lots_of_no = {no_of_hh[k]: sorted(v, key=lkey) for k, v in hh_lots.items()}
    return lots, no_of_lot, index, lots_of_no


def _last_row(ws):
    return max(r for r in range(1, ws.max_row + 1)
               if any(not _blank(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)))


def run(out_dir, moe_path, synth_path):
    import openpyxl
    res = []

    def rec(n, ok, msg=""):
        res.append((n, bool(ok), msg))
        print(f"{'PASS' if ok else 'FAIL'} {n} {msg}")

    sha, tab = moe(moe_path)
    got = {c: tab.get(c, (None,))[0] for c in STROKE_ANCHORS}
    rec("Q0", sha == MOE_SHA and got == STROKE_ANCHORS, f"sha256 {sha[:12]}…；筆劃 {got}")
    lots, no_of_lot, index, lots_of_no = truth(moe_path)
    anc = {ln: no_of_lot.get(ln) for ln in HH_ANCHORS}

    # Q1 歸戶圖之歸戶號
    try:
        import ezdxf
        from shapely.geometry import Point, Polygon
        msp = ezdxf.readfile(os.path.join(out_dir, FILE_DXF)).modelspace()
        polys = [Polygon([(p[0], p[1]) for p in e.get_points("xy")])
                 for e in msp if e.dxftype() == "LWPOLYLINE" and e.dxf.layer == L_CAD and e.closed]
        lot_poly = {}
        for e in msp:
            if e.dxftype() == "TEXT" and e.dxf.layer == L_CAD:
                pt = Point(e.dxf.align_point[0], e.dxf.align_point[1])
                hit = [i for i, pg in enumerate(polys) if pg.contains(pt)]
                if len(hit) == 1:
                    lot_poly[e.dxf.text] = hit[0]
        on = defaultdict(list)
        stray = []
        for e in msp:
            if e.dxftype() == "TEXT" and e.dxf.layer == L_OWN:
                pt = Point(e.dxf.align_point[0], e.dxf.align_point[1])
                hit = [ln for ln, i in lot_poly.items() if polys[i].contains(pt)]
                if len(hit) == 1:
                    on[hit[0]].append(e.dxf.text)
                else:
                    stray.append(e.dxf.text)
        wrong = [(ln, on.get(ln), no_of_lot[ln]) for ln in lots if on.get(ln) != [str(no_of_lot[ln])]]
        rec("Q1", set(lot_poly) == set(lots) and not wrong and not stray
            and sorted(set(no_of_lot.values())) == list(range(1, N_HH + 1)),
            f"註記 {sum(len(v) for v in on.values())}（期望 {len(lots)}）；不符 {wrong[:4]}；落宗外 {stray[:3]}")
    except Exception as e:  # noqa: BLE001
        rec("Q1", False, f"DXF 讀取失敗 {type(e).__name__}: {e}")
    rec("Q1b", anc == HH_ANCHORS, f"錨 {anc}")

    # Q2〜Q4 附表 26
    try:
        wb = openpyxl.load_workbook(os.path.join(out_dir, FILE_A26))
        if len(wb.worksheets) != 1:
            raise ValueError(f"工作表數 {len(wb.worksheets)}≠1")
        ws = wb.worksheets[0]
        merges = {str(m) for m in ws.merged_cells.ranges}
        r2 = [ws.cell(2, c).value for c in range(1, 9)]
        r3 = [ws.cell(3, c).value for c in range(1, 9)]
        rec("Q2", ws["A1"].value == "土地所有權人索引表" and merges == A26_MERGES and r2 == A26_ROW2
            and r3 == A26_ROW3 and str(ws.print_title_rows).replace("$", "") == "1:3" and ws.max_column == 8,
            f"題 {ws['A1'].value!r}；併格差 {sorted(merges ^ A26_MERGES)}；列2 {r2}；列3 {r3}；列印標題 {ws.print_title_rows}；欄數 {ws.max_column}")
        last = _last_row(ws)
        ok3, msg3 = a26_layout(ws, index)
        rec("Q3", ok3 and len(index) == N_INDEX, msg3)
        vals = [ws.cell(last, c).value for c in range(1, ws.max_column + 1) if not _blank(ws.cell(last, c).value)]
        rec("Q4", vals == A26_SIGN, f"末列 {vals}")
    except Exception as e:  # noqa: BLE001
        rec("Q2", False, f"附表 26 讀取失敗 {type(e).__name__}: {e}")

    # Q5 顏色對照表之歸戶號
    try:
        import ezdxf
        from shapely.geometry import Polygon
        wb = openpyxl.load_workbook(os.path.join(out_dir, FILE_COLOR), read_only=True)
        rows = list(wb["重劃前歸戶圖"].iter_rows(values_only=True))
        head, body = list(rows[0]), [r for r in rows[1:] if any(not _blank(v) for v in r)]
        nos = [r[1] for r in body]
        lot_ok = all(sorted(str(r[4]).split("、"), key=lkey) == lots_of_no.get(r[1]) for r in body)
        hcol = {}
        for e in msp:
            if e.dxftype() == "HATCH" and e.dxf.layer == L_OWN:
                rp = Polygon([(v[0], v[1]) for v in e.paths[0].vertices]).representative_point()
                hit = [ln for ln, i in lot_poly.items() if polys[i].contains(rp)]
                if len(hit) == 1:
                    hcol[hit[0]] = "{},{},{}".format(*e.rgb)
        col_bad = [r[1] for r in body if any(hcol.get(ln) != r[2] for ln in str(r[4]).split("、"))]
        rec("Q5", head == COLOR_HEAD and nos == list(range(1, N_HH + 1)) and lot_ok and not col_bad,
            f"表頭 {head}；歸戶號 {nos[:3]}…{nos[-2:]}（{len(nos)}）；地號合 {lot_ok}；色與圖不符 {col_bad[:3]}")
    except Exception as e:  # noqa: BLE001
        rec("Q5", False, f"顏色對照表讀取失敗 {type(e).__name__}: {e}")

    # Q6 跨頁之排列（示範 170 筆·synth_W-P002a.py）
    try:
        ws = openpyxl.load_workbook(synth_path).worksheets[0]
        ok6, msg6 = a26_layout(ws, SYNTH)
        vals = [ws.cell(_last_row(ws), c).value for c in range(1, 9) if not _blank(ws.cell(_last_row(ws), c).value)]
        rec("Q6", ok6 and vals == A26_SIGN and str(ws.print_title_rows).replace("$", "") == "1:3",
            f"{msg6}；末列 {vals}；列印標題 {ws.print_title_rows}")
    except Exception as e:  # noqa: BLE001
        rec("Q6", False, f"示範多頁之附表 26 讀取失敗 {type(e).__name__}: {e}")
    return res


N_Q = 8

if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    r = run(*sys.argv[1:4])
    sys.exit(0 if len(r) == N_Q and all(x[1] for x in r) else 1)
